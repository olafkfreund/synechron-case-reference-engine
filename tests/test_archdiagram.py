import io
import json
import re
from pathlib import Path
from xml.etree import ElementTree

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.archdiagram import H, W, layout, measure, seg_hits_rect, to_mermaid, to_png, to_pptx_shapes, to_svg, wrap
from app.schema import Architecture

ZONES = ["channels", "experience", "integration", "core", "data", "external", "platform"]
KINDS = ["user", "app", "service", "queue", "store", "external", "platform"]
FIX = Path(__file__).parent / "fixtures" / "arch"
SAMPLES = {p.stem: Architecture.model_validate_json(p.read_text()) for p in sorted(FIX.glob("*.json"))}


def arch(zones, n_flows=0, role="Handles the payment rail and settles every batch overnight, retried", name="Payments hub", **kw):
    """One component per entry in `zones`; flows spread across the diagram, platform ones included."""
    comps = [{"id": f"c{i}", "name": f"{name} {i}", "role": role, "zone": z, "kind": KINDS[i % 7]} for i, z in enumerate(zones)]
    modes = ["api", "event", "batch", "file", "ui", "data"]
    flows = [{"src": f"c{i % len(zones)}", "dst": f"c{(i * 5 + 3) % len(zones)}", "label": f"label number {i}", "mode": modes[i % 6]}
             for i in range(n_flows) if i % len(zones) != (i * 5 + 3) % len(zones)]
    return Architecture.model_validate({"components": comps, "flows": flows,
                                        "callouts": [{"text": "Cut onboarding from 12 days to 2", "outcome": 0}], **kw})


FULL = ["channels", "channels", "experience", "integration", "integration", "core", "core", "core", "data", "data", "external", "platform"]
LONG = arch(["channels", "core", "core", "data", "platform", "platform"], 6, name="Core banking ledger and posting engine",
            role="x" * 5 + " Handles the payment rail and settles every batch overnight",
            title="A very long architecture title about payments modernisation across the bank",
            )
CASES = {
    "full": arch(FULL, 14),
    "one-zone": arch(["core"] * 12, 14),
    "two-zones": arch(["channels"] * 6 + ["external"] * 6, 14),
    "platform-only": arch(["platform"] * 12, 6),
    "six-columns": arch(["channels", "experience", "integration", "core", "data", "external"] * 2, 14),
    "tiny": arch(["core"]),
    "long-text": LONG,
    **{f"sample-{k}": v for k, v in SAMPLES.items()},
}
FOOTER = "Source: Demo reference case with a rather long case name · reviewed 2026-01-01 and signed off by the delivery lead"


def cards(sc):
    return [b for b in sc.boxes if b.style.startswith("card-")]


def pills(sc):
    return [b for b in sc.boxes if b.style.startswith(("pill-", "marker-"))]


def test_sample_sizes():
    assert {k: (len(a.components), len(a.flows)) for k, a in SAMPLES.items()} == {"small": (5, 5), "typical": (9, 10), "max": (12, 14)}


def test_same_input_same_scene():
    assert layout(CASES["full"], FOOTER) == layout(CASES["full"], FOOTER)


@pytest.mark.parametrize("name", CASES)
def test_cards_do_not_overlap_and_everything_is_inside_the_canvas(name):
    sc = layout(CASES[name], FOOTER)
    assert len(cards(sc)) == len(CASES[name].components)
    cs = cards(sc)
    for n, a in enumerate(cs):
        for b in cs[n + 1:]:
            assert a.x + a.w <= b.x or b.x + b.w <= a.x or a.y + a.h <= b.y or b.y + b.h <= a.y, (a.ref, b.ref)
    for b in [*sc.zones, *sc.boxes]:
        assert 0 <= b.x and b.x + b.w <= W and 0 <= b.y and b.y + b.h <= H, b
    for ln in sc.lines:
        assert all(0 <= x <= W and 0 <= y <= H for x, y in ln.points), ln
    for t in sc.texts:
        assert 0 <= t.x <= W and 0 <= t.y <= H, t


@pytest.mark.parametrize("name", CASES)
def test_nothing_is_truncated_and_text_stays_inside_its_card(name):
    sc = layout(CASES[name], FOOTER)
    assert not [t.text for t in sc.texts if "…" in t.text]
    for c in cards(sc):
        for t in sc.texts:
            if c.x + 40 <= t.x <= c.x + c.w and c.y <= t.y <= c.y + c.h and t.style in ("name", "role", "role-sm", "chip", "note"):
                assert t.x + measure(t.text, t.style) * t.k <= c.x + c.w - 4, t
                assert t.y + 14 * t.k <= c.y + c.h, t  # the card grew to fit


@pytest.mark.parametrize("name", CASES)
def test_every_flow_label_is_visible(name):
    sc = layout(CASES[name], FOOTER)
    seen = " ".join(" ".join(t.text.split()) for t in sc.texts)
    for f in CASES[name].flows:
        assert f.label in seen, f.label  # a pill, a keyed marker, or a note on a platform chip


@pytest.mark.parametrize("name", CASES)
def test_no_line_crosses_a_card_it_does_not_connect(name):
    sc = layout(CASES[name], FOOTER)
    by = {b.ref: (b.x, b.y, b.w, b.h) for b in cards(sc)}
    for ln in sc.lines:
        if not ln.refs[0]:
            continue
        for a, b in zip(ln.points, ln.points[1:]):
            for ref, r in by.items():
                if ref not in ln.refs:
                    assert not seg_hits_rect(a, b, r), (ln.refs, ref)


@pytest.mark.parametrize("name", CASES)
def test_pills_overlap_neither_each_other_nor_cards(name):
    sc = layout(CASES[name], FOOTER)
    ps, cs = pills(sc), cards(sc)
    box = lambda b: (b.x, b.y, b.w, b.h)
    hit = lambda a, b: not (a.x + a.w <= b.x or b.x + b.w <= a.x or a.y + a.h <= b.y or b.y + b.h <= a.y)
    for n, a in enumerate(ps):
        assert not any(hit(a, c) for c in cs), a
        assert not any(hit(a, b) for b in ps[n + 1:]), a


@pytest.mark.parametrize("name", CASES)
def test_zones_hug_their_cards_and_platform_lines_stay_in_the_band(name):
    sc = layout(CASES[name], FOOTER)
    cs = cards(sc)
    for z in sc.zones:
        if z.ref == "platform":
            continue
        inside = [c for c in cs if z.x <= c.x and c.x + c.w <= z.x + z.w and z.y <= c.y and c.y + c.h <= z.y + z.h]
        assert sum(c.w * c.h for c in inside) / (z.w * z.h) >= 0.15, z.ref  # zones share the tallest zone's height
    band = next((z for z in sc.zones if z.ref == "platform"), None)
    plat = {c.ref for c in cs if band and band.x <= c.x and c.y >= band.y and c.y + c.h <= band.y + band.h}
    for ln in sc.lines:
        if plat & set(ln.refs):
            assert band and all(band.x <= x <= band.x + band.w and band.y <= y <= band.y + band.h for x, y in ln.points), ln.refs


def test_empty_zones_collapse_and_widths_follow_the_zones_in_use():
    two = layout(arch(["channels", "core"]))
    five = layout(arch(["channels", "experience", "integration", "core", "data"]))
    assert len(two.zones) == 2 and len(five.zones) == 5
    assert min(c.w / two.scale for c in cards(two)) >= 260 > min(c.w / five.scale for c in cards(five))
    assert not any(t.text == "PLATFORM" for t in two.texts)


def test_text_wraps_without_cutting():
    lines = wrap("one two three four five six seven eight nine ten eleven twelve", "role", 120)
    assert len(lines) > 2 and " ".join(lines) == "one two three four five six seven eight nine ten eleven twelve"
    assert all(measure(x, "role") <= 120 for x in lines)
    assert wrap("short", "role", 120) == ["short"]
    sc = layout(arch(["core"], name="Core banking ledger and posting engine"))
    assert sum(t.style == "name" for t in sc.texts) == 2 and 1 <= sum(t.style == "role" for t in sc.texts) <= 3


def test_barycentre_puts_connected_cards_level():
    a = Architecture.model_validate({"components": [
        {"id": "x1", "name": "X1", "zone": "channels", "kind": "app"}, {"id": "x2", "name": "X2", "zone": "channels", "kind": "app"},
        {"id": "y1", "name": "Y1", "zone": "core", "kind": "app"}, {"id": "y2", "name": "Y2", "zone": "core", "kind": "app"}],
        "flows": [{"src": "x1", "dst": "y2", "mode": "api"}, {"src": "x2", "dst": "y1", "mode": "api"}]})
    by = {b.ref: b for b in cards(layout(a))}
    assert (by["x1"].y < by["x2"].y) == (by["y2"].y < by["y1"].y) or by["x1"].y == by["y2"].y


# --- renderers ---------------------------------------------------------------------------------
def test_svg_parses_and_is_inert():
    a = Architecture.model_validate({"title": "Pay <&> \"go\"", "components": [{"id": "a", "name": "A<script>", "zone": "core", "kind": "app"}]})
    svg = to_svg(layout(a, "Source: x & y"), a.title, crop=False)
    root = ElementTree.fromstring(svg)
    assert root.get("viewBox") == "0 0 1600 900" and root.get("width") is None
    low = svg.lower()
    assert "<script" not in low and "href" not in low and "http" not in svg.replace("http://www.w3.org/2000/svg", "")
    assert "A&lt;script&gt;" in svg


def test_svg_uses_only_house_palette():
    svg = to_svg(layout(SAMPLES["typical"], "Source: Demo"))
    colours = {colour.upper() for colour in re.findall(r"#[0-9a-fA-F]{6}\b", svg)}
    assert colours <= {"#002535", "#0A5A78", "#FBF321", "#FFFFFF", "#000000", "#808285", "#F4F4F4"}


@pytest.mark.parametrize("name", SAMPLES)
def test_png_is_2x_and_text_is_drawn_in_every_card(name):
    sc = layout(SAMPLES[name], "Source: Demo")
    im = Image.open(io.BytesIO(to_png(sc, crop=False))).convert("RGB")
    assert im.size == (3200, 1800)
    for c in cards(sc):
        box = im.crop((int(c.x * 2) + 6, int(c.y * 2) + 6, int((c.x + c.w) * 2) - 6, int((c.y + c.h) * 2) - 6))
        dark = sum(1 for p in box.get_flattened_data() if sum(p) < 450)
        assert dark / (box.width * box.height) > 0.01, c.ref


@pytest.mark.parametrize("name", SAMPLES)
def test_pptx_shapes_one_freeform_per_flow_and_text_reads_back(name):
    a = SAMPLES[name]
    sc = layout(a, "Source: Demo")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    to_pptx_shapes(slide, sc, Inches(.4), Inches(.3), Inches(9), Inches(5))
    shapes = list(slide.shapes)
    assert sum(s.name.startswith("arch:card-") for s in shapes) == len(a.components)
    plat = {c.id for c in a.components if c.zone == "platform"}
    drawn = [f for f in a.flows if (f.src in plat) == (f.dst in plat)]  # a flow to a platform chip is a note, not a line
    assert sum(s.name.startswith("arch:flow-") for s in shapes) == len(drawn) + len({f.mode for f in drawn})  # plus the legend samples
    assert {s.name for s in shapes if s.name.startswith("arch:icon-")} >= {f"arch:icon-{c.kind}" for c in a.components}
    assert sum(s.name.startswith(("arch:zone", "arch:card-")) for s in shapes) >= len(a.components) + len(sc.zones)
    text = " ".join(s.text_frame.text for s in shapes if s.has_text_frame)
    assert all(c.name.split()[0] in text for c in a.components) and a.title.split()[0] in text
    assert "tailEnd" in slide.shapes._spTree.xml


def test_mermaid_snapshot_and_escaping():
    assert to_mermaid(SAMPLES["small"]) == (FIX / "small.mmd").read_text()
    a = Architecture.model_validate({"components": [
        {"id": "a b", "name": 'Say "hi" [x]', "zone": "core", "kind": "app"}, {"id": "c", "name": "C|D<E>", "zone": "data", "kind": "store"}],
        "flows": [{"src": "a b", "dst": "c", "label": "a|b", "mode": "batch"}]})
    m = to_mermaid(a)
    assert "#34;hi#34; #91;x#93;" in m and "C#124;D#60;E#62;" in m and 'n0 -.->|"a#124;b"| n1' in m


def test_mermaid_cannot_break_its_fence_or_its_lines_and_ids_never_collide():
    a = Architecture.model_validate({"components": [
        {"id": "a-b", "name": "Tick ``` fence", "role": "line one\nline two\r\nthree", "zone": "core", "kind": "app"},
        {"id": "a_b", "name": "Second", "zone": "core", "kind": "app"}],
        "flows": [{"src": "a-b", "dst": "a_b", "label": "x\ny", "mode": "api"}, {"src": "a_b", "dst": "a-b", "label": "back", "mode": "api"}]})
    m = to_mermaid(a)
    assert "`" not in m and "#96;#96;#96;" in m
    assert all(line.startswith(("flowchart", "  ")) for line in m.splitlines())  # no stray line from a newline in the text
    assert "line one line two three" in m and "x y" in m
    assert 'n0 -->|"x y"| n1' in m and 'n1 -->|"back"| n0' in m  # "a-b" and "a_b" are two nodes



def test_every_kind_has_an_icon_in_svg_and_png():
    a = arch(["channels", "experience", "integration", "core", "data", "external", "platform"])
    sc = layout(a)
    assert {sh.kind for sh in sc.shapes} == set(KINDS)
    svg = to_svg(sc)
    assert "<ellipse" in svg and "<polygon" in svg and "<polyline" in svg


@pytest.mark.parametrize("name", CASES)
def test_body_is_scaled_to_fill_the_area_and_zones_are_aligned(name):
    sc = layout(CASES[name], FOOTER)
    ax0, ay0, ax1, ay1 = sc.area
    parts = [*sc.zones, *(b for b in sc.boxes if b.style.startswith(("card-", "pill-", "marker-")) and b.y < ay1)]
    x0, y0 = min(b.x for b in parts), min(b.y for b in parts)
    x1, y1 = max(b.x + b.w for b in parts), max(b.y + b.h for b in parts)
    assert ax0 - 1 <= x0 and x1 <= ax1 + 1 and ay0 - 1 <= y0 and y1 <= ay1 + 1
    if sc.scale < 1.6 - 1e-6:  # not held back by the cap: the limiting side is nearly full
        assert max((x1 - x0) / (ax1 - ax0), (y1 - y0) / (ay1 - ay0)) >= 0.75
    cols = [z for z in sc.zones if z.ref != "platform"]
    assert len({(round(z.y, 1), round(z.h, 1)) for z in cols}) <= 1


@pytest.mark.parametrize("name", SAMPLES)
def test_documents_are_cropped_to_the_content_and_pptx_keeps_the_canvas(name):
    sc = layout(SAMPLES[name], "Source: Demo")
    bx0, by0, bx1, by1 = sc.body
    content = sc.area[1] + (by1 - by0) + (H - sc.area[3])  # header band + body + footer band
    im = Image.open(io.BytesIO(to_png(sc)))
    assert im.size[0] <= 2 * W and im.size[1] <= 2 * (content + 64) + 2 and im.size[1] < 2 * H
    root = ElementTree.fromstring(to_svg(sc))
    vw, vh = map(float, root.get("viewBox").split()[2:])
    assert (vw * 2, vh * 2) == im.size  # the preview and the documents show the same picture
    assert sc.height == H and sc.width == W  # the scene itself, which the pptx slide fills, is unchanged
    assert Image.open(io.BytesIO(to_png(sc, crop=False))).size == (3200, 1800)


def test_cropping_keeps_every_card_and_pill_inside_the_picture():
    from app.archdiagram import cropped
    sc = cropped(layout(SAMPLES["typical"], "Source: Demo"))
    for b in [*sc.zones, *sc.boxes]:
        assert 0 <= b.x and b.x + b.w <= sc.width and 0 <= b.y and b.y + b.h <= sc.height, b
    assert len(cards(sc)) == len(SAMPLES["typical"].components)


@pytest.mark.parametrize("name", CASES)
def test_cards_in_a_row_share_a_height_that_fits_their_text(name):
    sc = layout(CASES[name])
    rows = {}
    for c in cards(sc):
        z = next((z for z in sc.zones if z.x <= c.x < z.x + z.w and z.y <= c.y < z.y + z.h and z.ref != "platform"), None)
        if z:
            rows.setdefault((z.ref, round(c.y, 1)), set()).add(round(c.h, 1))
    assert all(len(h) == 1 for h in rows.values())
    if name == "sample-small":
        assert min(c.h / sc.scale for c in cards(sc)) < 78  # one-line roles no longer leave a fixed 78 px


def test_worst_case_layouts_stay_fast_and_the_router_gives_up_on_budget():
    import time
    from app.archdiagram import _Grid, _SCENES
    for name in ("one-zone", "six-columns", "full", "platform-only"):
        _SCENES.clear()
        t = time.perf_counter()
        layout(CASES[name], FOOTER)
        assert time.perf_counter() - t < 3.0, name  # a loose guard against hangs; route(budget=...) proves the cap deterministically
    assert _Grid([], [], (0, 0, 400, 400)).route((1, 1), 0, (40, 40), 0, budget=3) is None  # out of budget: the caller draws an elbow
    assert layout(CASES["one-zone"], FOOTER) is layout(CASES["one-zone"], FOOTER)  # cached, and deterministic below


def test_fonts_are_loaded_once_per_style_and_size():
    from app.archdiagram import _font_px
    to_png(layout(SAMPLES["small"], "Source: Demo"))
    before = _font_px.cache_info()
    to_png(layout(SAMPLES["small"], "Source: Demo"))
    assert _font_px.cache_info().misses == before.misses  # the second render loads no font


def test_a_long_unvalidated_architecture_still_lays_out_and_renders():
    from app.schema import ArchComponent
    long = "An expanded client label " * 6
    a = Architecture.model_construct(title=long, by_reviewer=False, flags=[], callouts=[], flows=[],
                                     components=[ArchComponent.model_construct(id="a", name=long, role=long, zone="core", kind="app", cite="")])
    sc = layout(a, "Source: " + long)
    assert not [t for t in sc.texts if "…" in t.text] and len(cards(sc)) == 1
    assert to_png(sc) and "<svg" in to_svg(sc)
