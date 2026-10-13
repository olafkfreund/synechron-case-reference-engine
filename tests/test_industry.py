import json
import re
import zipfile
from io import BytesIO

import pytest

from app import db, render
from tests.test_auth import env  # noqa: F401
from tests.test_render import doc_text, pdf_text, slides  # noqa: F401
from tests.test_research import as_user, reg  # noqa: F401
from tests.test_review import DOCS, make  # noqa: F401
from tests.test_search import approved, data  # noqa: F401

G = ["g-user", DOCS]
Q1 = "Entra supports passwordless sign in out of the box"
Q2 = "Tenants need a conditional access policy configured first"


def claim(quote=Q1, type="out_of_the_box", publisher="learn.microsoft.com", url="https://learn.microsoft.com/entra"):
    return {"statement": "AI summary that must never appear", "quote": quote, "type": type, "url": url,
            "publisher": publisher, "retrieved_at": "2026-10-01T00:00:00+00:00", "rank": 1}

OLD = ("Supported out of the box", "Available through configuration", "Common industry practice", "Vendor statement")


@pytest.fixture
def rows(env):  # noqa: F811
    db.init()
    made = []

    def add(claims, by="u1", status="done", case_id=None):
        with db.connect() as c:
            made.append(c.execute("insert into research(query, created_by, status, results, case_id) values ('q',%s,%s,%s,%s) returning id",
                                  (by, status, json.dumps({"pages": [], "skipped": [], "claims": claims}), case_id)).fetchone()[0])
        return made[-1]
    yield add
    with db.connect() as c:
        c.execute("delete from research where id = any(%s)", (made,))


def gen(rid=None, ids=(), fmt="md", ack=False, user="u1", **kw):
    form = {"format": fmt, "case_ids": list(ids), **({"research_id": rid} if rid else {}), **({"industry_context_ack": "on"} if ack else {})}
    return as_user(user, G).post("/generate", data=form, **kw)


def audit():
    with db.connect() as c:
        return c.execute("select case_ids, industry_context_ack, format from generations order by id desc limit 1").fetchone()


def test_no_case_needs_research_and_the_acknowledgement(rows):
    rid = rows([claim()])
    assert gen().status_code == 400                       # nothing at all
    assert gen(rid).status_code == 400                    # research, but no acknowledgement
    r = gen(rid, ack=True)
    assert r.status_code == 200 and audit() == ([], True, "md")


def test_with_a_case_the_ack_is_not_needed_and_stored_false(rows, approved):
    rid, cid = rows([claim()]), approved()
    assert gen(rid, [cid]).status_code == 200 and audit()[:2] == ([cid], False)
    assert gen(rid, [cid], ack=True).status_code == 200 and audit()[:2] == ([cid], False)
    assert gen(None, [cid]).status_code == 200 and "Industry context" not in gen(None, [cid]).text  # a case alone: no section
    assert gen(rid, [cid, cid, cid, 1, 2, 3][:0] or [cid]).status_code == 200


def test_research_must_be_yours_and_finished(rows):
    assert gen(rows([claim()], by="u2"), ack=True).status_code == 404
    assert gen(rows([claim()], status="running"), ack=True).status_code == 404
    assert gen(999999999, ack=True).status_code == 404
    assert gen(rows([]), ack=True).status_code == 400     # done, but no usable statements


def test_markdown_section_wording(rows, approved):
    rid = rows([claim(), claim(Q2, "configuration", "docs.vendor.io", "https://docs.vendor.io/x"),
                claim("Analysts say this is common practice for banks", "industry_practice"),
                claim("The vendor describes this as best in class", "vendor_claim")])
    t = gen(rid, [approved()], ack=True).text
    sec = t[t.index("# Industry context (public sources)"):]
    assert render.INDUSTRY_DISCLAIMER.replace("_", "\\_") in sec
    assert not [p for p in OLD if p in sec]
    assert "docs.vendor.io states:" in sec
    assert f"learn.microsoft.com states: “{Q1}” — https://learn.microsoft.com/entra, retrieved 2026-10-01" in sec.replace("\\", "")
    assert "AI summary" not in t
    assert not re.search(r"\b(we|our|ours|delivered)\b", sec.replace(Q1, ""), re.I), sec  # fixed wording only


def test_every_fixed_string_avoids_first_person_and_delivery_wording():
    fixed = [render.INDUSTRY_HEADING, render.INDUSTRY_DISCLAIMER, "states", "retrieved"]
    assert not [s for s in fixed if re.search(r"\b(we|our|ours|us|delivered)\b", s, re.I)]
    assert render.COMPANY in render.INDUSTRY_DISCLAIMER


def test_docx_and_pdf_have_the_quote_and_source(rows, approved):
    rid, cid = rows([claim()]), approved()
    t = doc_text(gen(rid, [cid], "docx").content)
    for want in (render.INDUSTRY_HEADING, render.INDUSTRY_DISCLAIMER, "learn.microsoft.com states:", Q1, "retrieved 2026-10-01", "Faster onboarding"):
        assert want in t, want
    text = " ".join(pdf_text(gen(rid, [cid], "pdf_docx").content)[0].split())  # the PDF wraps lines
    assert render.INDUSTRY_HEADING in text and Q1 in text and "learn.microsoft.com" in text
    assert "{{" not in t and "{%" not in t
    assert t.count("learn.microsoft.com") == 2  # "states:" and the URL: the publisher is not repeated
    assert not [p for p in OLD if p in t]


def test_pptx_industry_slide_keeps_the_quote_next_to_the_phrase(rows, approved):
    rid, cid = rows([claim(), claim(Q2, "configuration", "docs.vendor.io", "https://docs.vendor.io/x")]), approved()
    prs, out = slides(gen(rid, [cid], "pptx").content)
    assert len(out) == 2 and prs.slides[1].slide_layout.name == render.INDUSTRY_LAYOUT
    ind = out[1]
    assert ind["Title"] == [render.INDUSTRY_HEADING] and ind["Summary"] == [render.INDUSTRY_DISCLAIMER]
    assert len(ind["Statements"]) == 2
    for line, phrase, quote in zip(ind["Statements"], ("learn.microsoft.com states:", "docs.vendor.io states:"), (Q1, Q2)):
        assert line.startswith(phrase) and quote[:40] in line and line.count(phrase.split()[0]) == 1  # source once, quote next to it
    assert len(slides(gen(rid, [], "pptx", ack=True).content)[1]) == 1  # industry only
    assert gen(rid, [cid], "pdf_pptx").content.startswith(b"%PDF")


def test_a_claim_naming_a_protected_client_is_replaced_or_dropped_and_counted(rows, approved, reg, monkeypatch):
    name = reg
    evasion = name.replace("o", "ö")
    kept = gen(rows([claim(f"{evasion} runs this in production daily")]), ack=True)  # #166: replaced, not dropped
    assert kept.status_code == 200 and "a retailer runs this in production daily" in kept.text and "omitted" not in kept.text
    real = render.anonymise.apply
    monkeypatch.setattr(render.anonymise, "apply", lambda t, c: t if "ö" in t else real(t, c))  # misses only the accent: blocked() must still guard
    rid = rows([claim(f"{name} runs passwordless sign in out of the box"), claim(f"{evasion} runs this in production daily"),
                claim(Q2, "configuration")])
    r = gen(rid, [approved()], ack=True)
    assert r.status_code == 200 and "Faster onboarding" in r.text                 # the case output is still produced
    t = r.text
    assert name not in t and evasion not in t and "a retailer runs passwordless" in t and Q2 in t
    assert "1 public statement omitted: it named a protected client" in t
    only = gen(rows([claim(f"{evasion} runs this in production daily")]), ack=True)
    assert only.status_code == 400                                                  # nothing usable is left


def test_pptx_industry_slide_shows_the_omitted_note(rows, approved, reg, monkeypatch):
    rid = rows([claim(), claim(f"{reg.replace('o', 'ö')} runs this in production daily")])
    kept = slides(gen(rid, [approved()], "pptx").content)[1][1]["Statements"]  # #166: replaced, not dropped
    assert len(kept) == 2 and "omitted" not in "".join(kept)
    monkeypatch.setattr(render.anonymise, "apply", lambda t, c: t)  # apply() misses: blocked() alone must still guard
    _, out = slides(gen(rid, [approved()], "pptx").content)
    st = out[1]["Statements"]
    assert st[0].startswith("learn.microsoft.com states:") and st[-1] == "1 public statement omitted: it named a protected client"
    many = rows([claim(f"{Q1} {i}") for i in range(render.SLIDE_ITEMS + 1)] + [claim(f"{reg.replace('o', 'ö')} runs this daily")])
    st = slides(gen(many, [approved()], "pptx").content)[1][1]["Statements"]
    assert len(st) == render.SLIDE_ITEMS + 1 and st[-1].startswith("1 public statement omitted")  # the cap never cuts the note


def test_script_in_a_quote_is_escaped_everywhere(rows, approved):
    evil = "<script>alert(1)</script> & more"
    rid, cid = rows([claim(evil)]), approved()
    md = gen(rid, [cid]).text
    assert "<script>" not in md and "\\<script\\>" in md
    d = gen(rid, [cid], "docx").content
    xml = zipfile.ZipFile(BytesIO(d)).read("word/document.xml").decode()
    assert "<script>" not in xml and "&lt;script&gt;" in xml and evil in doc_text(d)
    p = gen(rid, [cid], "pptx").content
    slide_xml = zipfile.ZipFile(BytesIO(p)).read("ppt/slides/slide2.xml").decode()
    assert "<script>" not in slide_xml and "&lt;script&gt;" in slide_xml  # text, never markup


def test_preview_includes_industry_and_keeps_the_ack(rows, approved):
    rid, cid = rows([claim()]), approved()

    def prev(ids=(), ack=False):  # format is posted too, as the research form does
        form = {"format": "docx", "case_ids": list(ids), "research_id": rid, **({"industry_context_ack": "on"} if ack else {})}
        return as_user("u1", G).post("/preview", data=form)
    r = prev([cid])
    assert r.status_code == 200 and Q1 in r.text and "learn.microsoft.com states:" in r.text
    assert "&mdash; https://learn.microsoft.com/entra, retrieved" in r.text  # the publisher is not repeated after the dash
    assert f'name="research_id" value="{rid}"' in r.text
    assert prev().status_code == 400
    esc = as_user("u1", G).post("/preview", data={"case_ids": [cid], "research_id": rows([claim("<b>x</b> in the quote")])})
    assert "&lt;b&gt;x&lt;/b&gt; in the quote" in esc.text and "<b>x</b>" not in esc.text
    r = prev(ack=True)
    assert r.status_code == 200 and 'name="industry_context_ack"' in r.text


def test_research_page_offers_the_download(rows, approved):
    rid = rows([claim()])
    t = as_user("u1", G).get(f"/research/{rid}").text
    assert 'action="/generate"' in t and f'name="research_id" value="{rid}"' in t and 'name="industry_context_ack"' in t
    assert 'name="case_ids"' not in t  # no originating case
    assert 'action="/generate"' not in as_user("u1", G).get(f"/research/{rows([])}").text


def test_fallback_to_case_layout_without_industry_layout(rows, approved, tmp_path, monkeypatch):
    from pptx import Presentation
    prs = Presentation(render.MASTER)
    lay = next(l for l in prs.slide_layouts if l.name == render.INDUSTRY_LAYOUT)
    lay.name = "Something else"  # a marketing master without the optional layout
    alt = tmp_path / "m.pptx"
    prs.save(alt)
    monkeypatch.setattr(render, "MASTER", alt)
    rid, cid = rows([claim()]), approved()
    _, out = slides(gen(rid, [cid], "pptx").content)
    ind = out[1]
    assert ind["Outcomes"][0].startswith("learn.microsoft.com states:") and Q1[:40] in ind["Outcomes"][0]
    assert ind["Challenge"] == [""] and ind["Solution"] == [""]


def test_fallback_layout_shows_the_omitted_note(rows, approved, reg, tmp_path, monkeypatch):
    from pptx import Presentation
    prs = Presentation(render.MASTER)
    next(l for l in prs.slide_layouts if l.name == render.INDUSTRY_LAYOUT).name = "Something else"
    alt = tmp_path / "m.pptx"
    prs.save(alt)
    monkeypatch.setattr(render, "MASTER", alt)
    rid = rows([claim(), claim(f"{reg.replace('o', 'ö')} runs this in production daily")])
    assert "omitted" not in slides(gen(rid, [approved()], "pptx").content)[1][1]["Outcomes"][-1]  # #166: replaced, not dropped
    monkeypatch.setattr(render.anonymise, "apply", lambda t, c: t)  # apply() misses: blocked() alone must still guard
    _, out = slides(gen(rid, [approved()], "pptx").content)
    assert out[1]["Outcomes"][-1] == "1 public statement omitted: it named a protected client"


def test_citation_is_never_rewritten_and_client_site_is_dropped(rows, approved, reg):
    site = reg.lower() + ".com"  # the protected client's own website is the source
    rid, cid = rows([claim(publisher=site, url=f"https://{site}/p"), claim(Q2, "configuration")]), approved()
    md = gen(rid, [cid]).text
    assert site not in md and "a retailer.com" not in md and "1 public statement omitted" in md and Q2 in md


def test_audit_records_the_research_used(rows, approved):
    rid, cid = rows([claim()]), approved()
    gen(rid, [cid], "md")
    with db.connect() as c:
        assert c.execute("select research_id from generations order by id desc limit 1").fetchone()[0] == rid


def test_research_text_never_reaches_cases(rows, approved):
    rid, cid = rows([claim()]), approved()
    gen(rid, [cid], "docx")
    with db.connect() as c:
        data, search_text = c.execute("select data::text, search_text from cases where id=%s", (cid,)).fetchone()
    assert Q1[:30] not in data and Q1[:30] not in search_text
    rid2 = rows([claim()], case_id=cid)
    assert attach(cid, rid2, [0]).status_code == 303
    with db.connect() as c:
        data, search_text = c.execute("select data::text, search_text from cases where id=%s", (cid,)).fetchone()
    assert Q1[:30] not in data and Q1[:30] not in search_text


# ---- public sources attached to a case (#199) ----

RV = ["g-rev", DOCS]


def token(cid):
    with db.connect() as c:
        return c.execute("select md5(public_sources::text) from cases where id=%s", (cid,)).fetchone()[0]


def stored(cid):
    with db.connect() as c:
        return c.execute("select public_sources from cases where id=%s", (cid,)).fetchone()[0]


def attach(cid, rid, n, s=None, user="u1", groups=RV, **extra):
    return as_user(user, groups).post(f"/review/{cid}/sources", follow_redirects=False,
                                      data={"research_id": rid, "n": n, "s": s or token(cid), **extra})


def test_reviewer_attaches_chosen_claims_to_the_case(rows, approved):
    cid = approved()
    rid = rows([claim(), claim(Q2, "configuration", "docs.vendor.io", "https://docs.vendor.io/x")], case_id=cid)
    r = attach(cid, rid, [1])
    assert r.status_code == 303 and r.headers["location"] == f"/review/{cid}"
    (item,) = stored(cid)
    assert item["quote"] == Q2 and item["research_id"] == rid and item["attached_by"] == "u1" and item["attached_at"]
    assert "statement" not in item
    page = as_user("u1", RV).get(f"/review/{cid}").text
    assert "Public sources (not delivery evidence)" in page and Q2 in page and Q1 not in page
    assert 'action="/review/%d/sources"' % cid in as_user("u1", RV).get(f"/research/{rid}").text
    assert f'action="/review/{cid}/sources"' not in as_user("u1", G).get(f"/research/{rid}").text  # a plain user cannot attach


def test_attach_refuses_research_not_yours_not_done_or_from_another_case(rows, approved):
    cid, other = approved(), approved()
    for rid in (rows([claim()], by="u2", case_id=cid), rows([claim()], status="running", case_id=cid),
                rows([claim()], case_id=other), rows([claim()]), 999999999):
        assert attach(cid, rid, [0]).status_code == 404
    assert attach(cid, rows([claim()], case_id=cid), [0], groups=G).status_code == 403
    assert stored(cid) == []


def test_attach_copies_from_the_row_not_the_form(rows, approved):
    cid = approved()
    rid = rows([claim()], case_id=cid)
    attach(cid, rid, [0, 7, -1], quote="forged", url="https://evil.example/x", publisher="forged")
    (item,) = stored(cid)
    assert item["quote"] == Q1 and item["url"] == "https://learn.microsoft.com/entra" and item["publisher"] == "learn.microsoft.com"


def test_attach_skips_a_blocked_claim_and_stores_the_source_text_verbatim(rows, approved, reg):
    cid = approved()
    evasion = f"{reg.replace('o', 'ö')} runs this daily"
    rid = rows([claim(evasion), claim(publisher=reg.lower() + ".com", url=f"https://{reg.lower()}.com/p"),
                claim(Q2, "configuration")], case_id=cid)
    attach(cid, rid, [0, 1, 2])
    got = stored(cid)
    assert [c["quote"] for c in got] == [evasion, Q2] and reg not in json.dumps(got)  # stored as the source says; anonymised at render
    assert "a retailer runs this daily" in gen(None, [cid]).text


def test_attach_dedupes_and_caps_at_20(rows, approved):
    cid = approved()
    rid = rows([claim(f"{Q1} {i}") for i in range(21)], case_id=cid)
    assert attach(cid, rid, [0, 0, 1]).status_code == 303 and attach(cid, rid, [1, 2]).status_code == 303
    assert len(stored(cid)) == 3
    assert attach(cid, rid, list(range(3, 20))).status_code == 303 and len(stored(cid)) == 20
    before = stored(cid)
    assert attach(cid, rid, [20]).status_code == 400 and stored(cid) == before


def test_attach_and_remove_need_the_current_list(rows, approved):
    cid = approved()
    rid = rows([claim(), claim(Q2, "configuration")], case_id=cid)
    stale = token(cid)
    assert attach(cid, rid, [0], s=stale).status_code == 303
    assert attach(cid, rid, [1], s=stale).status_code == 409
    rm = lambda i, s: as_user("u1", RV).post(f"/review/{cid}/sources/remove", data={"i": i, "s": s}, follow_redirects=False)  # noqa: E731
    assert rm(0, stale).status_code == 409 and len(stored(cid)) == 1
    assert rm(5, token(cid)).status_code == 400
    assert rm(0, token(cid)).status_code == 303 and stored(cid) == []


def test_attached_sources_appear_in_every_format(rows, approved):
    cid = approved()
    attach(cid, rows([claim()], case_id=cid), [0])
    md = gen(None, [cid]).text
    assert "# Best practice and public sources" in md and "Industry context" not in md and Q1 in md
    assert render.INDUSTRY_DISCLAIMER.replace("_", "\\_") in md
    t = doc_text(gen(None, [cid], "docx").content)
    assert render.SOURCES_HEADING in t and Q1 in t and "learn.microsoft.com" in t
    _, out = slides(gen(None, [cid], "pptx").content)
    assert out[1]["Title"] == [render.SOURCES_HEADING] and Q1[:40] in out[1]["Statements"][0]
    assert audit()[0] == [cid]
    with db.connect() as c:
        assert c.execute("select research_id from generations order by id desc limit 1").fetchone()[0] is None


def test_attached_and_download_research_are_merged_without_duplicates(rows, approved):
    cid = approved()
    attach(cid, rows([claim()], case_id=cid), [0])
    rid = rows([claim(), claim(Q2, "configuration", "docs.vendor.io", "https://docs.vendor.io/x")])
    md = gen(rid, [cid]).text.replace("\\", "")
    assert md.count(Q1) == 1 and Q2 in md and md.index(Q1) < md.index(Q2) and "# Best practice and public sources" in md
    with db.connect() as c:
        assert c.execute("select research_id from generations order by id desc limit 1").fetchone()[0] == rid


def test_research_only_output_keeps_the_industry_heading(rows, approved):
    cid = approved()
    md = gen(rows([claim()]), [cid]).text
    assert "# Industry context (public sources)" in md and "Best practice" not in md


def test_client_protected_after_attach_is_dropped_at_render(rows, approved, reg):
    cid = approved()
    attach(cid, rows([claim(Q2, "configuration")], case_id=cid), [0])
    with db.connect() as c:  # the client joined the registry after the attach
        c.execute("update cases set public_sources = jsonb_build_array(public_sources->0 || jsonb_build_object('quote', %s::text), "
                  "public_sources->0)", (f"{reg} runs this daily",))
    md = gen(None, [cid]).text
    assert reg not in md and "a retailer runs this daily" in md and Q2 in md


def test_attached_sources_survive_reextraction_and_research_deletion(rows, approved):
    cid = approved()
    rid = rows([claim()], case_id=cid)
    attach(cid, rid, [0])
    with db.connect() as c:
        c.execute("delete from research where id=%s", (rid,))
        c.execute("update cases set data = data || '{\"summary\": \"re-extracted\"}'::jsonb where id=%s", (cid,))
    assert [c["quote"] for c in stored(cid)] == [Q1] and Q1 in gen(None, [cid]).text


def test_attach_does_not_reopen_an_approved_case(rows, approved):
    cid = approved()
    with db.connect() as c:
        before = c.execute("select status, review_due, approved_by, data::text, search_text from cases where id=%s", (cid,)).fetchone()
    attach(cid, rows([claim()], case_id=cid), [0])
    as_user("u1", RV).post(f"/review/{cid}/sources/remove", data={"i": 0, "s": token(cid)})
    attach(cid, rows([claim()], case_id=cid), [0])
    with db.connect() as c:
        assert c.execute("select status, review_due, approved_by, data::text, search_text from cases where id=%s", (cid,)).fetchone() == before


def _protect_attached_site(cid, reg):
    with db.connect() as c:  # the source's site now belongs to a protected client: blocked() drops it at render
        c.execute("update cases set public_sources = jsonb_build_array(public_sources->0 || jsonb_build_object('url', %s::text))",
                  (f"https://{reg.lower()}.com/p",))


def test_heading_stays_industry_when_every_attached_claim_is_dropped_at_render(rows, approved, reg):
    cid = approved()
    attach(cid, rows([claim(Q2, "configuration")], case_id=cid), [0])
    _protect_attached_site(cid, reg)
    md = gen(rows([claim()]), [cid]).text
    assert "# Industry context (public sources)" in md and "Best practice" not in md and Q1 in md and Q2 not in md


def test_review_page_lists_a_hidden_source_for_removal_without_showing_it(rows, approved, reg):
    cid = approved()
    attach(cid, rows([claim(Q2, "configuration")], case_id=cid), [0])
    _protect_attached_site(cid, reg)
    page = as_user("u1", RV).get(f"/review/{cid}").text
    assert "Withheld: names a protected client" in page and Q2 not in page and reg.lower() not in page
    assert 'name="i" value="0"' in page
