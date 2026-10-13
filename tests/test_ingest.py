import io
import re
import uuid

import boto3
import pytest
from docx import Document
from moto import mock_aws

from app import crawl, db, ingest as ing
from tests.test_auth import ADMIN, ORIGIN, SECRET, client

CASE = ing.Triage(kind="case", describes_delivered_work=True)
TRIAGE = {}


@pytest.fixture
def env(monkeypatch):
    for k in ("AWS_ACCESS_KEY_ID", "AWS_SECRET_ACCESS_KEY"):
        monkeypatch.setenv(k, "x")
    monkeypatch.setenv("AWS_DEFAULT_REGION", "us-east-1")
    monkeypatch.setenv("S3_BUCKET", "orig")
    monkeypatch.setattr(ing, "to_markdown", lambda data, name: data.decode())
    monkeypatch.setattr(ing, "complete_json", lambda *a, **k: TRIAGE["v"])
    TRIAGE["v"] = CASE
    db.init()
    with mock_aws():
        s3 = boto3.client("s3")
        s3.create_bucket(Bucket="orig")
        s3.create_bucket(Bucket="src")
        with db.connect() as c:
            sid = c.execute(
                "insert into sources(kind,name,config,acl_groups) values ('s3',%s,%s,%s) returning id",
                (uuid.uuid4().hex, '{"bucket":"src","prefix":"in/"}', ["g1"])).fetchone()[0]
        SID["v"] = sid
        yield s3, sid
        with db.connect() as c:
            c.execute("delete from jobs where (payload->>'document_id')::bigint in "
                      "(select id from documents where source_id=%s)", (sid,))
            c.execute("delete from cases where document_id is null and id in (select merged_into from cases "
                      "where document_id in (select id from documents where source_id=%s))", (sid,))  # merged rows first
            c.execute("delete from sources where id=%s", (sid,))


SID = {}
JOBS_OF = ("select count(*) from jobs where (payload->>'document_id')::bigint in "
           "(select id from documents where source_id=%s)")


def jobs():
    with db.connect() as c:
        return c.execute(JOBS_OF, (SID["v"],)).fetchone()[0]


def test_dedupe(env):
    _, sid = env
    assert ing.ingest(sid, "a", "a", b"one") == "new"
    assert ing.ingest(sid, "a", "a", b"one") == "skipped"
    with db.connect() as c:
        assert c.execute("select count(*) from documents where source_id=%s", (sid,)).fetchone()[0] == 1
    assert jobs() == 1


def test_same_bytes_rename_updates_title(env):
    _, sid = env
    assert ing.ingest(sid, "a", "old.docx", b"one") == "new"
    n = jobs()
    assert ing.ingest(sid, "a", "new.docx", b"one") == "skipped"
    with db.connect() as c:
        assert c.execute("select title from documents where source_id=%s and external_id='a'",
                         (sid,)).fetchone()[0] == "new.docx"
    assert jobs() == n


def test_new_version_updates_in_place_and_reopens_case(env):
    _, sid = env
    ing.ingest(sid, "a", "a", b"one")
    with db.connect() as c:
        did = c.execute("select id from documents where source_id=%s", (sid,)).fetchone()[0]
        c.execute("insert into cases(document_id,status) values (%s,'approved')", (did,))
    with db.connect() as c:
        c.execute("update sources set acl_groups='{g2}' where id=%s", (sid,))
    assert ing.ingest(sid, "a", "a", b"two") == "updated"
    with db.connect() as c:
        rows = c.execute("select text, acl_groups from documents where source_id=%s", (sid,)).fetchall()
        assert rows == [("two", ["g2"])]
        assert c.execute("select status from cases where document_id=%s", (did,)).fetchone()[0] == "extracted"


def test_non_case_routed_away(env):
    _, sid = env
    TRIAGE["v"] = ing.Triage(kind="proposal", describes_delivered_work=False)
    ing.ingest(sid, "a", "a", b"x")
    assert jobs() == 0
    TRIAGE["v"] = ing.Triage(kind="deck", describes_delivered_work=True)
    ing.ingest(sid, "b", "b", b"y")
    assert jobs() == 1


@pytest.mark.parametrize("kind,deliv,executed,flag,want", [
    ("case", False, False, False, "delivered"),
    ("proposal", True, False, False, "delivered"),
    ("deck", True, False, False, "delivered"),
    ("proposal", False, True, False, None),
    ("deck", False, False, False, None),
    ("other", True, True, True, None),
    ("contract", False, True, False, "engagement"),
    ("contract", True, False, False, None),
    ("contract", False, False, True, "engagement"),
])
def test_basis_for_routing_table(kind, deliv, executed, flag, want):
    t = ing.Triage(kind=kind, describes_delivered_work=deliv, executed=executed)
    assert ing.basis_for(t, {"executed_contracts": True} if flag else {}) == want


def test_contract_payload_and_source_flag(env):
    _, sid = env
    TRIAGE["v"] = ing.Triage(kind="contract", describes_delivered_work=False, executed=False)
    ing.ingest(sid, "a", "a", b"x")
    assert jobs() == 0
    with db.connect() as c:
        c.execute("update sources set config=config || '{\"executed_contracts\": true}' where id=%s", (sid,))
    ing.ingest(sid, "b", "b", b"y")
    TRIAGE["v"] = ing.Triage(kind="contract", describes_delivered_work=False, executed=True)
    ing.ingest(sid, "c", "c", b"z")
    with db.connect() as c:
        got = c.execute("select payload->>'basis', payload->>'basis_reason' from jobs where (payload->>'document_id')::bigint in "
                        "(select id from documents where source_id=%s) order by id", (sid,)).fetchall()
    assert got == [("engagement", "source marked executed"), ("engagement", "executed contract")]


def _toggle_during_triage(monkeypatch, sid, sql):
    """As an admin's save does while a document converts: the flag changes after the early read (#133)."""
    def triage(text, data_class):
        with db.connect() as c:
            c.execute(sql, (sid,))
        return ing.Triage(kind="contract", describes_delivered_work=False, executed=False)
    monkeypatch.setattr(ing, "triage_text", triage)


def _flag_jobs(sid):
    with db.connect() as c:
        return c.execute("select payload->>'basis', payload->>'basis_reason' from jobs where (payload->>'document_id')::bigint in "
                         "(select id from documents where source_id=%s) order by id", (sid,)).fetchall()


def test_executed_ticked_during_triage_queues_engagement(env, monkeypatch):
    _, sid = env
    _toggle_during_triage(monkeypatch, sid, "update sources set config=config || '{\"executed_contracts\": true}' where id=%s")
    ing.ingest(sid, "a", "a", b"x")
    assert _flag_jobs(sid) == [("engagement", "source marked executed")]


def test_executed_unticked_during_triage_queues_nothing(env, monkeypatch):
    _, sid = env
    with db.connect() as c:
        c.execute("update sources set config=config || '{\"executed_contracts\": true}' where id=%s", (sid,))
    _toggle_during_triage(monkeypatch, sid, "update sources set config=config - 'executed_contracts' where id=%s")
    ing.ingest(sid, "a", "a", b"x")
    assert jobs() == 0


def test_ingest_waits_for_an_uncommitted_tick(env, monkeypatch):
    """The flag read takes the share lock: an admin's save still in flight is waited for, not read stale (#133)."""
    import threading
    _, sid = env
    monkeypatch.setattr(ing, "triage_text", lambda text, data_class: ing.Triage(
        kind="contract", describes_delivered_work=False, executed=False))
    save = db.connect()  # stands in for the admin save: row locked, flag set, not yet committed
    try:
        save.execute("select 1 from sources where id=%s for update", (sid,))
        save.execute("update sources set config=config || '{\"executed_contracts\": true}' where id=%s", (sid,))
        t = threading.Thread(target=ing.ingest, args=(sid, "a", "a", b"x"))
        t.start()
        _wait_for_lock_wait(save)
        save.commit()
    finally:
        save.close()
    t.join(10)
    assert not t.is_alive() and _flag_jobs(sid) == [("engagement", "source marked executed")]


def test_data_class_raised_during_conversion_triages_under_new_class(env, monkeypatch):
    _, sid = env
    with db.connect() as c:
        c.execute("update sources set data_class='public' where id=%s", (sid,))

    def convert(data, name):  # the admin raises the class while Docling runs
        with db.connect() as c:
            c.execute("update sources set data_class='confidential' where id=%s", (sid,))
        return data.decode()
    seen = []
    monkeypatch.setattr(ing, "to_markdown", convert)
    monkeypatch.setattr(ing, "triage_text", lambda text, data_class: seen.append(data_class) or CASE)
    ing.ingest(sid, "a", "a", b"x")
    assert seen == ["confidential"]


def test_triage_input_includes_the_tail(monkeypatch):
    seen = []
    monkeypatch.setattr(ing, "complete_json", lambda m, s, text, *a, **k: seen.append(text) or CASE)
    ing.triage_text("H" * 9000 + "M" * 9000 + "SIGNED", "confidential")
    assert seen[0].endswith("SIGNED") and len(seen[0]) < 12100 and "M" * 3000 not in seen[0][8000:8010]
    ing.triage_text("short", "confidential")
    assert seen[1] == "short"


def test_crawl_cursor_and_deletion(env):
    s3, sid = env
    s3.put_object(Bucket="src", Key="in/a.docx", Body=b"aaa")
    s3.put_object(Bucket="src", Key="in/b.docx", Body=b"bbb")
    assert crawl.crawl_s3(sid) == {"new": 2, "updated": 0, "skipped": 0, "deleted": 0, "failed": 0,
                                   "skipped_type": 0, "skipped_too_large": 0}
    again = crawl.crawl_s3(sid)
    assert again["new"] == 0 and again["updated"] == 0 and again["skipped"] == 2  # inside cursor slack
    s3.delete_object(Bucket="src", Key="in/b.docx")
    assert crawl.crawl_s3(sid)["deleted"] == 1
    with db.connect() as c:
        assert c.execute("select deleted_at is not null from documents where external_id='in/b.docx'").fetchone()[0]
        assert c.execute("select acl_groups from documents where external_id='in/a.docx'").fetchone()[0] == ["g1"]


def test_upload_title_is_the_file_name(env, monkeypatch):
    for k, v in dict(SESSION_SECRET=SECRET, SESSION_HTTPS_ONLY="false", APP_ORIGIN=ORIGIN,
                     ROLE_ADMIN_GROUPS=ADMIN).items():
        monkeypatch.setenv(k, v)
    db.init()
    with db.connect() as c:
        usid = c.execute("insert into sources(kind,name,config) values ('upload',%s,%s) returning id",
                         (uuid.uuid4().hex, '{"bucket":"src","prefix":"up/"}')).fetchone()[0]
    try:
        for body in (b"PK\x03\x04one", b"PK\x03\x04two"):
            r = client([ADMIN]).post("/admin/upload", data={"source_id": usid}, files={"file": ("My Case.docx", body)})
            assert r.status_code == 202
        assert crawl.crawl_s3(usid)["new"] == 2
        with db.connect() as c:
            rows = c.execute("select title, external_id from documents where source_id=%s", (usid,)).fetchall()
        assert len(rows) == 2 and {t for t, _ in rows} == {"My_Case.docx"}
        assert len({e for _, e in rows}) == 2
        assert all(re.fullmatch(r"up/[0-9a-f]{32}/My_Case\.docx", e) for _, e in rows)
    finally:
        with db.connect() as c:
            c.execute("delete from jobs where payload->>'source_id' = %s", (str(usid),))
            c.execute("delete from jobs where (payload->>'document_id')::bigint in "
                      "(select id from documents where source_id=%s)", (usid,))
            c.execute("delete from sources where id=%s", (usid,))


def test_crawl_skips_type_and_size_without_downloading(env, monkeypatch):
    s3, sid = env
    monkeypatch.setenv("S3_MAX_BYTES", "10")
    for key, body in [("a.docx", b"aaa"), ("b.PDF", b"bbb"), ("c.mp4", b"c"), ("d.txt", b"d"), ("e.pdf", b"e" * 20)]:
        s3.put_object(Bucket="src", Key=f"in/{key}", Body=body)
    fetched, real = [], crawl.boto3.client

    def client(*a, **k):  # the real moto client, with downloads recorded
        c = real(*a, **k)
        get = c.get_object
        c.get_object = lambda **kw: (fetched.append(kw["Key"]), get(**kw))[1]
        return c
    monkeypatch.setattr(crawl.boto3, "client", client)
    counts = crawl.crawl_s3(sid)
    assert (counts["new"], counts["skipped_type"], counts["skipped_too_large"]) == (2, 2, 1)
    assert sorted(fetched) == ["in/a.docx", "in/b.PDF"]
    with db.connect() as c:
        assert sorted(r[0] for r in c.execute("select external_id from documents where source_id=%s", (sid,))) == [
            "in/a.docx", "in/b.PDF"]
    fetched.clear()
    again = crawl.crawl_s3(sid)
    assert (again["skipped_type"], again["skipped_too_large"], again["deleted"]) == (2, 1, 0)
    assert not set(fetched) & {"in/c.mp4", "in/d.txt", "in/e.pdf"}


def test_crawl_oversized_existing_doc_is_withdrawn(env, monkeypatch):
    s3, sid = env
    monkeypatch.setenv("S3_MAX_BYTES", "10")
    with db.connect() as c:
        c.execute("""update sources set config = config || '{"include_ext": ["txt", "docx"]}' where id=%s""", (sid,))
    s3.put_object(Bucket="src", Key="in/notes.txt", Body=b"notes")
    s3.put_object(Bucket="src", Key="in/a.docx", Body=b"small")
    assert crawl.crawl_s3(sid)["new"] == 2
    s3.put_object(Bucket="src", Key="in/a.docx", Body=b"now far too large")
    counts = crawl.crawl_s3(sid)
    assert (counts["skipped_too_large"], counts["deleted"]) == (1, 1)
    with db.connect() as c:
        assert c.execute("select deleted_at is not null from documents where external_id='in/a.docx'").fetchone()[0]


def test_crawl_s3_type_narrowed_withdraws(env):
    s3, sid = env
    s3.put_object(Bucket="src", Key="in/a.docx", Body=b"aaa")
    s3.put_object(Bucket="src", Key="in/b.pdf", Body=b"bbb")
    assert crawl.crawl_s3(sid)["new"] == 2
    with db.connect() as c:
        c.execute("""update sources set config = config || '{"include_ext": ["docx"]}' where id=%s""", (sid,))
    counts = crawl.crawl_s3(sid)
    assert (counts["deleted"], counts["skipped_type"]) == (1, 1)
    with db.connect() as c:
        rows = dict(c.execute("select external_id, deleted_at is not null from documents where source_id=%s", (sid,)))
    assert rows == {"in/a.docx": False, "in/b.pdf": True}


FUTURE = "2999-01-01T00:00:00+00:00"  # every stored key is now older than the cursor slack


def withdrawn_then_changed(s3, sid, monkeypatch):
    """in/a.txt ingested, then changed while over the cap (so withdrawn), then the cap raised and the cursor moved on."""
    monkeypatch.setenv("S3_MAX_BYTES", "10")
    with db.connect() as c:
        c.execute("""update sources set config = config || '{"include_ext": ["txt"]}' where id=%s""", (sid,))
    s3.put_object(Bucket="src", Key="in/keep.txt", Body=b"keep")  # a non-empty listing, so the sweep runs
    s3.put_object(Bucket="src", Key="in/a.txt", Body=b"old text")
    assert crawl.crawl_s3(sid)["new"] == 2
    s3.put_object(Bucket="src", Key="in/a.txt", Body=b"brand new longer text")
    assert crawl.crawl_s3(sid)["deleted"] == 1
    set_cursor(sid, FUTURE)
    monkeypatch.setenv("S3_MAX_BYTES", "1000")


def doc_state(sid, key):
    with db.connect() as c:
        return c.execute("select deleted_at is null, text from documents where source_id=%s and external_id=%s",
                         (sid, key)).fetchone()


def test_raised_cap_reconverts_a_file_changed_while_too_large(env, monkeypatch):
    s3, sid = env
    withdrawn_then_changed(s3, sid, monkeypatch)
    assert crawl.crawl_s3(sid)["updated"] == 1
    live, text = doc_state(sid, "in/a.txt")
    assert live and "brand new" in text


def test_widened_include_ext_reconverts_a_file_changed_while_excluded(env):
    s3, sid = env
    with db.connect() as c:
        c.execute("""update sources set config = config || '{"include_ext": ["txt", "docx"]}' where id=%s""", (sid,))
    s3.put_object(Bucket="src", Key="in/b.txt", Body=b"old")
    s3.put_object(Bucket="src", Key="in/keep.docx", Body=b"keep")  # a non-empty listing, so the sweep runs
    assert crawl.crawl_s3(sid)["new"] == 2
    with db.connect() as c:
        c.execute("""update sources set config = config || '{"include_ext": ["docx"]}' where id=%s""", (sid,))
    s3.put_object(Bucket="src", Key="in/b.txt", Body=b"changed")
    assert crawl.crawl_s3(sid)["deleted"] == 1
    set_cursor(sid, FUTURE)
    with db.connect() as c:
        c.execute("""update sources set config = config || '{"include_ext": ["txt", "docx"]}' where id=%s""", (sid,))
    assert crawl.crawl_s3(sid)["updated"] == 1
    live, text = doc_state(sid, "in/b.txt")
    assert live and text == "changed"


def test_withdrawn_key_that_fails_on_return_stays_withdrawn(env, monkeypatch):
    s3, sid = env
    withdrawn_then_changed(s3, sid, monkeypatch)
    def boom(*a, **k):
        raise ValueError("down")
    monkeypatch.setattr(crawl, "ingest", boom)
    assert crawl.crawl_s3(sid)["failed"] == 1
    live, _ = doc_state(sid, "in/a.txt")
    assert not live
    with db.connect() as c:
        assert c.execute("select last_counts->'retry_ids' from sources where id=%s", (sid,)).fetchone()[0] == ["in/a.txt"]


def test_unchanged_live_key_still_skips_the_download(env, monkeypatch):
    s3, sid = env
    s3.put_object(Bucket="src", Key="in/a.docx", Body=b"aaa")
    assert crawl.crawl_s3(sid)["new"] == 1
    set_cursor(sid, FUTURE)
    fetched, real = [], crawl.boto3.client
    def client(*a, **k):
        c = real(*a, **k)
        get = c.get_object
        c.get_object = lambda **kw: (fetched.append(kw["Key"]), get(**kw))[1]
        return c
    monkeypatch.setattr(crawl.boto3, "client", client)
    crawl.crawl_s3(sid)
    assert fetched == []


def test_empty_listing_with_only_withdrawn_documents_is_not_a_mass_delete(env):
    s3, sid = env
    s3.put_object(Bucket="src", Key="in/a.docx", Body=b"aaa")
    assert crawl.crawl_s3(sid)["new"] == 1
    with db.connect() as c:
        c.execute("update documents set deleted_at=now() where source_id=%s", (sid,))
    s3.delete_object(Bucket="src", Key="in/a.docx")
    assert crawl.crawl_s3(sid)["empty_listing"] is True


def test_real_docling_conversion():
    buf = io.BytesIO()
    d = Document()
    d.add_heading("Bank onboarding", 1)
    d.add_paragraph("We cut onboarding from 12 days to 3 days.")
    d.save(buf)
    md = ing.to_markdown(buf.getvalue(), "case.docx")
    assert "Bank onboarding" in md and "12 days to 3 days" in md


def set_cursor(sid, iso):
    with db.connect() as c:
        c.execute("update sources set cursor=%s where id=%s", (iso, sid))


def test_cursor_skips_known_old_keys_but_fetches_unknown(env):
    s3, sid = env
    s3.put_object(Bucket="src", Key="in/a.docx", Body=b"aaa")
    crawl.crawl_s3(sid)
    set_cursor(sid, "2999-01-01T00:00:00+00:00")  # everything is now "old"
    s3.put_object(Bucket="src", Key="in/late.docx", Body=b"late")  # multipart-style: dated behind cursor
    r = crawl.crawl_s3(sid)
    assert r["skipped"] == 0 and r["new"] == 1  # a.docx not re-downloaded, unknown late.docx fetched


def test_bad_document_does_not_stop_crawl(env, monkeypatch):
    s3, sid = env
    real = ing.to_markdown
    def boom(data, name):
        if data == b"bad":
            raise ValueError("corrupt SECRET-CLIENT file")
        return real(data, name)
    monkeypatch.setattr(ing, "to_markdown", boom)
    s3.put_object(Bucket="src", Key="in/a.docx", Body=b"bad")
    s3.put_object(Bucket="src", Key="in/b.docx", Body=b"good")
    r = crawl.crawl_s3(sid)
    assert r["failed"] == 1 and r["new"] == 1
    with db.connect() as c:
        cursor, counts = c.execute("select cursor, last_counts from sources where id=%s", (sid,)).fetchone()
    assert cursor and counts["failed_keys"] == [{"key": "in/a.docx", "error": "ValueError"}]
    assert "SECRET" not in str(counts)


def test_failed_new_version_is_retried_past_the_cursor(env, monkeypatch):
    s3, sid = env
    real = ing.to_markdown
    s3.put_object(Bucket="src", Key="in/a.docx", Body=b"v1")
    crawl.crawl_s3(sid)
    def boom(data, name):
        if data == b"v2":
            raise ValueError("corrupt")
        return real(data, name)
    monkeypatch.setattr(ing, "to_markdown", boom)
    s3.put_object(Bucket="src", Key="in/a.docx", Body=b"v2")
    assert crawl.crawl_s3(sid)["failed"] == 1
    with db.connect() as c:
        assert c.execute("select last_counts->'retry_ids' from sources where id=%s", (sid,)).fetchone()[0] == ["in/a.docx"]
    set_cursor(sid, "2999-01-01T00:00:00+00:00")  # later uploads moved the cursor far past it
    monkeypatch.setattr(ing, "to_markdown", real)
    assert crawl.crawl_s3(sid)["updated"] == 1  # main: skipped on the cursor, never downloaded
    with db.connect() as c:
        assert c.execute("select text from documents where source_id=%s and external_id='in/a.docx'",
                         (sid,)).fetchone()[0] == "v2"
    r = crawl.crawl_s3(sid)  # succeeded, so off the retry list: skipped on the cursor again
    assert r["updated"] == 0 and r["skipped"] == 0


def test_acl_change_and_reappearing_key_apply_without_reconversion(env):
    s3, sid = env
    s3.put_object(Bucket="src", Key="in/a.docx", Body=b"aaa")
    s3.put_object(Bucket="src", Key="in/b.docx", Body=b"bbb")
    crawl.crawl_s3(sid)
    s3.delete_object(Bucket="src", Key="in/b.docx")
    crawl.crawl_s3(sid)
    s3.put_object(Bucket="src", Key="in/b.docx", Body=b"bbb")
    with db.connect() as c:
        c.execute("update sources set acl_groups='{g9}' where id=%s", (sid,))
    set_cursor(sid, "2999-01-01T00:00:00+00:00")
    r = crawl.crawl_s3(sid)  # the withdrawn key is downloaded again (#189); same bytes, so the checksum skips it
    assert (r["skipped"], r["new"], r["updated"]) == (1, 0, 0)
    with db.connect() as c:
        rows = c.execute("select acl_groups, deleted_at is null from documents where source_id=%s", (sid,)).fetchall()
    assert rows == [(["g9"], True), (["g9"], True)]


def test_empty_listing_is_not_a_mass_delete(env):
    s3, sid = env
    s3.put_object(Bucket="src", Key="in/a.docx", Body=b"aaa")
    crawl.crawl_s3(sid)
    s3.delete_object(Bucket="src", Key="in/a.docx")
    r = crawl.crawl_s3(sid)
    assert r["deleted"] == 0 and r["empty_listing"] is True
    with db.connect() as c:
        assert c.execute("select deleted_at is null from documents where source_id=%s", (sid,)).fetchone()[0]


def test_concurrent_crawl_returns_running(env):
    _, sid = env
    with db.connect() as other:
        other.execute("select pg_advisory_lock(2, %s)", (sid,))
        assert crawl.crawl_s3(sid) == {"status": "running"}


def test_changed_document_no_longer_a_case_retires_it(env):
    _, sid = env
    ing.ingest(sid, "a", "a", b"one")
    with db.connect() as c:
        did = c.execute("select id from documents where source_id=%s", (sid,)).fetchone()[0]
        c.execute("insert into cases(document_id,status) values (%s,'approved')", (did,))
    TRIAGE["v"] = ing.Triage(kind="other", describes_delivered_work=False)
    before = jobs()
    ing.ingest(sid, "a", "a", b"two")
    with db.connect() as c:
        assert c.execute("select status from cases where document_id=%s", (did,)).fetchone()[0] == "rejected"
    assert jobs() == before



def test_changed_document_clears_the_untick_mark(env):  # #160: a reject by ingest is not an untick reject
    _, sid = env
    ing.ingest(sid, "a", "a", b"one")
    with db.connect() as c:
        did = c.execute("select id from documents where source_id=%s", (sid,)).fetchone()[0]
        c.execute("insert into cases(document_id,status,retired_by) values (%s,'rejected','untick')", (did,))
    TRIAGE["v"] = ing.Triage(kind="other", describes_delivered_work=False)
    ing.ingest(sid, "a", "a", b"two")
    with db.connect() as c:
        assert c.execute("select status, retired_by from cases where document_id=%s", (did,)).fetchone() == ("rejected", None)

def merged_pair(sid):
    """Two ingested documents "a" and "b", each with an engagement case, merged into one approved case."""
    ing.ingest(sid, "a", "a", b"one")
    ing.ingest(sid, "b", "b", b"one")
    with db.connect() as c:
        dids = [r[0] for r in c.execute("select id from documents where source_id=%s order by external_id", (sid,))]
        new = c.execute("insert into cases(document_id, member_count, status) values (null, 2, 'approved') "
                        "returning id").fetchone()[0]
        for d in dids:
            c.execute("insert into cases(document_id, status, merged_into) values (%s, 'approved', %s)", (d, new))
    return new, dids


def status(cid):
    with db.connect() as c:
        return c.execute("select status from cases where id=%s", (cid,)).fetchone()[0]


def test_changed_member_reopens_its_approved_merged_case(env):
    _, sid = env
    new, _ = merged_pair(sid)
    assert ing.ingest(sid, "a", "a", b"one") == "skipped" and status(new) == "approved"  # nothing changed
    ing.ingest(sid, "a", "a", b"two")
    assert status(new) == "extracted"


def test_member_that_is_no_longer_an_engagement_is_rejected_and_blocks_approval(env):
    _, sid = env
    new, dids = merged_pair(sid)
    TRIAGE["v"] = ing.Triage(kind="other", describes_delivered_work=False)
    ing.ingest(sid, "a", "a", b"two")
    with db.connect() as c:
        assert c.execute("select status, merged_into from cases where document_id=%s", (dids[0],)).fetchone() == ("rejected", new)
    assert status(new) == "extracted"  # approval is then refused by test_review's rejected-member test


def test_unrelated_document_leaves_merged_cases_alone(env):
    _, sid = env
    new, _ = merged_pair(sid)
    ing.ingest(sid, "c", "c", b"one")
    ing.ingest(sid, "c", "c", b"two")
    assert status(new) == "approved"


def test_executed_flag_must_be_true_not_truthy():
    t = ing.Triage(kind="contract", describes_delivered_work=False, executed=False)
    assert ing.basis_for(t, {"executed_contracts": "false"}) is None
    assert ing.basis_for(t, {"executed_contracts": True}) == "engagement"


def test_member_change_during_merged_approve_reopens_it(env):
    """A member change that lands while M is being approved waits for it, then reopens M (#144)."""
    import threading
    _, sid = env
    new, _ = merged_pair(sid)
    with db.connect() as c:
        c.execute("update cases set status='extracted' where id=%s", (new,))
    first = db.connect()  # stands in for approve: load() holds M's row lock
    first.execute("select 1 from cases where id=%s for update", (new,))
    t = threading.Thread(target=ing.ingest, args=(sid, "a", "a", b"two"))
    t.start()
    try:
        _wait_for_lock_wait(first)
        first.execute("update cases set status='approved' where id=%s", (new,))
        first.commit()
    finally:
        first.close()  # releases the lock (rolls back on failure), or teardown blocks on it
    t.join(10)
    assert not t.is_alive() and status(new) == "extracted"


def _wait_for_lock_wait(conn, timeout=10):
    """Until another backend is waiting on a row lock (the ingest blocked by the open save)."""
    import time
    end = time.time() + timeout
    with db.connect(autocommit=True) as c:
        while time.time() < end:
            if c.execute("select count(*) from pg_stat_activity where wait_event_type = 'Lock' "
                         "and pid <> pg_backend_pid()").fetchone()[0]:
                return
            time.sleep(0.05)
    raise AssertionError("ingest never waited for the save's lock")


def test_ingest_waits_for_a_saving_admin_and_writes_the_new_groups(env):
    import threading
    _, sid = env
    first = db.connect()
    first.execute("select 1 from sources where id=%s for update", (sid,))
    first.execute("update sources set acl_groups='{new}' where id=%s", (sid,))
    t = threading.Thread(target=ing.ingest, args=(sid, "a", "a", b"one"))
    t.start()
    try:
        _wait_for_lock_wait(first)
        first.commit()
    finally:
        first.close()  # releases the lock (rolls back on failure), or teardown blocks on it
    t.join(10)
    with db.connect() as c:
        assert c.execute("select acl_groups from documents where source_id=%s", (sid,)).fetchone()[0] == ["new"]


def test_unchanged_reingest_waits_for_a_saving_admin(env):
    import threading
    _, sid = env
    ing.ingest(sid, "a", "a", b"one")
    first = db.connect()
    first.execute("select 1 from sources where id=%s for update", (sid,))
    first.execute("update sources set acl_groups='{new}' where id=%s", (sid,))
    t = threading.Thread(target=ing.ingest, args=(sid, "a", "a", b"one"))  # same bytes: the skipped branch
    t.start()
    try:
        _wait_for_lock_wait(first)
        first.commit()
    finally:
        first.close()  # releases the lock (rolls back on failure), or teardown blocks on it
    t.join(10)
    with db.connect() as c:
        assert c.execute("select acl_groups from documents where source_id=%s", (sid,)).fetchone()[0] == ["new"]


def test_migration_repair_waits_for_a_saving_admin(env):
    """End to end: a migration during an open save ends with the save's groups. The schema's earlier
    `alter table sources/documents` locks already serialise it; the repair's `for share` is the guard
    if those statements are ever removed (a standalone repair without it can write back old groups)."""
    import threading
    _, sid = env
    ing.ingest(sid, "a", "a", b"one")
    with db.connect() as c:
        c.execute("update documents set acl_groups='{stale,g1}' where source_id=%s", (sid,))  # pre-#58 damage
    first = db.connect()
    first.execute("select 1 from sources where id=%s for update", (sid,))
    first.execute("update sources set acl_groups='{new}' where id=%s", (sid,))
    first.execute("update documents set acl_groups='{new}' where source_id=%s", (sid,))
    t = threading.Thread(target=db.init)
    t.start()
    try:
        _wait_for_lock_wait(first)
        first.commit()
    finally:
        first.close()
    t.join(30)
    with db.connect() as c:
        assert c.execute("select acl_groups from documents where source_id=%s", (sid,)).fetchone()[0] == ["new"]


def test_source_deleted_before_insert_raises(env, monkeypatch):
    _, sid = env

    def triage_then_delete(text, data_class):
        with db.connect() as c:
            c.execute("delete from sources where id=%s", (sid,))
        return CASE
    monkeypatch.setattr(ing, "triage_text", triage_then_delete)
    with pytest.raises(LookupError):
        ing.ingest(sid, "a", "a", b"one")

def _change_groups_after_first_ingest(monkeypatch, sid):
    """As an admin's save does mid-crawl: new groups on the source and on its documents already written."""
    real, done = crawl.ingest, []

    def wrapped(*a, **k):
        out = real(*a, **k)
        if not done:
            done.append(1)
            with db.connect() as c:
                c.execute("update sources set acl_groups='{g-new}' where id=%s", (sid,))
                c.execute("update documents set acl_groups='{g-new}' where source_id=%s", (sid,))
        return out
    monkeypatch.setattr(crawl, "ingest", wrapped)


def test_s3_crawl_does_not_write_back_groups_changed_mid_crawl(env, monkeypatch):
    s3, sid = env
    for k in "abc":
        s3.put_object(Bucket="src", Key=f"in/{k}.docx", Body=k.encode())
    _change_groups_after_first_ingest(monkeypatch, sid)
    crawl.crawl_s3(sid)
    with db.connect() as c:
        rows = c.execute("select acl_groups from documents where source_id=%s", (sid,)).fetchall()
    assert rows == [(["g-new"],)] * 3


def stamp_old(sid):
    with db.connect() as c:
        c.execute("update documents set updated_at = now() - interval '2 days' where source_id=%s", (sid,))


def updated_age(sid):
    with db.connect() as c:
        return c.execute("select now() - updated_at from documents where source_id=%s", (sid,)).fetchone()[0].days


def test_changed_content_moves_a_document_up_but_unchanged_does_not(env):  # #215
    _, sid = env
    ing.ingest(sid, "a", "a", b"one", source_version="v1")
    stamp_old(sid)
    assert ing.ingest(sid, "a", "a", b"one", source_version="v1") == "skipped"  # same bytes
    assert updated_age(sid) == 2
    assert ing.skip_unchanged(sid, "a", "v1", "a")  # crawl skips the download
    assert updated_age(sid) == 2
    assert ing.ingest(sid, "a", "a", b"two", source_version="v2") == "updated"
    assert updated_age(sid) == 0


def images_of(sid):
    with db.connect() as c:
        return c.execute("select i.width, i.height, i.hint, i.png from document_images i join documents d on d.id = i.document_id "
                         "where d.source_id=%s order by i.id", (sid,)).fetchall()


def test_image_candidates_are_stored_with_the_registry_and_replaced_on_a_new_version(env, monkeypatch):
    _, sid = env
    seen = []

    def fake(data, name, clients=()):
        seen.append((name, [c["name"] for c in clients]))
        return [(b"png-" + data, 500, 400, f"hint {data.decode()}")]
    monkeypatch.setattr(ing, "document_images_from", fake)
    with db.connect() as c:
        cid = c.execute("insert into clients(name, aliases, anonymised_label) values (%s, '{}', 'a bank') returning id",
                        (f"Zorpale {uuid.uuid4().hex[:6]}",)).fetchone()[0]
    try:
        ing.ingest(sid, "a", "a.docx", b"one")
        assert images_of(sid) == [(500, 400, "hint one", b"png-one")]
        assert seen[0][0] == "a.docx" and any(n.startswith("Zorpale") for n in seen[0][1])  # the registry is passed in
        ing.ingest(sid, "a", "a.docx", b"two")
        assert images_of(sid) == [(500, 400, "hint two", b"png-two")]  # replaced, not added
        ing.ingest(sid, "a", "a.docx", b"two")  # same bytes: left as they are
        assert len(seen) == 2 and len(images_of(sid)) == 1
    finally:
        with db.connect() as c:
            c.execute("delete from clients where id=%s", (cid,))


def test_a_document_that_gets_no_case_gets_no_images(env, monkeypatch):
    _, sid = env
    monkeypatch.setattr(ing, "document_images_from", lambda *a, **k: [(b"p", 500, 400, "")])
    ing.ingest(sid, "a", "a.docx", b"one")
    assert len(images_of(sid)) == 1
    TRIAGE["v"] = ing.Triage(kind="proposal", describes_delivered_work=False)
    ing.ingest(sid, "a", "a.docx", b"two")  # now routed away: the old candidates go
    assert images_of(sid) == []


def test_an_image_failure_never_fails_the_ingest(env, monkeypatch, capsys):
    _, sid = env

    def boom(*a, **k):
        raise RuntimeError("echo Zorbulon")
    monkeypatch.setattr(ing, "document_images_from", boom)
    assert ing.ingest(sid, "a", "a.docx", b"one") == "new" and jobs() == 1
    out = capsys.readouterr().out
    assert "RuntimeError" in out and "Zorbulon" not in out and images_of(sid) == []


def test_a_document_ingested_before_images_gets_them_once_on_a_same_bytes_crawl(env, monkeypatch):
    _, sid = env
    calls = []
    monkeypatch.setattr(ing, "document_images_from", lambda *a, **k: calls.append(1) or [(b"p", 500, 400, "h")])
    ing.ingest(sid, "a", "a.docx", b"one")
    with db.connect() as c:  # as if it was ingested before phase C: a case, no candidates
        did = c.execute("select id from documents where source_id=%s", (sid,)).fetchone()[0]
        c.execute("insert into cases(document_id,status) values (%s,'extracted')", (did,))
        c.execute("delete from document_images where document_id=%s", (did,))
    calls.clear()
    assert ing.ingest(sid, "a", "a.docx", b"one") == "skipped"
    assert len(calls) == 1 and len(images_of(sid)) == 1
    assert ing.ingest(sid, "a", "a.docx", b"one") == "skipped"
    assert len(calls) == 1  # the rows exist now: nothing is decoded again


def test_same_bytes_without_a_case_stores_no_images(env, monkeypatch):
    _, sid = env
    TRIAGE["v"] = ing.Triage(kind="proposal", describes_delivered_work=False)
    ing.ingest(sid, "a", "a.docx", b"one")
    calls = []
    monkeypatch.setattr(ing, "document_images_from", lambda *a, **k: calls.append(1) or [])
    assert ing.ingest(sid, "a", "a.docx", b"one") == "skipped" and calls == [] and images_of(sid) == []


def test_stale_candidates_are_gone_the_moment_the_new_version_is_saved(env, monkeypatch):
    _, sid = env
    monkeypatch.setattr(ing, "document_images_from", lambda *a, **k: [(b"p", 500, 400, "h")])
    ing.ingest(sid, "a", "a.docx", b"one")
    assert len(images_of(sid)) == 1
    seen = []

    def spy(doc_id, data, name, wanted):  # what a reader sees after the commit, before images are stored
        seen.append(len(images_of(sid)))
    monkeypatch.setattr(ing, "store_images", spy)
    ing.ingest(sid, "a", "a.docx", b"two")
    assert seen == [0]
