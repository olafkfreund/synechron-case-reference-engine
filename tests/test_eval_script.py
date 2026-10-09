import importlib.util
import io
from pathlib import Path
from types import SimpleNamespace

from docx import Document

from app import llm

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("eval_extraction", ROOT / "scripts" / "eval_extraction.py")
ev = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ev)

SECRET = "Zorbulon Reconciliation Platform cut settlement time from 12 days to 3 days"


def docs_dir(tmp_path):
    d = Document()
    d.add_heading("Zorbulon case", 1)
    d.add_paragraph(SECRET + ".")
    buf = io.BytesIO()
    d.save(buf)
    (tmp_path / "case.docx").write_bytes(buf.getvalue())
    return tmp_path


def fake(calls):
    def completion(**kw):
        calls.append(kw)
        body = ('{"kind": "case", "describes_delivered_work": true}' if "kind" in kw["messages"][0]["content"][0]["text"]
                else '{"items": [{"field": "title", "value": "Zorbulon", "quote": "%s"}, '
                     '{"field": "outcome", "value": "settlement time: 12 days to 3 days", "quote": "%s"}, '
                     '{"field": "bogus", "value": "x", "quote": "y"}], "summary": "Settlement fell."}' % (SECRET, SECRET))
        return SimpleNamespace(choices=[SimpleNamespace(finish_reason="stop", message=SimpleNamespace(content=body))])
    return completion


def test_repo_path_refused(capsys):
    assert ev.main(["--docs", str(ROOT / "tests"), "--models", "bedrock/m"]) == 2
    assert "outside the repository" in capsys.readouterr().err


def test_repo_path_refused_via_dotdot_and_symlink(tmp_path, monkeypatch):
    monkeypatch.chdir(ROOT)
    assert ev.main(["--docs", "tests/../tests", "--models", "bedrock/m"]) == 2
    (tmp_path / "link").symlink_to(ROOT / "tests")
    assert ev.main(["--docs", str(tmp_path / "link"), "--models", "bedrock/m"]) == 2


def test_inherited_destination_cannot_override_policy(tmp_path, monkeypatch, capsys):
    calls = []
    monkeypatch.setattr(llm.litellm, "completion", fake(calls))
    monkeypatch.setenv("EXTRACT_MODEL_OPTIONS", '{"destination": "local"}')
    assert ev.main(["--docs", str(docs_dir(tmp_path)), "--models", "openai/gpt-x"]) == 2
    assert "destination" in capsys.readouterr().err and not calls


def test_third_party_model_on_confidential_is_refused(tmp_path, monkeypatch, capsys):
    calls = []
    monkeypatch.setattr(llm.litellm, "completion", fake(calls))
    assert ev.main(["--docs", str(docs_dir(tmp_path)), "--models", "openai/gpt-x"]) == 2
    assert "data policy" in capsys.readouterr().err and not calls


def test_output_has_metrics_and_no_document_text(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(llm.litellm, "completion", fake([]))
    assert ev.main(["--docs", str(docs_dir(tmp_path)), "--models", "bedrock/m"]) == 0
    out = capsys.readouterr()
    assert "kind=case" in out.out and "sourced=2/2" in out.out and "1 malformed" in out.out and "== bedrock/m" in out.out
    assert "bad_metrics=0" in out.out and "mean bad_metrics" in out.out
    for text in (out.out, out.err):
        assert "Zorbulon" not in text and "settlement" not in text.lower()
