import shutil
from pathlib import Path
import subprocess
import uuid
from io import BytesIO

import pytest
from PIL import Image
from psycopg.types.json import Jsonb

from app import db, diagrams, office
from app.schema import Sourced

TITLE = Sourced[str](value="T", source_quote="q").model_dump()


@pytest.fixture
def make():
    db.init()
    sids = []

    def make(raw, diagram_entry=True):
        """A case with one pending diagram holding `raw`; returns (case_id, diagram_id)."""
        with db.connect() as c:
            sid = c.execute("insert into sources(kind,name) values ('s3',%s) returning id", (uuid.uuid4().hex,)).fetchone()[0]
            sids.append(sid)
            did = c.execute("insert into documents(source_id,external_id,title,checksum,text) "
                            "values (%s,'k','d','x','t') returning id", (sid,)).fetchone()[0]
            cid = c.execute("insert into cases(document_id,data) values (%s,%s) returning id",
                            (did, Jsonb({"title": TITLE}))).fetchone()[0]
            gid = c.execute("insert into diagrams(case_id,raw,created_by) values (%s,%s,'t') returning id", (cid, raw)).fetchone()[0]
            if diagram_entry:
                c.execute("update cases set data = data || %s where id=%s", (Jsonb({"diagrams": [
                    {"id": gid, "caption": "c", "status": "pending", "reason": "", "width": 0, "height": 0}]}), cid))
            return cid, gid
    yield make
    with db.connect() as c:
        for sid in sids:
            c.execute("delete from sources where id=%s", (sid,))


def encode(fmt, size=(40, 20), **kw):
    b = BytesIO()
    Image.new("RGB", size, "red").save(b, fmt, **kw)
    return b.getvalue()


def state(cid, gid):
    with db.connect() as c:
        entry = c.execute("select data->'diagrams'->0 from cases where id=%s", (cid,)).fetchone()[0]
        raw, png = c.execute("select raw, png from diagrams where id=%s", (gid,)).fetchone()
    return entry, raw, png


def version(cid):
    with db.connect() as c:
        return c.execute("select md5(data::text) from cases where id=%s", (cid,)).fetchone()[0]


def test_png_and_jpeg_become_png_without_metadata(make):
    exif = Image.Exif()
    exif[0x013B] = "Some Owner"  # Artist
    icc = b"\x00" * 128 + b"desc Some ICC Client Ltd"  # a profile whose description is client text
    for fmt in ("PNG", "JPEG"):
        cid, gid = make(encode(fmt, exif=exif.tobytes(), icc_profile=icc) + b"trailing")
        diagrams.process(gid)
        entry, stored, png = state(cid, gid)
        assert (entry["status"], entry["width"], entry["height"], entry["reason"]) == ("ready", 40, 20, "")
        assert stored is None
        im = Image.open(BytesIO(bytes(png)))
        assert im.format == "PNG" and not im.getexif() and "icc_profile" not in im.info
        assert bytes(png).endswith(b"IEND\xaeB`\x82")
        assert b"Some Owner" not in bytes(png) and b"Some ICC Client" not in bytes(png)


def test_exif_orientation_is_applied(make):
    exif = Image.Exif()
    exif[0x0112] = 6  # rotate 90 degrees to display
    cid, gid = make(encode("JPEG", (40, 20), exif=exif.tobytes()))
    diagrams.process(gid)
    entry, _, png = state(cid, gid)
    assert (entry["width"], entry["height"]) == (20, 40) and Image.open(BytesIO(bytes(png))).size == (20, 40)


def test_decoder_crash_rejects_the_diagram(make, monkeypatch):
    cid, gid = make(encode("PNG"))
    monkeypatch.setattr(diagrams, "_decode", lambda raw: (_ for _ in ()).throw(SyntaxError("boom")))
    diagrams.process(gid)
    entry, raw, png = state(cid, gid)
    assert (entry["status"], entry["reason"]) == ("rejected", diagrams.UNREADABLE) and raw is None and png is None


def test_case_gone_is_a_noop(make):
    cid, gid = make(encode("PNG"))
    with db.connect() as c:
        c.execute("delete from cases where id=%s", (cid,))
    diagrams.process(gid)  # no error


def test_pixel_bomb_rejected_before_decode(make, monkeypatch):
    made = [make(encode("PNG", size)) for size in ((9000, 10), (5000, 5000))]  # a side too long; too many pixels
    monkeypatch.setattr(Image.Image, "load", lambda self: pytest.fail("decoded"))
    for cid, gid in made:
        diagrams.process(gid)
        entry, raw, png = state(cid, gid)
        assert (entry["status"], entry["reason"]) == ("rejected", diagrams.TOO_LARGE)
        assert raw is None and png is None


@pytest.mark.parametrize("raw,reason", [
    (b"not an image at all", diagrams.UNREADABLE),
    (encode("GIF"), diagrams.NOT_IMAGE),
    (encode("PNG")[:30], diagrams.UNREADABLE),
])
def test_corrupt_or_wrong_format_rejected_with_fixed_reason(make, raw, reason):
    cid, gid = make(raw)
    diagrams.process(gid)
    entry, stored, png = state(cid, gid)
    assert (entry["status"], entry["reason"]) == ("rejected", reason)
    assert stored is None and png is None


def test_ready_changes_case_version(make):
    cid, gid = make(encode("PNG"))
    before = version(cid)
    diagrams.process(gid)
    assert version(cid) != before


def test_removed_while_processing_deletes_row(make):
    cid, gid = make(encode("PNG"), diagram_entry=False)
    diagrams.process(gid)
    with db.connect() as c:
        assert c.execute("select count(*) from diagrams where id=%s", (gid,)).fetchone()[0] == 0


def test_diagram_deleted_before_processing_is_a_noop(make):
    cid, gid = make(encode("PNG"))
    with db.connect() as c:
        c.execute("delete from diagrams where id=%s", (gid,))
    before = version(cid)
    diagrams.process(gid)
    assert version(cid) == before


# --- images inside a docx/pptx (#212): the zip is hostile ---
import os
import struct
import zipfile
import zlib


def noise_png(side=500):
    b = BytesIO()
    Image.frombytes("RGB", (side, side), os.urandom(side * side * 3)).save(b, "PNG")
    return b.getvalue()


def bomb_png(side=9000):
    """A valid PNG header for a 9000 x 9000 image whose pixels are all zero: ~250 KB, 243 MB decoded."""
    def chunk(tag, body):
        return struct.pack(">I", len(body)) + tag + body + struct.pack(">I", zlib.crc32(tag + body))
    z, out = zlib.compressobj(9), []
    row = b"\x00" * (1 + side * 3)
    for _ in range(side):
        out.append(z.compress(row))
    out.append(z.flush())
    return (diagrams.MAGIC[0] + chunk(b"IHDR", struct.pack(">IIBBBBB", side, side, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", b"".join(out)) + chunk(b"IEND", b""))


def pack(entries, name="x.docx"):
    b = BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in entries.items():
            z.writestr(n, d)
    return b.getvalue()


def test_good_image_found_and_small_or_fake_ones_filtered():
    good = noise_png(500)
    z = pack({"word/media/a.png": good, "word/media/logo.png": noise_png(200),   # 120 KB but only 200 px
              "word/media/fake.png": b"A" * 40_000,                                # not an image
              "word/media/junk.png": diagrams.MAGIC[0] + b"J" * 40_000,            # magic, no image
              "word/media/tiny.png": good[:100], "word/other/b.png": good})       # too small; wrong folder
    got = diagrams.document_images_from(z, "Case.DOCX")
    assert len(got) == 1 and got[0][1:3] == (500, 500) and got[0][0].startswith(diagrams.MAGIC[0]) and got[0][3] == ""
    assert diagrams.document_images_from(z, "case.txt") == [] and diagrams.document_images_from(z, "case.pdf") == []


def test_bomb_is_refused_and_does_not_hide_a_good_image():
    z = pack({"word/media/bomb.png": bomb_png(), "word/media/ok.png": noise_png(450)})
    got = diagrams.document_images_from(z, "x.docx")
    assert [g[1:3] for g in got] == [(450, 450)]


def test_oversized_entry_is_skipped_before_it_is_read(monkeypatch):
    z = pack({"word/media/huge.png": diagrams.MAGIC[0] + b"\0" * (diagrams.MAX_BYTES + 10), "word/media/ok.png": noise_png(450)})
    opened = []
    real = zipfile.ZipFile.open
    monkeypatch.setattr(zipfile.ZipFile, "open", lambda self, n, *a, **k: (opened.append(getattr(n, "filename", n)), real(self, n, *a, **k))[1])
    assert [g[1:3] for g in diagrams.document_images_from(z, "x.docx")] == [(450, 450)]
    assert "word/media/huge.png" not in opened and "word/media/ok.png" in opened


def test_entry_count_past_the_cap_reads_nothing(monkeypatch):
    monkeypatch.setattr(diagrams, "MAX_ENTRIES", 5)
    z = pack({**{f"word/other/{i}.txt": b"x" for i in range(6)}, "word/media/ok.png": noise_png(450)})
    assert diagrams.document_images_from(z, "x.docx") == []


def test_at_most_ten_largest_first_and_corrupt_zip_is_empty():
    z = pack({f"ppt/media/{i}.png": noise_png(420 + i * 10) for i in range(12)})
    got = diagrams.document_images_from(z, "x.pptx")
    assert len(got) == diagrams.MAX_FOUND and got[0][1] == 530 and [g[1] for g in got] == sorted((g[1] for g in got), reverse=True)
    assert diagrams.document_images_from(b"PK\x03\\x04 not a zip", "x.docx") == [] and diagrams.document_images_from(b"", "x.pptx") == []


def test_docx_alt_text_hint_is_anonymised_and_cut():
    import docx
    d = docx.Document()
    run = d.add_paragraph().add_run()
    run.add_picture(BytesIO(noise_png(500)))
    pr = run._r.xpath(".//wp:docPr")[0]
    pr.set("descr", "Target architecture for Zorpale Bank " + "x" * 400)
    out = BytesIO()
    d.save(out)
    clients = [{"name": "Zorpale Bank", "aliases": [], "anonymised_label": "a retail bank", "referenceable": False}]
    (png, w, h, hint), = diagrams.document_images_from(out.getvalue(), "d.docx", clients)
    assert "Zorpale" not in hint and hint.startswith("Target architecture for a retail bank") and len(hint) <= diagrams.MAX_CAPTION
    assert "Zorpale" in diagrams.document_images_from(out.getvalue(), "d.docx")[0][3]  # no registry, no anonymising


def test_pptx_hint_falls_back_to_the_slide_title_and_filename_alt_text_is_ignored():
    import pptx
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Payments platform overview"
    slide.shapes.add_picture(BytesIO(noise_png(500)), 0, 0)  # python-pptx fills descr with "image.png"
    out = BytesIO()
    prs.save(out)
    (_, _, _, hint), = diagrams.document_images_from(out.getvalue(), "d.pptx")
    assert hint == "Payments platform overview"


def docx_with_alt(alt, side=500):
    import docx
    d = docx.Document()
    run = d.add_paragraph().add_run()
    run.add_picture(BytesIO(noise_png(side)))
    run._r.xpath(".//wp:docPr")[0].set("descr", alt)
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def rewrite(data, part, fn):
    src, out = zipfile.ZipFile(BytesIO(data)), BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for i in src.infolist():
            body = src.read(i.filename)
            z.writestr(i.filename, fn(body) if i.filename == part else body)
    return out.getvalue()


CLIENTS = [{"name": "Zorpale Bank", "aliases": [], "anonymised_label": "a retail bank", "referenceable": False}]


def test_invisible_characters_cannot_hide_a_client_name_from_the_anonymiser():
    # U+E000 (private use) is not a format character, so the anonymiser alone would not see through it
    data = docx_with_alt("Plan for Zor\u200bpale Bank core, Zor\ue000pale Bank too")
    (_, _, _, hint), = diagrams.document_images_from(data, "d.docx", CLIENTS)
    assert "Zorpale" not in hint and hint.count("a retail bank") == 2 and "\u200b" not in hint and "\ue000" not in hint


def test_xml_with_a_doctype_or_a_nul_byte_gives_no_hint_but_keeps_the_image():
    data = docx_with_alt("Architecture overview")
    assert diagrams.document_images_from(data, "d.docx")[0][3] == "Architecture overview"
    doctype = rewrite(data, "word/document.xml", lambda b: b.replace(b"<w:document", b'<!DOCTYPE d [<!ENTITY a "b">]><w:document', 1))
    utf16 = rewrite(data, "word/document.xml", lambda b: b.decode().split("?>", 1)[1].encode("utf-16"))  # no encoding declaration: it parses, and hides a DOCTYPE
    for bad in (doctype, utf16):
        (png, w, h, hint), = diagrams.document_images_from(bad, "d.docx")
        assert (w, h, hint) == (500, 500, "")


def test_xml_part_past_max_xml_and_the_total_budget_give_no_hint(monkeypatch):
    data = docx_with_alt("Architecture overview")
    monkeypatch.setattr(diagrams, "MAX_XML", 100)
    assert diagrams.document_images_from(data, "d.docx")[0][3] == ""
    monkeypatch.undo()
    monkeypatch.setattr(diagrams, "MAX_XML_TOTAL", 600)  # the rels file fits, the document part does not
    assert diagrams.document_images_from(data, "d.docx")[0][3] == ""
    monkeypatch.setattr(diagrams, "MAX_XML_TOTAL", 10_000_000)
    assert diagrams.document_images_from(data, "d.docx")[0][3] == "Architecture overview"


def test_decoding_stops_once_enough_images_are_kept(monkeypatch):
    z = pack({f"word/media/{i}.png": noise_png(420 + i) for i in range(13)})
    calls = []
    real = diagrams._decode
    monkeypatch.setattr(diagrams, "_decode", lambda raw: (calls.append(1), real(raw))[1])
    assert len(diagrams.document_images_from(z, "x.docx")) == diagrams.MAX_FOUND and len(calls) == diagrams.MAX_FOUND


def test_a_declared_size_that_lies_is_bounded_and_loses_only_that_entry(monkeypatch):
    z = pack({"word/media/liar.png": diagrams.MAGIC[0] + os.urandom(diagrams.MAX_BYTES + 5000), "word/media/ok.png": noise_png(450)})
    real = zipfile.ZipFile.infolist

    def lying(self):
        infos = real(self)
        for i in infos:
            if i.filename.endswith("liar.png"):
                i.file_size = 40_000  # the real entry is over 5 MB
        return infos
    monkeypatch.setattr(zipfile.ZipFile, "infolist", lying)
    assert [g[1:3] for g in diagrams.document_images_from(z, "x.docx")] == [(450, 450)]


# --- vector (EMF/WMF) candidates (#234)
EMF = b"\x01\x00\x00\x00" + b"\0" * 36 + b" EMF" + b"\0" * 2500
WMF = b"\xd7\xcd\xc6\x9a" + b"\0" * 2500


def page_png(box=(100, 200, 700, 800), size=(794, 1123), color="white"):
    """A page-sized render with one dark drawing in it, as soffice makes."""
    im = Image.new("RGB", size, color)
    if box:
        im.paste((20, 40, 160), box)
    b = BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def fake_convert(monkeypatch, png, calls=None):
    def convert(files, fmt, timeout, wait, limit=None):
        if calls is not None:
            calls.append((dict(files), fmt, timeout, wait, limit))
        return {n: png for n in files}
    monkeypatch.setattr(diagrams.office, "convert", convert)


def test_emf_and_wmf_are_found_by_magic_cropped_and_svg_never_converted(monkeypatch):
    calls = []
    fake_convert(monkeypatch, page_png(), calls)
    z = pack({"word/media/a.bin": EMF, "word/media/b.dat": WMF, "word/media/fake.emf": noise_png(450)[:60] + b"\0" * 3000,
              "word/media/c.svg": b"<svg/>" + b" " * 3000})
    got = diagrams.document_images_from(z, "x.docx")
    assert len(calls) == 1 and sorted(calls[0][0]) == ["v0.emf", "v1.wmf"]  # the larger entry first
    assert calls[0][1:] == ("png", diagrams.VEC_TIMEOUT, diagrams.VEC_WAIT, diagrams.MAX_BYTES)
    assert len(got) == 2 and all((w, h) == (616, 616) for _, w, h, _ in got)  # 600 px of drawing + 8 px margin a side


def test_blank_page_and_thin_drawing_are_skipped(monkeypatch):
    for png in (page_png(box=None), page_png(box=(100, 200, 700, 300))):
        fake_convert(monkeypatch, png)
        assert diagrams.document_images_from(pack({"word/media/a.emf": EMF}), "x.docx") == []


def test_vector_size_and_count_caps(monkeypatch):
    calls = []
    fake_convert(monkeypatch, page_png(), calls)
    z = pack({"word/media/tiny.emf": EMF[:100] + b"\0" * 1000, "word/media/huge.emf": EMF + b"\0" * diagrams.MAX_BYTES})
    assert diagrams.document_images_from(z, "x.docx") == [] and calls == []
    calls.clear()
    z = pack({f"word/media/{i}.emf": EMF + b"\0" * i for i in range(14)})
    diagrams.document_images_from(z, "x.docx")
    assert len(calls[0][0]) == diagrams.MAX_VECTORS


def test_vectors_merge_by_area_and_are_not_converted_with_ten_rasters(monkeypatch, capsys):
    calls = []
    fake_convert(monkeypatch, page_png(), calls)
    z = pack({"word/media/r.png": noise_png(450), "word/media/v.emf": EMF})
    got = diagrams.document_images_from(z, "x.docx")
    assert [g[1] for g in got] == [616, 450]
    calls.clear()
    z = pack({f"word/media/{i}.png": noise_png(420 + i) for i in range(10)} | {"word/media/v.emf": EMF})
    assert len(diagrams.document_images_from(z, "x.docx")) == 10 and calls == []


def test_converter_failure_keeps_rasters_and_logs_the_class_only(monkeypatch, capsys):
    def boom(*a, **k):
        raise RuntimeError("timeout SECRET-NAME")
    monkeypatch.setattr(diagrams.office, "convert", boom)
    got = diagrams.document_images_from(pack({"word/media/r.png": noise_png(450), "word/media/v.emf": EMF}), "x.docx")
    assert [g[1] for g in got] == [450]
    out = capsys.readouterr().out
    assert "RuntimeError" in out and "SECRET" not in out


def test_vector_hint_comes_from_the_alt_text_and_is_anonymised(monkeypatch):
    import docx
    fake_convert(monkeypatch, page_png())
    d = docx.Document()
    d.add_paragraph().add_run().add_picture(BytesIO(noise_png(500)))
    pr = d.element.xpath(".//wp:docPr")[0]
    pr.set("descr", "Flow at Zorpale Bank")
    out = BytesIO()
    d.save(out)
    z = rewrite_png_to_emf(out.getvalue())
    clients = [{"name": "Zorpale Bank", "aliases": [], "anonymised_label": "a retail bank", "referenceable": False}]
    (png, w, h, hint), = diagrams.document_images_from(z, "d.docx", clients)
    assert hint == "Flow at a retail bank"


def rewrite_png_to_emf(data):
    """The docx with its one media part replaced by EMF bytes (the relationship still names it)."""
    b = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as src, zipfile.ZipFile(b, "w") as dst:
        for i in src.infolist():
            dst.writestr(i.filename, EMF if i.filename.startswith("word/media/") else src.read(i.filename))
    return b.getvalue()


@pytest.mark.skipif(not shutil.which("soffice"), reason="needs LibreOffice")
def test_real_conversion_of_a_generated_emf_gives_a_cropped_candidate(tmp_path):
    svg = tmp_path / "d.svg"
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="900" height="600">'
                   '<rect x="0" y="0" width="900" height="600" fill="#102a80"/></svg>')
    emf = office.convert({"d.svg": svg.read_bytes()}, "emf", 60, 30)["d.svg"]
    assert diagrams._vector_ext(emf[:48]) == "emf"
    got = diagrams.document_images_from(pack({"word/media/d.emf": emf + b"\0" * diagrams.VEC_MIN_BYTES}), "x.docx")
    assert len(got) == 1 and min(got[0][1:3]) >= diagrams.MIN_SIDE and got[0][1:3] != (794, 1123)


def test_a_slow_office_is_killed_and_leaves_no_child(tmp_path, monkeypatch):
    script = tmp_path / "slow.sh"
    script.write_text("#!/bin/sh\nsleep 30 &\nsleep 30\n")
    script.chmod(0o755)
    monkeypatch.setattr(office, "SOFFICE", str(script))
    made = []
    real = office.tempfile.mkdtemp
    monkeypatch.setattr(office.tempfile, "mkdtemp", lambda **k: made.append(Path(real(**k))) or str(made[-1]))
    with pytest.raises(RuntimeError, match="timeout"):
        office.convert({"a.emf": EMF}, "png", 0.5, 1)
    assert subprocess.run(["pgrep", "-f", "sleep 30"], capture_output=True).returncode == 1
    assert office.SLOTS._value == 2 and not made[0].exists()  # the slot is back and the temp dir is gone


@pytest.mark.skipif(not shutil.which("soffice"), reason="needs LibreOffice")
def test_malformed_and_truncated_emf_fail_safely_and_keep_the_rasters():
    bad = {"word/media/a.emf": EMF,  # all zeros after the header
           "word/media/b.emf": EMF[:44] + os.urandom(3000),  # random records
           "word/media/c.wmf": WMF[:200] + b"\xff" * 2500}  # truncated placeholder header, junk body
    got = diagrams.document_images_from(pack({"word/media/r.png": noise_png(450)} | bad), "x.docx")
    assert [g[1] for g in got] == [450]


def test_a_failed_conversion_is_reported_apart_from_no_vectors(monkeypatch):
    z = pack({"word/media/r.png": noise_png(450), "word/media/v.emf": EMF})
    for err in ("busy", "timeout", "exit 1"):
        def boom(*a, **k):
            raise RuntimeError(err)
        monkeypatch.setattr(diagrams.office, "convert", boom)
        meta = {}
        assert [g[1] for g in diagrams.document_images_from(z, "x.docx", meta=meta)] == [450]
        assert meta["vector_failed"] is True and len(meta["src_hashes"]) == 1
    fake_convert(monkeypatch, page_png())
    meta = {}
    got = diagrams.document_images_from(z, "x.docx", meta=meta)
    assert meta["vector_failed"] is False and len(meta["src_hashes"]) == len(got) == 2 and all(len(h) == 64 for h in meta["src_hashes"])
    meta = {}
    diagrams.document_images_from(pack({"word/media/r.png": noise_png(450)}), "x.docx", meta=meta)
    assert meta["vector_failed"] is False  # nothing to convert is not a failure


def test_office_limit_drops_an_oversized_output_and_to_pdf_passes_none(tmp_path, monkeypatch):
    script = tmp_path / "fake.sh"
    script.write_text('#!/bin/sh\nwhile [ "$1" != "--outdir" ]; do shift; done\nout=$2; shift 2\nmkdir -p "$out"\n'
                      'for f in "$@"; do b=$(basename "$f"); b=${b%.*}\n'
                      'case $b in big*) n=100;; *) n=10;; esac\nhead -c $n /dev/zero > "$out/$b.png"; done\n')
    script.chmod(0o755)
    monkeypatch.setattr(office, "SOFFICE", str(script))
    files = {"big.x": b"1", "small.x": b"1"}
    assert sorted(office.convert(files, "png", 10, 1, limit=50)) == ["small.x"]
    assert sorted(office.convert(files, "png", 10, 1)) == ["big.x", "small.x"]
    seen = []
    monkeypatch.setattr(office, "convert", lambda *a, **k: seen.append((a, k)) or {"in.docx": b"%PDF"})
    from app import render
    assert render.to_pdf(b"x", ".docx") == b"%PDF" and "limit" not in seen[0][1] and len(seen[0][0]) == 4
