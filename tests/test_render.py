import html
import re
import zipfile
from io import BytesIO

import pytest
from docx import Document
from lxml import etree

from app import db, office, render
from app.schema import Outcome, Prose, ReferenceCase, Sourced
from tests.test_auth import client, env  # noqa: F401
from tests.test_review import DOCS, make  # noqa: F401
from tests.test_search import approved, data, link, reg  # noqa: F401

R = ["g-user", DOCS]


def gen(ids, fmt="md", groups=R, **kw):
    return client(groups).post("/generate", data={"case_ids": ids, "format": fmt}, **kw)


def prev(ids, groups=R, **kw):
    return client(groups).post("/preview", data={"case_ids": ids, "format": "md"}, **kw)  # format is ignored


def doc_text(content):
    from docx.oxml.ns import qn
    d = Document(BytesIO(content))
    parts = ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in d.element.body.iter(qn("w:p"))]
    parts += [p.text for p in d.sections[0].header.paragraphs] + [p.text for p in d.sections[0].footer.paragraphs]
    return "\n".join(parts)


def full(**kw):
    c = data(challenge=Sourced[str](value="Slow manual KYC & checks", source_quote="Hidden quote text here"),
             organisations=["Secret Org Ltd"], needs_attention=["Hidden attention note"], **kw)
    return c.model_copy(update={"duration_months": Sourced[int](value=18, source_quote="q")})


def test_docx_content_and_leaks(approved, reg):
    _, label_id = reg("Zorp", "a UK retailer")
    cid = approved(full())
    link(cid, label_id)
    r = gen([cid], "docx")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    t = doc_text(r.content)
    for want in ("Faster onboarding", "a UK retailer", "Slow manual KYC & checks", "18 months", "Kubernetes",
                 "onboarding: 12 to 3 days", "Onboarding fell from 12 days", "Synechron \u00b7 Confidential"):
        assert want in t, want
    for bad in ("{{", "{%", "Hidden quote", "Secret Org", "Hidden attention", "DRAFT TEMPLATE"):
        assert bad not in t, bad


def add_diagram(cid, caption="System overview", size=(40, 20), status="ready"):
    import json
    from PIL import Image
    buf = BytesIO()
    Image.new("RGB", size, "blue").save(buf, "PNG")
    with db.connect() as c:
        did = c.execute("insert into diagrams(case_id, png, created_by) values (%s,%s,'t') returning id",
                        (cid, buf.getvalue())).fetchone()[0]
        entry = {"id": did, "caption": caption, "status": status, "reason": "", "width": size[0], "height": size[1]}
        c.execute("update cases set data = data || jsonb_build_object('diagrams', coalesce(data->'diagrams', '[]') || %s::jsonb) "
                  "where id=%s", (json.dumps([entry]), cid))
    return did


def test_only_ready_diagrams_reach_outputs(approved):
    cid = approved(full())
    add_diagram(cid, "Ready one")
    add_diagram(cid, "Pending one", status="pending")  # all three have bytes, so only the status keeps them out
    add_diagram(cid, "Rejected one", status="rejected")
    out = {"md": gen([cid], "md").text, "preview": prev([cid]).text, "docx": doc_text(gen([cid], "docx").content)}
    for name, text in out.items():
        assert "Ready one" in text and "Pending one" not in text and "Rejected one" not in text, name
    assert out["preview"].count("<figure>") == 1
    xml = zipfile.ZipFile(BytesIO(gen([cid], "docx").content)).read("word/document.xml").decode()
    assert xml.count("<w:drawing>") == 1


def test_docx_embeds_diagram_with_caption(approved):
    cid = approved(full())
    add_diagram(cid, "System overview", (2000, 1000))  # wider than 160 mm at 96 dpi: scaled down
    r = gen([cid], "docx")
    z = zipfile.ZipFile(BytesIO(r.content))
    xml = z.read("word/document.xml").decode()
    assert "<w:drawing>" in xml and [n for n in z.namelist() if n.startswith("word/media/") and n.endswith(".png")]
    assert 'cx="5760000"' in xml  # 160 mm
    assert "System overview" in doc_text(r.content) and "{{" not in xml
    assert "<w:drawing>" not in zipfile.ZipFile(BytesIO(gen([approved(full())], "docx").content)).read("word/document.xml").decode()


def test_markdown_lists_diagram_captions(approved):
    cid = approved(full())
    assert "## Diagrams" not in gen([cid], "md").text
    add_diagram(cid, "# Big *flow*")
    m = gen([cid], "md").text
    assert "## Diagrams\n\n- \\# Big \\*flow\\* (in the Word, PowerPoint and PDF versions)\n" in m


def test_caption_is_anonymised_and_withheld_fail_closed(approved, reg, monkeypatch):
    name, label_id = reg("Zorp", "a UK retailer")
    cid = approved(full())
    link(cid, label_id)
    add_diagram(cid, f"Flow at {name}")
    m = gen([cid], "md").text
    assert name not in m and "Flow at a UK retailer" in m
    monkeypatch.setattr(render.anonymise, "apply", lambda t, c: t)  # a missed replacement must withhold, not leak
    r = gen([cid], "md")
    assert r.status_code == 409 and r.json()["detail"] == render.WITHHELD and name not in r.text


def engagement():
    return full().model_copy(update={"basis": "engagement", "basis_reason": "executed contract"})


def test_engagement_has_note_and_no_outcomes_in_docx_md_pptx(approved):
    cid = approved(engagement())  # the record still holds an outcome: output must drop it
    note = render.ENGAGEMENT_NOTE
    t = doc_text(gen([cid], "docx").content)
    assert note in t and "Outcomes" not in t and "12 to 3 days" not in t
    m = gen([cid], "md").text
    assert note in m and "Outcomes" not in m and "12 to 3" not in m
    _, out = slides(gen([cid], "pptx").content)
    assert out[0]["Outcomes"] == [note]


def test_delivered_has_no_note(approved):
    cid = approved(full())
    assert render.ENGAGEMENT_NOTE not in doc_text(gen([cid], "docx").content)
    assert render.ENGAGEMENT_NOTE not in gen([cid], "md").text
    assert "Outcomes" in gen([cid], "md").text


def test_empty_metric_renders_value_only(approved):
    cid = approved(full().model_copy(update={"outcomes": [Outcome(metric="", value="cut by half", source_quote="q")]}))
    t = gen([cid]).text
    assert "- cut by half" in t and ": cut by half" not in t


def test_markdown_fields_and_escaping(approved):
    cid = approved(full(title="A *bold* [title]\n# x"))
    t = gen([cid]).text
    assert "# A \\*bold\\* \\[title\\] \\# x" in t and "- **Industry:** Banking" in t and "- Kubernetes" in t
    assert "Hidden quote" not in t and "Secret Org" not in t


def test_client_display_rules(approved, reg):
    ref_name, ref_id = reg("Globex", "a manufacturer", True)
    _, anon_id = reg("Initech", "a software firm")
    a, b, c = approved(), approved(), approved()
    link(a, ref_id)
    link(b, anon_id)
    t = gen([a, b, c]).text
    assert f"*{ref_name}*" in t and "*a software firm*" in t and "*a client*" in t
    with db.connect() as conn:
        row = conn.execute("select user_id, format, case_ids, anonymised from generations order by id desc limit 1").fetchone()
    assert row == ("u1", "md", [a, b, c], True)
    gen([a])
    with db.connect() as conn:
        assert conn.execute("select anonymised from generations order by id desc limit 1").fetchone()[0] is False


def test_protected_name_replaced_or_withheld(approved, reg, monkeypatch):
    name, cid_client = reg("Zorp", "a UK retailer")
    ok = approved(full().model_copy(update={"challenge": Sourced[str](value=f"Built for {name} in 2024", source_quote="q")}))
    t = gen([ok]).text
    assert name not in t and "Built for a UK retailer in 2024" in t
    evasion = name.replace("o", "ö")  # #166: apply() replaces what blocked() finds
    ev = approved(full().model_copy(update={"challenge": Sourced[str](value=f"Built for {evasion}", source_quote="q")}))
    r = gen([ev])
    assert r.status_code == 200 and "Built for a UK retailer" in r.text and evasion not in r.text
    monkeypatch.setattr(render.anonymise, "apply", lambda t, c: t)  # force apply() to miss: the fail-closed path
    bad = approved(full().model_copy(update={"challenge": Sourced[str](value=f"Built for {name}", source_quote="q")}))
    r = gen([bad])
    assert r.status_code == 409 and r.json()["detail"] == render.WITHHELD and name not in r.text
    with db.connect() as conn:
        assert conn.execute("select count(*) from generations where %s = any(case_ids)", (bad,)).fetchone()[0] == 0


def test_referenceable_name_containing_protected_name_uses_its_label(approved, reg):
    p, _ = reg("Zorp", "a retailer")
    n, ref = reg("Zorp", "a logistics firm", True, suffix=" Logistics")
    cid = approved(full().model_copy(update={"challenge": Sourced[str](value=f"Built for {n} in 2024", source_quote="q")}))
    link(cid, ref)
    t = gen([cid]).text
    assert "*a logistics firm*" in t and "Built for a logistics firm in 2024" in t
    assert p not in t and "a retailer Logistics" not in t
    with db.connect() as conn:
        assert conn.execute("select anonymised from generations order by id desc limit 1").fetchone()[0] is True


def test_visibility_and_validation(approved):
    ok = approved()
    assert gen([approved(acl=("g-other",))]).status_code == 404
    assert gen([approved(due="-1 day")]).status_code == 404
    assert gen([ok, 999999999]).status_code == 404
    assert gen([approved(status="extracted", due=None)]).status_code == 404
    assert gen([ok], "pdf").status_code == 400
    assert gen([ok, ok, ok, ok]).status_code == 200  # duplicates collapse
    assert gen([ok, approved(), approved(), approved()]).status_code == 400


def test_post_without_origin_refused(approved):
    cid = approved()
    r = client(R, origin=None).post("/generate", data={"case_ids": [cid], "format": "md"})
    assert r.status_code == 403


def test_preview_shows_the_download_content(approved):
    cid = approved(full())
    r = prev([cid])
    assert r.status_code == 200
    page = " ".join(html.unescape(re.sub(r"<[^>]+>", "", r.text)).split())
    lines = [ln for ln in gen([cid]).text.splitlines() if ln.strip() and ln.strip() != "---"]
    assert len(lines) > 8
    for ln in lines:
        want = " ".join(re.sub(r"^[#\s-]+|[*\\]", "", ln).split())
        assert want in page, want


def test_preview_shows_ready_diagrams_from_the_user_route(approved):
    a, b = approved(full()), approved(full())
    did = add_diagram(b, "Flow <b>&</b>", (900, 100))
    page = prev([a, b]).text
    url = f"/cases/{b}/diagrams/{did}.png"
    assert f'<a href="{url}" target="_blank" rel="noopener" class="zoom" aria-label="Open Flow &lt;b&gt;&amp;&lt;/b&gt; full size"><img src="{url}" alt="Flow &lt;b&gt;&amp;&lt;/b&gt;"></a>' in page
    assert page.count(f'href="{url}"') == 2 and "Open full size" in page and 'width="480"' not in page
    assert "<figcaption>Flow &lt;b&gt;&amp;&lt;/b&gt;</figcaption>" in page and page.count("<figure>") == 1
    assert client(R).get(url).status_code == 200  # the same user can load the image
    assert "<figure>" not in prev([a]).text


def test_preview_is_not_logged(approved):
    cid = approved()

    def count():
        with db.connect() as conn:
            return conn.execute("select count(*) from generations").fetchone()[0]
    before = count()
    assert prev([cid]).status_code == 200 and count() == before
    assert gen([cid]).status_code == 200 and count() == before + 1


def test_preview_enforces_generate_rules(approved):
    ok = approved()
    for bad in (prev([approved(acl=("g-other",))]), prev([approved(due="-1 day")]), prev([ok, 999999999]),
                prev([approved(status="extracted", due=None)])):
        assert bad.status_code == 404 and bad.json()["detail"] == "no such case"  # the route, not an unknown path
    assert prev([ok, ok, ok, ok]).status_code == 200  # duplicates collapse
    assert prev([ok, approved(), approved(), approved()]).status_code == 400
    assert client(R, origin=None).post("/preview", data={"case_ids": [ok], "format": "md"}).status_code == 403


def test_preview_escapes_values(approved):
    t = prev([approved(full(title="<script>x</script>"))]).text
    assert "&lt;script&gt;" in t and "<script>x" not in t


def test_preview_withholds_a_protected_client(approved, reg, monkeypatch):
    name, _ = reg("Zorp", "a UK retailer")
    monkeypatch.setattr(render.anonymise, "apply", lambda t, c: t)  # force apply() to miss: the fail-closed path
    bad = approved(full().model_copy(update={"challenge": Sourced[str](value=f"Built for {name}", source_quote="q")}))
    r = prev([bad])
    assert r.status_code == 409 and name not in r.text


def test_docx_metadata_is_neutral(approved):
    props = Document(BytesIO(gen([approved()], "docx").content)).core_properties
    assert (props.title, props.author, props.comments) == ("Reference cases", "Reference Engine", "")


def test_markdown_values_cannot_create_structure():
    for raw in ("# Heading", "- # x", "2) item", "+ plus", "| a | b |", "~~strike~~"):
        out = render.md(raw)
        assert not out.startswith(("#", "-", "+", "2)", "|", "~")), (raw, out)


from pptx import Presentation  # noqa: E402

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
NAMES = ("Title", "Client", "Summary", "Challenge", "Solution", "Outcomes", "Technology")


def test_markdown_escapes_ampersand():
    assert render.md("Zo&#114;p") == "Zo\\&\\#114;p"  # neither & nor # left bare: no entity is decoded
    assert render.md("R&D") == "R\\&D"


def body(prs):
    """The slides between the brand cover and the closing slide (#282)."""
    s = list(prs.slides)
    assert s[0].slide_layout.name == render.COVER and s[-1].slide_layout.name == render.CLOSING
    return s[1:-1]


def slides(content):
    """[{placeholder name: [paragraph texts]}] per content slide, names taken from the slide's layout."""
    prs = Presentation(BytesIO(content))
    out = []
    for s in body(prs):
        names = {p.placeholder_format.idx: p.name for p in s.slide_layout.placeholders}
        out.append({names[p.placeholder_format.idx]: [x.text for x in p.text_frame.paragraphs] for p in s.placeholders})
    return prs, out


def test_pptx_fills_named_placeholders(approved, reg):
    _, cid_client = reg("Zorp", "a UK retailer")
    a = approved(full().model_copy(update={"solution": Sourced[str](value="Cloud KYC", source_quote="q")}))
    b = approved(data(title="Second case"))
    link(a, cid_client)
    r = gen([a, b], "pptx")
    assert r.status_code == 200 and r.headers["content-type"] == PPTX
    assert r.headers["content-disposition"].endswith('reference-cases.pptx"')
    prs, out = slides(r.content)
    assert len(out) == 2 and all(set(s) == set(NAMES) for s in out)
    first = out[0]
    assert first["Title"] == ["Faster onboarding"] and first["Client"] == ["a UK retailer"]
    assert first["Summary"] == ["Onboarding fell from 12 days to 3 days."]
    assert first["Challenge"] == ["Slow manual KYC & checks"] and first["Solution"] == ["Cloud KYC"]
    assert first["Outcomes"] == ["onboarding: 12 to 3 days"] and first["Technology"] == ["Kubernetes"]
    assert out[1]["Solution"] == [""] and out[1]["Client"] == ["a client"]  # no content: empty, not prompt text
    everything = "\n".join(t for s in out for ps in s.values() for t in ps)
    for bad in ("Click to add", "Hidden quote", "Secret Org", "Hidden attention"):
        assert bad not in everything, bad
    cp = prs.core_properties
    assert (cp.title, cp.author, cp.last_modified_by, cp.comments) == ("Reference cases", "Reference Engine", "Reference Engine", "")


def test_pptx_slide_shaping(approved):
    long = "word " * 120
    c = data(tech=("Kubernetes", "Kafka", "Terraform"), challenge=Sourced[str](value=long.strip(), source_quote="q"))
    c = c.model_copy(update={"outcomes": [Outcome(metric=f"m{i}", value=f"{i}", source_quote="q") for i in range(7)]})
    _, out = slides(gen([approved(c)], "pptx").content)
    assert out[0]["Technology"] == ["Kubernetes, Kafka, Terraform"]  # one line
    assert len(out[0]["Outcomes"]) == render.SLIDE_OUTCOMES
    assert out[0]["Challenge"][0].endswith(" …") and len(out[0]["Challenge"][0]) <= render.SLIDE_TEXT + 2


def test_pptx_protected_name_withheld(approved, reg, monkeypatch):
    name, _ = reg("Zorp", "a UK retailer")
    monkeypatch.setattr(render.anonymise, "apply", lambda t, c: t)
    bad = approved(full().model_copy(update={"challenge": Sourced[str](value=name, source_quote="q")}))
    r = gen([bad], "pptx")
    assert r.status_code == 409 and r.json()["detail"] == render.WITHHELD


def test_pptx_missing_placeholder_or_layout_fails_loudly(tmp_path, monkeypatch):
    prs = Presentation(render.MASTER)
    layout = next(l for l in prs.slide_layouts if l.name == "Reference case")
    next(p for p in layout.placeholders if p.name == "Solution").name = "Soluton"
    broken = tmp_path / "broken.pptx"
    prs.save(broken)
    monkeypatch.setattr(render, "MASTER", broken)
    with pytest.raises(ValueError, match="missing placeholders: Solution"):
        render.to_pptx([])
    layout.name = "Other"
    prs.save(broken)
    with pytest.raises(ValueError, match="no slide layout named 'Reference case'"):
        render.to_pptx([])


def test_master_is_the_brand_template_with_blue_cover_and_closing():  # #282
    prs = Presentation(render.MASTER)
    layouts = {l.name: l for l in prs.slide_layouts}
    assert set(layouts) == {render.COVER, render.CLOSING, render.LAYOUT, render.INDUSTRY_LAYOUT, render.DIAGRAM_LAYOUT}
    for name in (render.COVER, render.CLOSING):  # brand v2 cover option 2: yellow device on Synechron blue
        assert '<a:srgbClr val="002535"/>' in etree.tostring(layouts[name]._element.cSld.bg).decode()
    text = " ".join(sh.text_frame.text for l in prs.slide_layouts for sh in l.shapes if sh.has_text_frame)
    assert "DRAFT" not in text
    assert prs.slide_width == 12192000  # the official 16:9 template


def test_control_character_cannot_rebuild_a_protected_name(approved, reg):
    name, _ = reg("Zorp", "a UK retailer")
    sneaky = name[:2] + "\x01" + name[2:]  # stripped after the check, this used to become the name again
    cid = approved(full().model_copy(update={"challenge": Sourced[str](value=sneaky, source_quote="q")}))
    for fmt in ("md", "docx", "pptx"):
        r = gen([cid], fmt)
        assert r.status_code == 200, fmt  # stripped first, then replaced by the label
        body = r.content if fmt == "md" else b"".join(
            z.read(n) for z in [zipfile.ZipFile(BytesIO(r.content))] for n in z.namelist())
        assert name.encode() not in body and b"a UK retailer" in body, fmt


def test_lowercase_short_alias_is_rewritten_not_withheld():
    cl = {"id": 1, "name": "Union Bank of Zeta", "aliases": ["UBZ"], "anonymised_label": "a Swiss bank", "referenceable": False}
    out = render.protect([{"text": "Contact ops@ubz.example and the Ubz team"}], [cl])
    assert "ubz" not in out[0]["text"].lower()


def test_master_layout_has_unique_idx_and_no_slides():
    prs = Presentation(render.MASTER)
    layout = next(l for l in prs.slide_layouts if l.name == "Reference case")
    idx = [p.placeholder_format.idx for p in layout.placeholders]
    assert len(idx) == len(set(idx)) and len(prs.slides) == 0
    assert {p.name for p in layout.placeholders} == set(NAMES)


def test_extra_layout_placeholder_is_removed_and_sample_slides_refused(tmp_path, monkeypatch, approved):
    import copy
    prs = Presentation(render.MASTER)
    layout = next(l for l in prs.slide_layouts if l.name == "Reference case")
    extra = copy.deepcopy(next(p for p in layout.placeholders if p.name == "Summary")._element)
    extra.nvSpPr.cNvPr.set("name", "Logo")
    extra.nvSpPr.nvPr.ph.set("idx", "99")
    layout.shapes._spTree.append(extra)
    extended = tmp_path / "extended.pptx"
    prs.save(extended)
    monkeypatch.setattr(render, "MASTER", extended)
    out = Presentation(BytesIO(render.to_pptx([render.section(data(), "a client")])))
    assert len(list(body(out)[0].placeholders)) == len(NAMES)  # the logo placeholder was removed
    prs.slides.add_slide(layout)  # a master that ships its own sample slide
    prs.save(extended)
    with pytest.raises(ValueError, match="no slides"):
        render.to_pptx([])


def test_broken_template_writes_no_audit_row(approved, monkeypatch):
    cid = approved()
    monkeypatch.setattr(render, "MASTER", render.MASTER.with_name("missing.pptx"))
    with db.connect() as c:
        before = c.execute("select count(*) from generations").fetchone()[0]
    with pytest.raises(Exception):
        gen([cid], "pptx")
    with db.connect() as c:
        assert c.execute("select count(*) from generations").fetchone()[0] == before


import threading  # noqa: E402

import pypdfium2 as pdfium  # noqa: E402


def pdf_text(content):
    pdf = pdfium.PdfDocument(content)
    return "\n".join(p.get_textpage().get_text_range() for p in pdf), pdf.get_metadata_dict()


@pytest.mark.parametrize("fmt", ["pdf_docx", "pdf_pptx"])
def test_pdf_from_route_has_label_not_name(approved, reg, fmt):
    name, cid_client = reg("Zorp", "a UK retailer")
    c = approved(full().model_copy(update={"challenge": Sourced[str](value=f"Built for {name} in 2024", source_quote="q")}))
    link(c, cid_client)
    r = gen([c], fmt)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"].endswith('reference-cases.pdf"') and r.content.startswith(b"%PDF")
    text, meta = pdf_text(r.content)
    assert "Faster onboarding" in text and "a UK retailer" in text and name not in text
    assert name not in " ".join(meta.values()) and "Secret Org" not in text and "Hidden quote" not in text
    assert meta.get("Author") == "Reference Engine"


def test_pdf_failures_are_clear_and_leak_nothing(monkeypatch):
    docx = render.to_docx([])
    monkeypatch.setattr(render, "PDF_TIMEOUT", 0.01)
    with pytest.raises(render.PdfError, match="timed out"):
        render.to_pdf(docx, ".docx")
    monkeypatch.setattr(render, "PDF_TIMEOUT", 60)
    monkeypatch.setattr(office, "SOFFICE", "false")
    with pytest.raises(render.PdfError, match=r"failed \(exit 1\)"):
        render.to_pdf(docx, ".docx")
    monkeypatch.setattr(office, "SOFFICE", "true")  # exits 0 but writes nothing
    with pytest.raises(render.PdfError, match="failed"):
        render.to_pdf(docx, ".docx")


def test_pdf_route_reports_503_and_no_audit_row(approved, monkeypatch):
    cid = approved()
    monkeypatch.setattr(office, "SOFFICE", "false")
    r = gen([cid], "pdf_docx")
    assert r.status_code == 503 and "PDF conversion failed" in r.json()["detail"]
    with db.connect() as conn:
        assert conn.execute("select count(*) from generations where %s = any(case_ids)", (cid,)).fetchone()[0] == 0


def test_parallel_conversions_do_not_collide():
    docx = render.to_docx([])
    results, errors = [], []

    def run():
        try:
            results.append(render.to_pdf(docx, ".docx"))
        except Exception as e:  # noqa: BLE001
            errors.append(e)
    ts = [threading.Thread(target=run) for _ in range(3)]
    [t.start() for t in ts]
    [t.join() for t in ts]
    assert not errors and len(results) == 3 and all(r.startswith(b"%PDF") for r in results)


def test_pdf_conversion_does_not_fetch_linked_content():
    """A marketing template could link images by URL; soffice must not fetch them."""
    import http.server
    import threading as th
    hits = []

    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            hits.append(self.path)
            self.send_response(404)
            self.end_headers()

        def log_message(self, *a):
            pass
    srv = http.server.HTTPServer(("127.0.0.1", 0), H)
    th.Thread(target=srv.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{srv.server_port}/linked.png"
    d = Document()
    d.add_paragraph("before")
    # a linked (not embedded) picture, as Word writes it: external relationship + r:link on the blip
    rid = d.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image", is_external=True)
    from docx.oxml import parse_xml
    d.paragraphs[0]._p.append(parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:drawing>'
        '<wp:inline><wp:extent cx="914400" cy="914400"/><wp:docPr id="1" name="linked"/><a:graphic>'
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic>'
        '<pic:nvPicPr><pic:cNvPr id="1" name="linked"/><pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip r:link="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic>'
        '</wp:inline></w:drawing></w:r>'))
    buf = BytesIO()
    d.save(buf)
    try:
        assert render.to_pdf(buf.getvalue(), ".docx").startswith(b"%PDF")
    finally:
        srv.shutdown()
    assert hits == []


def test_pdf_busy_when_slots_are_taken(monkeypatch):
    monkeypatch.setattr(render, "PDF_WAIT", 0.1)
    held = [office.SLOTS.acquire() for _ in range(2)]
    try:
        with pytest.raises(render.PdfError, match="busy"):
            render.to_pdf(b"x", ".docx")
    finally:
        for _ in held:
            office.SLOTS.release()


def test_invisible_character_in_name_is_rewritten_not_shipped():
    cl = [{"id": 1, "name": "Zorp Bank", "aliases": ["Zorp"], "anonymised_label": "a bank", "referenceable": False}]
    assert render.protect([{"summary": "Run for Zo\u00adrp"}], cl)[0]["summary"] == "Run for a bank"


def sections(**kw):
    S = lambda v: Sourced[str](value=v, source_quote="q")  # noqa: E731
    return full(implementation=S("Phased rollout"), configuration=S("SSO via OIDC"), lessons_learned=S("Start earlier"), **kw)


def test_new_sections_in_word_and_markdown(approved):
    cid = approved(sections())
    for t in (doc_text(gen([cid], "docx").content), gen([cid], "md").text):
        for want in ("Implementation approach", "Phased rollout", "Configuration and integration", "SSO via OIDC",
                     "Lessons learned", "Start earlier"):
            assert want in t, want


def test_engagement_output_has_no_lessons_learned(approved):
    cid = approved(sections(basis="engagement", basis_reason="executed contract"))
    m = gen([cid], "md").text
    assert "Phased rollout" in m and "Lessons learned" not in m and "Start earlier" not in m


def test_pptx_unchanged_by_new_sections(approved):
    plain, rich = approved(full()), approved(sections())
    _, a = slides(gen([plain], "pptx").content)
    _, b = slides(gen([rich], "pptx").content)
    assert a == b


def test_new_section_naming_client_is_anonymised(approved, reg):
    name, label_id = reg("Zorp", "a UK retailer")
    cid = approved(full(implementation=Sourced[str](value=f"Rolled out at {name}", source_quote="q")))
    link(cid, label_id)
    t = gen([cid]).text
    assert name not in t and "Rolled out at a UK retailer" in t


def test_pptx_diagram_slide_per_diagram(approved):
    cid = approved(full())
    add_diagram(cid, "System overview", (400, 100))  # wider than the box ratio: fitted by width, not cropped
    add_diagram(cid, "Second", (50, 100))
    prs, out = slides(gen([cid], "pptx").content)
    assert [set(s) for s in out] == [set(NAMES), {"Title", "Caption"}, {"Title", "Caption"}]
    assert out[1]["Caption"] == ["System overview"] and out[1]["Title"] == ["Faster onboarding"]
    pics = [[sh for sh in s.shapes if sh.shape_type == 13] for s in body(prs)]
    assert [len(p) for p in pics] == [0, 1, 1]
    wide, tall = pics[1][0], pics[2][0]
    assert abs(wide.width / wide.height - 4) < .01 and abs(tall.width / tall.height - .5) < .01
    assert wide.width <= 12.5 * 914400 and tall.height <= 5 * 914400


def test_pptx_without_diagram_layout_fails_only_with_diagrams(tmp_path, monkeypatch):
    prs = Presentation(render.MASTER)
    next(l for l in prs.slide_layouts if l.name == "Diagram").name = "Other"
    odd = tmp_path / "odd.pptx"
    prs.save(odd)
    monkeypatch.setattr(render, "MASTER", odd)
    sec = lambda diagrams: dict(title="T", client="C", summary="S", note="", blocks=[], lists=[], diagrams=diagrams)  # noqa: E731
    d = dict(id=1, caption="c", width=10, height=10)
    render.to_pptx([sec([])])  # no diagrams: unaffected
    render.to_pptx([sec([d])])  # a diagram with no image (rejected or missing) is skipped
    with pytest.raises(ValueError, match="no slide layout named 'Diagram'"):
        render.to_pptx([sec([d])], images={1: b"x"})


# --- technical references (#212 phase D) ---
def with_refs(*values, **kw):
    return full(**kw).model_copy(update={"tech_references": [Sourced[str](value=v, source_quote="q") for v in values]
                                         + [Sourced[str](value="Unsourced ref", source_quote="q", unsourced=True)]})


def test_technical_references_render_in_md_docx_and_preview_without_unsourced_ones(approved):
    cid = approved(with_refs("ISO 20022", "OAuth 2.0"))
    md, preview, docx = gen([cid], "md").text, prev([cid]).text, doc_text(gen([cid], "docx").content)
    for name, text in (("md", md), ("preview", preview), ("docx", docx)):
        assert "Technical references" in text and "ISO 20022" in text and "OAuth 2.0" in text, name
        assert "Unsourced ref" not in text, name
    assert md.index("Technology") < md.index("Technical references")
    assert "Technical references" not in gen([approved(full())], "md").text  # none: the list is dropped


def test_technical_references_are_escaped_and_the_pptx_has_none(approved):
    cid = approved(with_refs("<script>x</script> & y"))
    assert "&lt;script&gt;" in prev([cid]).text and "<script>x" not in prev([cid]).text
    assert "\\<script" in gen([cid], "md").text or "&lt;script" in gen([cid], "md").text
    from pptx import Presentation
    deck = Presentation(BytesIO(gen([cid], "pptx").content))
    assert "script" not in " ".join(s.text_frame.text for sl in deck.slides for s in sl.shapes if s.has_text_frame)


def test_a_reference_naming_a_protected_client_is_anonymised_or_withheld(approved, reg, monkeypatch):
    name, label_id = reg("Zorp", "a UK retailer")
    cid = approved(with_refs(f"{name} payments API v2"))
    link(cid, label_id)
    m = gen([cid], "md").text
    assert name not in m and "a UK retailer payments API v2" in m
    monkeypatch.setattr(render.anonymise, "apply", lambda t, c: t)  # a missed replacement must withhold
    r = gen([cid], "md")
    assert r.status_code == 409 and name not in r.text


def test_section_prose_replaces_value_only_when_clean():
    from app.schema import Prose, ReferenceCase
    c = ReferenceCase.model_validate({"title": Sourced[str](value="T", source_quote="T").model_dump()})
    c.challenge = Sourced[str](value="one line", source_quote="one line")
    c.solution = Sourced[str](value="sol line", source_quote="sol line")
    c.prose = {"challenge": Prose(text="A\n\n\nB\n\n"), "solution": Prose(text="X", flags=["number"])}
    b = {x["heading"]: x for x in render.section(c, "C")["blocks"]}
    assert b["Challenge"]["paragraphs"] == ["A", "B"]
    assert b["Solution"]["text"] == "sol line" and b["Solution"]["paragraphs"] == ["sol line"]


def test_placeholder_metric_is_hidden_in_outputs():
    from app.schema import ReferenceCase
    c = ReferenceCase(title=Sourced[str](value="T", source_quote="T"), outcomes=[
        Outcome(metric="metric", value="flagged 17 percent more", source_quote="q"),
        Outcome(metric="Fraud", value="up 17%", source_quote="q")])
    out = render.section(c, "C")["lists"][0]["bullets"]
    assert out == ["flagged 17 percent more", "Fraud: up 17%"]


def test_docx_templates_loop_paragraphs():
    from docxtpl import DocxTemplate
    sec = {"title": "T", "client": "C", "summary": "S", "note": "", "details": [], "lists": [], "diagrams": [],
           "blocks": [{"heading": "Challenge", "text": "A\n\nB", "paragraphs": ["A", "B"]}]}
    for path in (render.TEMPLATE, render.TEMPLATE.with_name("case_study.docx")):
        tpl = DocxTemplate(path)
        tpl.render({"cases": [{**sec, "footer": "F"}], "industry": None}, autoescape=True)
        buf = BytesIO()
        tpl.save(buf)
        paras = [p.text for p in Document(buf).paragraphs]
        assert "A" in paras and "B" in paras, path


def test_case_study_docx_pdf_footer_and_one_case(approved, reg):
    name, label_id = reg("Zorp", "a UK retailer")
    c = approved(full().model_copy(update={"challenge": Sourced[str](value=f"Built for {name}", source_quote="q")}))
    link(c, label_id)
    r = gen([c], "case_study_docx")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('case-study.docx"')
    t = doc_text(r.content)
    assert "Faster onboarding" in t and "At a glance" in t and "Based on 1 source document reviewed" in t
    assert name not in t and "{{" not in t and "{%" not in t
    assert [p.text for p in Document(BytesIO(r.content)).paragraphs] and Document(BytesIO(r.content)).core_properties.author == "Reference Engine"
    pdf = gen([c], "case_study_pdf")
    assert pdf.headers["content-disposition"].endswith('case-study.pdf"') and pdf.content.startswith(b"%PDF")
    assert gen([c, approved(full())], "case_study_docx").status_code == 400
    one = html.unescape(prev([c]).text)
    note = re.search(r"Case study downloads also end with: (Based on 1 source document reviewed[^<]*)", one)
    assert note and note[1] in t and one.count("Based on 1 source document") == 1  # the note only, not the bid section
    assert "Case study downloads also end" not in html.unescape(prev([c, approved(full())]).text)


def test_case_study_withheld_gives_409(approved, reg, monkeypatch):
    c = approved(full())
    monkeypatch.setattr(render.anonymise, "blocked", lambda *a, **k: True)
    assert gen([c], "case_study_docx").status_code == 409



def test_case_study_docx_full_and_anonymised(approved, reg):
    from app.schema import Prose
    name, label_id = reg("Zorp", "a UK retailer")
    c = with_refs(f"{name} payments API v2")
    c = c.model_copy(update={"prose": {"challenge": Prose(text=f"{name} was slow.\n\nSecond paragraph.")}})
    cid = approved(c)
    link(cid, label_id)
    add_diagram(cid, f"Flow at {name}")
    r = gen([cid], "case_study_docx")
    assert r.status_code == 200
    t = doc_text(r.content)
    for want in ("Diagrams", "Technical references", "a UK retailer was slow.", "Second paragraph.",
                 "Flow at a UK retailer", "payments API v2", "Based on 1 source document"):
        assert want in t, want
    assert name not in t and "{{" not in t and "{%" not in t
    assert any(n.startswith("word/media/") for n in zipfile.ZipFile(BytesIO(r.content)).namelist())


def test_by_reviewer_prose_is_used_despite_flags_and_other_prose_is_not():
    from app.schema import Prose, ReferenceCase
    c = ReferenceCase.model_validate({"title": Sourced[str](value="T", source_quote="T").model_dump()})
    c.challenge = Sourced[str](value="headline", source_quote="headline")
    for mine, want in ((True, "my text"), (False, "headline")):
        c.prose = {"challenge": Prose(text="my text", flags=["a name"], by_reviewer=mine)}
        assert render.section(c, "C")["blocks"][0]["text"] == want


def test_headline_value_is_one_unsplit_paragraph_in_md_and_docx(approved):
    value = "  first part\n\nsecond part  "
    cid = approved(full().model_copy(update={"challenge": Sourced[str](value=value, source_quote="q")}))
    c = ReferenceCase.model_validate(db.connect().execute("select data from cases where id=%s", (cid,)).fetchone()[0])
    assert render.section(c, "C")["blocks"][0]["paragraphs"] == [c.challenge.value]
    assert "first part second part" in gen([cid], "md").text  # md() folds whitespace, as on main
    paras = [p.text for p in Document(BytesIO(gen([cid], "docx").content)).paragraphs]
    assert not {"first part", "second part"} & {p.strip() for p in paras}
    assert any("first part" in p and "second part" in p for p in paras)


# --- solution architecture (#246 step A5) ---
def with_arch(*names, flags=(), label="Open account", by_reviewer=False, **kw):
    from app.schema import Architecture
    names = names or ("Customer portal", "Account engine")
    comps = [{"id": f"c{i}", "name": n, "role": "Does the work", "zone": z, "kind": "service"}
             for i, (n, z) in enumerate(zip(names, ("channels", "core", "data", "external")))]
    flows = [{"src": "c0", "dst": "c1", "label": label, "mode": "api"}]
    arch = Architecture.model_validate({"components": comps, "flows": flows, "flags": list(flags), "by_reviewer": by_reviewer,
                                        "callouts": [{"text": "Onboarding in 3 days", "outcome": 0}]})
    return full(**kw).model_copy(update={"architecture": arch})


def pptx_arch(content):
    prs = Presentation(BytesIO(content))
    return [s for s in body(prs) if any(sh.name.startswith("arch:") for sh in s.shapes)]


def test_architecture_is_in_every_format_and_the_preview(approved):
    cid = approved(with_arch())
    md = gen([cid], "md").text
    assert "## Architecture" in md and "```mermaid" in md and "Account engine" in md and "Open account" in md
    z = zipfile.ZipFile(BytesIO(gen([cid], "docx").content))
    assert any(n.startswith("word/media/") for n in z.namelist())  # the PNG; this case has no other picture
    assert "Architecture" in doc_text(gen([cid], "docx").content)
    cs = zipfile.ZipFile(BytesIO(gen([cid], "case_study_docx").content))
    assert any(n.startswith("word/media/") for n in cs.namelist())
    slides_ = pptx_arch(gen([cid], "pptx").content)
    assert len(slides_) == 1
    text = " ".join(sh.text_frame.text for sh in slides_[0].shapes if sh.has_text_frame)
    assert "solution architecture" in text and "Account engine" in text and "Open account" in text
    p = prev([cid]).text
    assert 'class="zoom" aria-label="Open Solution architecture full size"><svg xmlns=' in p and "Account engine" in p and "Open account" in p
    assert all(n in md and n in p for n in ("Customer portal", "Account engine", "Open account"))  # the preview shows what is downloaded


def test_architecture_names_are_anonymised_like_the_rest(approved, reg):
    name, label_id = reg("Zorp", "a UK retailer")
    cid = approved(with_arch(f"{name} ledger", "Account engine"))
    link(cid, label_id)
    outs = [gen([cid], "md").text, prev([cid]).text, doc_text(gen([cid], "docx").content)]
    pp = pptx_arch(gen([cid], "pptx").content)[0]
    outs.append(" ".join(sh.text_frame.text for sh in pp.shapes if sh.has_text_frame))
    for o in outs:
        assert name not in o and "a UK retailer ledger" in o or o is outs[2], o[:200]  # the docx has no text of the drawing, only its PNG
    assert all(name not in o for o in outs)


def test_a_blocked_architecture_is_left_out_not_the_whole_output(approved, reg, monkeypatch):
    name, _ = reg("Zorp", "a UK retailer")
    cid = approved(with_arch(f"{name} ledger", "Account engine"))
    real = render.anonymise.apply
    monkeypatch.setattr(render.anonymise, "apply", lambda t, c: t if "ledger" in t else real(t, c))  # apply() misses the component
    r = gen([cid], "md")
    assert r.status_code == 200 and name not in r.text and "## Architecture" not in r.text
    p = prev([cid])
    assert p.status_code == 200 and "left out" in p.text and 'class="paper"' not in p.text and name not in p.text


def test_unapproved_or_flagged_architecture_is_not_output(approved):
    assert gen([approved(with_arch(), status="extracted", due=None)], "md").status_code == 404
    flagged = approved(with_arch(flags=["a number not in the quote"]))
    assert "## Architecture" not in gen([flagged], "md").text and 'class="paper"' not in prev([flagged]).text
    assert not pptx_arch(gen([flagged], "pptx").content)


def test_a_reviewer_confirmed_architecture_renders_despite_flags_but_a_price_never_does(approved):
    ok = approved(with_arch(flags=["a number not in the quote"], by_reviewer=True))
    assert "## Architecture" in gen([ok], "md").text
    priced = approved(with_arch(flags=["component 'X': contains a price or payment term"], by_reviewer=True))
    assert "## Architecture" not in gen([priced], "md").text


def test_architecture_preview_escapes_a_hostile_label(approved):
    cid = approved(with_arch("<script>alert(1)</script>", "Account engine", label="<img src=x onerror=1>"))
    p = prev([cid]).text
    assert 'class="zoom" aria-label="Open Solution architecture full size"><svg xmlns=' in p and "<script>alert" not in p and "<img src=x" not in p and "&lt;script&gt;" in p
    md = gen([cid], "md").text
    assert "<script>" not in md and "#60;script#62;" in md


def test_an_architecture_meets_the_completeness_diagram_item():
    from app import completeness
    item = lambda c: next(i for i in completeness.checklist(c) if i["key"] == "diagram")
    assert not item(full())["ok"] and item(with_arch())["ok"]


def test_architecture_png_route(approved, reg):
    from PIL import Image
    cid = approved(with_arch())
    r = client(R).get(f"/cases/{cid}/architecture.png")
    assert r.status_code == 200 and r.headers["content-type"] == "image/png"
    assert r.headers["cache-control"] == "private, no-store" and r.headers["x-content-type-options"] == "nosniff"
    assert "default-src 'none'" in r.headers["content-security-policy"]
    im = Image.open(BytesIO(r.content))
    assert im.width == 3200 and im.height < 1800  # cropped to the content
    url = f"/cases/{cid}/architecture.png"
    page = prev([cid]).text
    assert f'<a href="{url}" target="_blank" rel="noopener" class="zoom"' in page and page.count(f'href="{url}"') == 2 and "Open full size" in page


def test_architecture_png_route_answers_404_alike_and_403_without_a_role(approved, reg, monkeypatch):
    ok = approved(with_arch())
    miss = [client(R).get("/cases/999999/architecture.png"),
            client(["g-user"]).get(f"/cases/{ok}/architecture.png"),  # no access to the source document
            client(R).get(f"/cases/{approved(with_arch(), status='extracted', due=None)}/architecture.png"),  # not approved
            client(R).get(f"/cases/{approved(with_arch(flags=['unresolved']))}/architecture.png"),  # flagged
            client(R).get(f"/cases/{approved()}/architecture.png")]  # none
    assert [r.status_code for r in miss] == [404] * 5 and len({r.text for r in miss}) == 1
    name, _ = reg("Zorp", "a UK retailer")
    blocked = approved(with_arch(f"{name} ledger", "Account engine"))
    real = render.anonymise.apply
    monkeypatch.setattr(render.anonymise, "apply", lambda t, c: t if "ledger" in t else real(t, c))  # apply() misses: blocked() must omit it
    r = client(R).get(f"/cases/{blocked}/architecture.png")
    assert r.status_code == 404 and r.text == miss[0].text
    assert client([]).get(f"/cases/{ok}/architecture.png").status_code == 403


def test_a_long_expanded_client_label_does_not_drop_the_architecture(approved, reg):
    name, label_id = reg("Zorp", "a UK retailer with a very long anonymised label that expands every client name past forty characters")
    cid = approved(with_arch(f"{name} ledger", "Account engine", label=f"Open {name} account"))
    link(cid, label_id)
    md = gen([cid], "md").text
    assert "## Architecture" in md and name not in md and "a very long anonymised label" in md
    p = prev([cid]).text
    assert 'class="paper"' in p and "left out" not in p and name not in p
    assert gen([cid], "docx").status_code == 200 and pptx_arch(gen([cid], "pptx").content)
    assert client(R).get(f"/cases/{cid}/architecture.png").status_code == 200


def test_a_client_named_like_a_zone_does_not_rewrite_the_structure(approved):
    with db.connect() as c:  # reg() adds a random tag to every name; a client really called "core" needs a row of its own
        c.execute("delete from clients where name = 'core'")
        label_id = c.execute("insert into clients(name, anonymised_label, referenceable) values ('core', 'a UK retailer', false) returning id").fetchone()[0]
    try:
        cid = approved(with_arch("core ledger", "Account engine"))
        link(cid, label_id)
        md = gen([cid], "md").text
        assert "## Architecture" in md and 'subgraph core["Core"]' in md  # the zone literal is untouched
        assert "a UK retailer ledger" in md and "core ledger" not in md
        assert 'class="paper"' in prev([cid]).text
    finally:
        with db.connect() as c:
            c.execute("update cases set client_id = null where client_id = %s", (label_id,))
            c.execute("delete from clients where id = %s", (label_id,))


def test_a_structure_that_cannot_be_kept_is_left_out_with_a_note(approved, monkeypatch):
    cid = approved(with_arch())
    monkeypatch.setattr(render, "architecture_of", lambda sec: None)
    p = prev([cid]).text
    assert "left out" in p and 'class="paper"' not in p


def test_callouts_follow_the_shown_outcomes_and_a_contract_has_none(approved):
    from app.schema import Architecture, Outcome
    base = with_arch()
    outs = [Outcome(metric="hidden", value="x", source_quote="q", unsourced=True),
            Outcome(metric="onboarding", value="12 to 3 days", source_quote="q")]
    arch = Architecture.model_validate({**base.architecture.model_dump(), "callouts": [{"text": "Onboarding in 3 days", "outcome": 1},
                                                                                       {"text": "Hidden claim", "outcome": 0}]})
    case = base.model_copy(update={"outcomes": outs, "architecture": arch})
    d = render.architecture_dict(case)
    assert d["callouts"] == [{"text": "Onboarding in 3 days", "outcome": 0}]  # numbered by place among the shown outcomes; unsourced dropped
    eng = render.architecture_dict(case.model_copy(update={"basis": "engagement"}))
    assert eng["callouts"] == [] and eng["components"]
    cid = approved(case.model_copy(update={"basis": "engagement"}))
    assert "Onboarding in 3 days" not in prev([cid]).text.split("<figure")[1].split("</figure>")[0] and "Hidden claim" not in gen([cid], "md").text


@pytest.mark.parametrize("fmt", ["md", "docx", "pptx", "case_study_docx"])
def test_produce_matches_generate_bytes(approved, fmt):
    from app.main import User, roles_for
    cid = approved()
    r = gen([cid], fmt)
    with db.connect() as c:
        body, mime, name = render.produce(c, User("u1", "U", frozenset(R), roles_for(R)), fmt, [cid], None, False, [])
    def same(a, b):  # an Office file stamps its creation time, which can differ by a second between the two builds
        if not a.startswith(b"PK"):
            return a == b
        za, zb = zipfile.ZipFile(BytesIO(a)), zipfile.ZipFile(BytesIO(b))
        return za.namelist() == zb.namelist() and all(za.read(n) == zb.read(n) for n in za.namelist() if n != "docProps/core.xml")
    assert r.status_code == 200 and same(body, r.content) and mime == r.headers["content-type"]
    assert r.headers["content-disposition"] == f'attachment; filename="{name}"'


def test_section_of_a_capability_case_has_the_note_and_no_results_or_lessons():  # #213
    case = ReferenceCase(title=Sourced[str](value="Offer", source_quote="q"), basis="capability",
                         outcomes=[Outcome(metric="m", value="3 days", source_quote="q")],
                         lessons_learned=Sourced[str](value="Start early", source_quote="q"))
    s = render.section(case, "a client")
    assert s["note"] == render.CAPABILITY_NOTE
    assert "Outcomes" not in [x["heading"] for x in s["lists"]] and "Lessons learned" not in [b["heading"] for b in s["blocks"]]


def capability():
    return full().model_copy(update={"basis": "capability"})


def test_capability_goes_in_its_own_section_after_the_cases(approved):  # #213
    cap, dlv = approved(capability()), approved(full())
    note = render.CAPABILITY_NOTE
    t = doc_text(gen([cap, dlv], "docx").content)
    assert t.count("Our capability") == 1 and t.count(note) == 1
    assert t.index("Our capability") > t.rindex("Outcomes") and "Outcomes" in t  # the delivered case keeps its results, before the heading
    assert t.index(note) > t.index("Our capability")
    m = gen([cap, dlv], "md").text
    assert m.count("# Our capability") == 1 and m.count(note) == 1 and m.index("Our capability") > m.index("Outcomes")
    assert "Our capability" not in gen([dlv], "md").text
    only = gen([cap], "md").text
    assert "Our capability" in only and "Outcomes" not in only and "12 to 3" not in only
    p = prev([cap, dlv]).text
    assert "Our capability" in p and note in p
    assert "Our capability" in prev([cap]).text  # a lone capability previews: no case-study build for it


@pytest.mark.parametrize("fmt", ["case_study_docx", "case_study_pdf", "pptx", "pdf_pptx"])
def test_capability_refuses_case_study_and_slide_formats(approved, fmt):
    r = gen([approved(capability())], fmt)
    assert r.status_code == 400 and r.json()["detail"] == "capability documents can't be case studies"
    assert gen([approved(full())], "md").status_code == 200


def test_section_never_outputs_capability_text_that_claims_a_result():  # #213
    from app.writeup import RESULT_FLAG
    case = ReferenceCase(title=Sourced[str](value="Offer", source_quote="q"), basis="capability", summary="It reduced cost.",
                         solution=Sourced[str](value="Delivered a ledger", source_quote="q"),
                         prose={"challenge": Prose(text="This resulted in speed.", by_reviewer=True, flags=[RESULT_FLAG])})
    s = render.section(case, "a client")
    assert s["summary"] == "" and not s["blocks"]
    assert render.section(case.model_copy(update={"basis": "delivered"}), "a client")["summary"] == "It reduced cost."


PALETTE = {"002535", "0A5A78", "FBF321", "FFFFFF", "000000", "808285", "F4F4F4"}


def test_docx_templates_use_aptos_and_brand_palette():
    for path in (render.TEMPLATE, render.TEMPLATE_CASE_STUDY, render.BRIEF_TEMPLATE):
        d = Document(path)
        for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption", "List Bullet", "Header", "Footer"):
            assert d.styles[name].font.name == "Aptos", (path.name, name)
        with zipfile.ZipFile(path) as z:
            styles = z.read("word/styles.xml").decode()
            assert "Aptos" in styles and "Calibri" not in styles and "Cambria" not in styles
            cols = {c.upper() for c in re.findall(r'(?:w:val|w:color|w:fill)="([0-9A-Fa-f]{6})"', styles)}
            assert cols and cols <= PALETTE, (path.name, cols - PALETTE)
            for part in ("word/document.xml", "word/header1.xml", "word/footer1.xml"):
                xml = z.read(part).decode()
                assert "DRAFT TEMPLATE" not in xml, (path.name, part)
                pcols = {c.upper() for c in re.findall(r'(?:w:color[^>]*w:val|w:color|w:fill)="([0-9A-Fa-f]{6})"', xml)}
                assert pcols <= PALETTE, (path.name, part, pcols - PALETTE)
            assert "<w:drawing>" in z.read("word/header1.xml").decode()
            footer = z.read("word/footer1.xml").decode()
            assert "Synechron \u00b7 Confidential" in footer and "PAGE" in footer



@pytest.mark.parametrize("fmt", ["pptx", "docx"])
def test_pdfs_embed_only_aptos(fmt):  # #282: a font the container lacks falls back silently (DejaVu)
    sec = render.section(full(), "a UK retailer")
    office = render.to_pptx([sec]) if fmt == "pptx" else render.to_docx([sec])
    fonts = set(re.findall(rb"/BaseFont\s*/[A-Z]{6}\+([A-Za-z0-9-]+)", render.to_pdf(office, "." + fmt)))
    assert fonts and all(f.startswith(b"Aptos") for f in fonts), fonts
