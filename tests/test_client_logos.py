import uuid
from io import BytesIO

import pytest
from PIL import Image

from app import clients, db, diagrams
from tests.test_auth import ADMIN, REV, USER, client, env  # noqa: F401


def img(fmt="PNG", size=(200, 100)):
    b = BytesIO()
    Image.new("RGB", size, "red").save(b, fmt)
    return b.getvalue()


@pytest.fixture
def cl(env):
    db.init()
    made = []

    def make(ref=True, allowed=True, png=True):
        with db.connect() as c:
            cid = c.execute("insert into clients(name,anonymised_label,referenceable,logo_allowed) values (%s,'a bank',%s,%s) returning id",
                            (uuid.uuid4().hex, ref, allowed)).fetchone()[0]
            if png:
                c.execute("insert into client_logos(client_id,png,width,height) values (%s,%s,200,100)", (cid, img()))
        made.append(cid)
        return cid
    yield make
    with db.connect() as c:
        for i in made:
            c.execute("delete from clients where id=%s", (i,))
            c.execute("delete from client_logo_changes where client_id=%s", (i,))
        c.execute("delete from jobs where kind='logo'")


@pytest.mark.parametrize("ref,allowed,decoded,shown", [
    (True, True, True, True), (True, True, False, False), (True, False, True, False), (False, True, True, False),
    (False, False, True, False), (True, False, False, False), (False, True, False, False), (False, False, False, False),
])
def test_logo_for_truth_table(cl, ref, allowed, decoded, shown):
    cid = cl(ref, allowed, png=False)
    with db.connect() as c:
        c.execute("insert into client_logos(client_id,png,raw) values (%s,%s,%s)", (cid, img() if decoded else None, None if decoded else b"x"))
        got = clients.logo_for(c, cid)
    assert (got is not None) is shown


def test_public_route_follows_the_rule(cl):
    ok, protected, nolicence = cl(), cl(ref=False), cl(allowed=False)
    u = client([USER])
    r = u.get(f"/logos/{ok}.png")
    assert r.status_code == 200 and r.headers["content-type"] == "image/png" and r.headers["x-content-type-options"] == "nosniff"
    assert u.get(f"/logos/{protected}.png").status_code == 404 and u.get(f"/logos/{nolicence}.png").status_code == 404
    assert client().get(f"/logos/{ok}.png").status_code == 401


def test_admin_preview_is_admin_only_and_ignores_the_rule(cl):
    cid = cl(allowed=False)
    assert client([ADMIN]).get(f"/admin/clients/{cid}/logo.png").status_code == 200
    assert client([REV]).get(f"/admin/clients/{cid}/logo.png").status_code == 403


def test_upload_queues_a_job_and_audits(cl):
    cid = cl(png=False)
    r = client([ADMIN]).post(f"/admin/clients/{cid}/logo", files={"file": ("l.png", img())}, follow_redirects=False)
    assert r.status_code == 303
    with db.connect() as c:
        assert c.execute("select raw is not null, png is null from client_logos where client_id=%s", (cid,)).fetchone() == (True, True)
        assert c.execute("select count(*) from jobs where kind='logo' and payload->>'client_id'=%s", (str(cid),)).fetchone()[0] == 1
        assert c.execute("select action, changed_by from client_logo_changes where client_id=%s", (cid,)).fetchall() == [("upload", "u1")]
    assert "Client logo changes" in client([ADMIN]).get("/admin/audit").text


def test_upload_refusals(cl):
    cid = cl(png=False)
    a = client([ADMIN])
    url = f"/admin/clients/{cid}/logo"
    assert a.post(url, files={"file": ("l.svg", b"<svg onload=x>")}).status_code == 400
    assert a.post(url, files={"file": ("l.png", b"GIF89a")}).status_code == 400
    assert a.post(url, files={"file": ("l.png", img() + b"0" * diagrams.LOGO_MAX_BYTES)}).status_code == 413
    assert a.post("/admin/clients/999999999/logo", files={"file": ("l.png", img())}).status_code == 404
    assert client([REV]).post(url, files={"file": ("l.png", img())}).status_code == 403
    with db.connect() as c:
        assert not c.execute("select 1 from client_logos where client_id=%s", (cid,)).fetchone()
        assert not c.execute("select 1 from client_logo_changes where client_id=%s", (cid,)).fetchone()


def test_remove_deletes_and_audits(cl):
    cid = cl()
    assert client([REV]).post(f"/admin/clients/{cid}/logo/remove").status_code == 403
    assert client([ADMIN]).post(f"/admin/clients/{cid}/logo/remove", follow_redirects=False).status_code == 303
    assert client([ADMIN]).post(f"/admin/clients/{cid}/logo/remove").status_code == 404
    with db.connect() as c:
        assert not c.execute("select 1 from client_logos where client_id=%s", (cid,)).fetchone()
        assert c.execute("select action from client_logo_changes where client_id=%s", (cid,)).fetchall() == [("remove",)]


def test_guard_caps_the_logo_body_before_parsing(cl):
    cid = cl(png=False)
    r = client([ADMIN]).post(f"/admin/clients/{cid}/logo", files={"file": ("l.png", img() + b"0" * 2 * diagrams.LOGO_MAX_BYTES)})
    assert r.status_code == 413 and "larger than" in r.text


def test_registry_page_shows_the_logo_state(cl):
    ready, nologo = cl(), cl(png=False)
    with db.connect() as c:
        c.execute("insert into client_logos(client_id,reason) values (%s,%s)", (nologo, diagrams.LOGO_TOO_SMALL))
    t = client([ADMIN]).get("/admin/clients").text
    assert "Logo ready" in t and f'src="/admin/clients/{ready}/logo.png"' in t and f'action="/admin/clients/{ready}/logo/remove"' in t
    assert "Rejected: " + diagrams.LOGO_TOO_SMALL in t and f'src="/admin/clients/{nologo}/logo.png"' not in t


# --- outputs (step 4)
import zipfile  # noqa: E402

from pptx import Presentation  # noqa: E402

from tests.test_render import R, full, gen, prev  # noqa: E402,F401
from tests.test_review import DOCS, make  # noqa: E402,F401
from tests.test_search import approved, link, reg  # noqa: E402,F401


@pytest.fixture
def case_with(approved, reg):
    def build(ref, allowed, png=True):
        name, client_id = reg("Zorp", "a UK retailer", referenceable=ref)
        with db.connect() as c:
            c.execute("update clients set logo_allowed=%s where id=%s", (allowed, client_id))
            if png:
                c.execute("insert into client_logos(client_id,png,width,height) values (%s,%s,200,100)", (client_id, img()))
        cid = approved(full())
        link(cid, client_id)
        return cid, client_id
    yield build
    with db.connect() as c:
        c.execute("delete from client_logos where client_id in (select id from clients where name like 'Zorp%')")


def drawings(content):
    return zipfile.ZipFile(BytesIO(content)).read("word/document.xml").decode().count("<w:drawing>")


def pictures(content):
    return sum(sh.shape_type == 13 for s in Presentation(BytesIO(content)).slides for sh in s.shapes)


def test_an_allowed_client_gets_its_logo_in_every_format(case_with):
    cid, client_id = case_with(True, True)
    assert drawings(gen([cid], "docx").content) == 1
    assert drawings(gen([cid], "case_study_docx").content) == 1
    assert pictures(gen([cid], "pptx").content) == 1
    assert f'src="/logos/{client_id}.png"' in prev([cid]).text
    assert f'src="/logos/{client_id}.png"' in client(R).post("/search", data={"bid_text": "onboarding"}).text
    assert b"logo" not in gen([cid], "md").content.lower()


@pytest.mark.parametrize("ref,allowed,png", [(True, False, True), (True, True, False), (False, True, True), (False, False, True)])
def test_no_logo_without_permission_a_logo_or_for_a_protected_client(case_with, ref, allowed, png):
    cid, client_id = case_with(ref, allowed, png)
    for fmt in ("docx", "case_study_docx"):
        body = gen([cid], fmt).content
        assert drawings(body) == 0
        t = "\n".join(p.text for p in __import__("docx").Document(BytesIO(body)).paragraphs)
        assert "{{" not in t and "{%" not in t
    first = __import__("docx").Document(BytesIO(gen([cid], "case_study_docx").content)).paragraphs[0]
    assert first.text.strip()  # the title, not an empty paragraph left by the logo condition
    assert pictures(gen([cid], "pptx").content) == 0
    assert f"/logos/{client_id}.png" not in prev([cid]).text
    assert f"/logos/{client_id}.png" not in client(R).post("/search", data={"bid_text": "onboarding"}).text


def test_a_name_protect_rewrites_gets_no_logo(approved, reg):
    reg("Zorp", "a retailer")  # protected; the referenceable client's name contains it, so outputs show its label instead
    _, client_id = reg("Zorp", "a logistics firm", referenceable=True, suffix=" Logistics")
    with db.connect() as c:
        c.execute("update clients set logo_allowed=true where id=%s", (client_id,))
        c.execute("insert into client_logos(client_id,png,width,height) values (%s,%s,200,100)", (client_id, img()))
    cid = approved(full())
    link(cid, client_id)
    for fmt in ("docx", "case_study_docx"):
        assert drawings(gen([cid], fmt).content) == 0
    assert pictures(gen([cid], "pptx").content) == 0
    assert f"/logos/{client_id}.png" not in prev([cid]).text
    assert f"/logos/{client_id}.png" not in client(R).post("/search", data={"bid_text": "onboarding"}).text
    with db.connect() as c:  # the display rule alone allows it: only the masked checks keep it out
        assert clients.logo_for(c, client_id)
        c.execute("delete from client_logos where client_id=%s", (client_id,))


def test_a_long_title_stops_before_the_logo_on_the_slide(case_with):
    cid, _ = case_with(True, True)
    with db.connect() as c:
        c.execute("update cases set data = jsonb_set(data, '{title,value}', to_jsonb(%s::text)) where id=%s", ("A very long title " * 12, cid))
    prs = Presentation(BytesIO(gen([cid], "pptx").content))
    slide = prs.slides[0]
    logo = next(sh for sh in slide.shapes if sh.shape_type == 13)
    idx = next(p.placeholder_format.idx for p in slide.slide_layout.placeholders if p.name == "Title")
    title = next(p for p in slide.placeholders if p.placeholder_format.idx == idx)
    assert title.left + title.width <= logo.left
