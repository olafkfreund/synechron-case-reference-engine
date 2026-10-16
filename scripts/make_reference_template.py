"""Write the branded Word templates (brand/reference.docx, brand/case_study.docx, brand/brief.docx).
All three use the Synechron house style (Aptos, the 7-colour palette, wordmark header, confidential footer)
and keep the same context: cases[] with title, group ("Our capability" on the first capability document, empty elsewhere), client, note (empty unless an engagement), summary, details[{label,value}],
blocks[{heading,text,paragraphs[str]}], lists[{heading,bullets[]}], diagrams[{image,caption}], arch_image (an InlineImage of the solution architecture; absent when there is none), footer (the case-study document only; empty elsewhere), logo (an InlineImage of the client's logo; absent unless the client is referenceable and logo-allowed); and industry (or None): heading, disclaimer,
groups[{title, lead, quotes[{quote, n, mono, rich}]}] (rich: a RichText of the note, only when mono), sources[{n, title, publisher, url, date}], note, show_head (false when only library briefs follow: skip the heading, disclaimer, groups and sources),
briefs[{title, disclaimer, takeaways[str], themes[{heading, text}], sources[{n, title, publisher, url, date}], note}] (library briefs, each with its own sources)
(statements[{text}] is the flat fallback). The optional second path gets the brief template (brand/brief.docx): brief with title, disclaimer,
takeaways[str], themes[{heading, text}], diagram (an InlineImage; falsy when there is none) and credit, sources[{n, title, publisher, url, date}], appendix[{n, quote, mono, rich}], note.
Usage: make_reference_template.py <out.docx> [<brief.docx>]
       make_reference_template.py --case-study <case_study.docx>   (one case per document: same context, cases has one item)"""
from io import BytesIO
from pathlib import Path
import sys
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

BRAND = Path(__file__).resolve().parent.parent / "brand"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# Strict 7-colour palette: #002535, #0A5A78, #FBF321, #FFFFFF, #000000, #808285, #F4F4F4
NAVY_HEX, TEAL_HEX, YELLOW_HEX = "002535", "0A5A78", "FBF321"
WHITE_HEX, BLACK_HEX, GREY_HEX, LIGHT_GREY_HEX = "FFFFFF", "000000", "808285", "F4F4F4"
NAVY = RGBColor(0x00, 0x25, 0x35)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x80, 0x82, 0x85)

KEEP_STYLES = frozenset({
    "Normal", "Header", "HeaderChar", "Footer", "FooterChar",
    "Heading1", "Heading1Char", "Heading2", "Heading2Char",
    "Title", "TitleChar", "Caption", "ListBullet", "ListParagraph",
    "DefaultParagraphFont", "TableNormal", "NoList",
})


def _set_rfonts(rPr, font="Aptos"):
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts xmlns:w="{W}"/>')
        rPr.insert(0, rFonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rFonts.attrib.pop(qn(f"w:{attr}"), None)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rFonts.set(qn(f"w:{attr}"), font)


def _style_rpr(style, size_pt, color_rgb, bold=False):
    style.font.name = "Aptos"
    style.font.size = Pt(size_pt)
    style.font.color.rgb = color_rgb
    style.font.bold = bold
    rPr = style.element.get_or_add_rPr()
    _set_rfonts(rPr, "Aptos")
    color = rPr.find(qn("w:color"))
    if color is not None:
        for attr in ("themeColor", "themeShade", "themeTint"):
            color.attrib.pop(qn(f"w:{attr}"), None)


def style_doc(d):
    styles_el = d.styles.element
    for s in list(styles_el.findall(qn("w:style"))):
        if s.get(qn("w:styleId")) not in KEEP_STYLES:
            styles_el.remove(s)
        else:
            for rsid in list(s.findall(qn("w:rsid"))):
                s.remove(rsid)

    doc_defaults = styles_el.find(qn("w:docDefaults"))
    if doc_defaults is not None:
        rPr_def = doc_defaults.find(f".//{qn('w:rPr')}")
        if rPr_def is not None:
            _set_rfonts(rPr_def, "Aptos")
            for tag in ("sz", "szCs"):
                el = rPr_def.find(qn(f"w:{tag}"))
                if el is not None:
                    el.set(qn("w:val"), "21")

    _style_rpr(d.styles["Normal"], 10.5, BLACK, bold=False)
    _style_rpr(d.styles["List Paragraph"], 10.5, BLACK, bold=False)
    _style_rpr(d.styles["List Bullet"], 10.5, BLACK, bold=False)
    _style_rpr(d.styles["Title"], 20, NAVY, bold=True)
    _style_rpr(d.styles["Heading 1"], 15, NAVY, bold=True)
    _style_rpr(d.styles["Heading 2"], 12.5, NAVY, bold=True)
    _style_rpr(d.styles["Caption"], 9, GREY, bold=False)
    _style_rpr(d.styles["Header"], 9, GREY, bold=False)
    _style_rpr(d.styles["Footer"], 8, GREY, bold=False)

    title_pPr = d.styles["Title"].element.find(qn("w:pPr"))
    if title_pPr is not None and (pBdr := title_pPr.find(qn("w:pBdr"))) is not None:
        title_pPr.remove(pBdr)

    for sid, (pt, col, bold) in {
        "TitleChar": (20, NAVY, True),
        "Heading1Char": (15, NAVY, True),
        "Heading2Char": (12.5, NAVY, True),
        "HeaderChar": (9, GREY, False),
        "FooterChar": (8, GREY, False),
        "DefaultParagraphFont": (10.5, BLACK, False),
    }.items():
        el = styles_el.find(f"{qn('w:style')}[@{qn('w:styleId')}='{sid}']")
        if el is not None:
            rPr = el.get_or_add_rPr()
            _set_rfonts(rPr, "Aptos")
            for child in list(rPr):
                if child.tag in (qn("w:color"), qn("w:sz"), qn("w:szCs"), qn("w:b"), qn("w:bCs")):
                    rPr.remove(child)
            if bold:
                rPr.append(parse_xml(f'<w:b xmlns:w="{W}"/>'))
                rPr.append(parse_xml(f'<w:bCs xmlns:w="{W}"/>'))
            rPr.append(parse_xml(f'<w:color xmlns:w="{W}" w:val="{col}"/>'))
            half_pt = str(int(pt * 2))
            rPr.append(parse_xml(f'<w:sz xmlns:w="{W}" w:val="{half_pt}"/>'))
            rPr.append(parse_xml(f'<w:szCs xmlns:w="{W}" w:val="{half_pt}"/>'))

    num_el = d.part.numbering_part.element
    for rFonts in num_el.findall(f".//{qn('w:rFonts')}"):
        rFonts.set(qn("w:ascii"), "Aptos")
        rFonts.set(qn("w:hAnsi"), "Aptos")
        rFonts.set(qn("w:cs"), "Aptos")
    for lvlText in num_el.findall(f".//{qn('w:lvlText')}"):
        if lvlText.get(qn("w:val")) == "\uf0b7":
            lvlText.set(qn("w:val"), "\u2022")

    hp = d.sections[0].header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.add_run().add_picture(str(BRAND / "wordmark-black.png"), width=Mm(30))

    footer(d.sections[0].footer.paragraphs[0])
    return d


def footer(fp):
    """The confidential line and the page number, grey 8 pt."""
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = fp.add_run("Synechron \u00b7 Confidential\t\t")  # the Footer style: centre tab, then the right one
    r.font.name, r.font.size, r.font.color.rgb = "Aptos", Pt(8), GREY
    for part in (
        '<w:fldChar w:fldCharType="begin"/>',
        '<w:instrText xml:space="preserve"> PAGE </w:instrText>',
        '<w:fldChar w:fldCharType="separate"/>',
        '<w:fldChar w:fldCharType="end"/>',
    ):
        fr = fp.add_run()
        fr.font.name, fr.font.size, fr.font.color.rgb = "Aptos", Pt(8), GREY
        _set_rfonts(fr._r.get_or_add_rPr(), "Aptos")
        fr._r.append(parse_xml(f'<w:r xmlns:w="{W}">{part}</w:r>')[0])


def yellow_rule(p):
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:pBdr xmlns:w="{W}"><w:bottom w:val="single" w:sz="6" w:space="4" w:color="{YELLOW_HEX}"/></w:pBdr>'
    ))
    return p


def _borderless_cell(table, fill_hex, pad_v=120, pad_h=160):
    table._tbl.tblPr.append(parse_xml(f'<w:tblW xmlns:w="{W}" w:w="5000" w:type="pct"/>'))
    table._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders xmlns:w="{W}">'
        '<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '</w:tblBorders>'
    ))
    cell = table.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd xmlns:w="{W}" w:val="clear" w:color="auto" w:fill="{fill_hex}"/>'))
    tcPr.append(parse_xml(
        f'<w:tcBorders xmlns:w="{W}">'
        '<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '</w:tcBorders>'
    ))
    tcPr.append(parse_xml(
        f'<w:tcMar xmlns:w="{W}">'
        f'<w:top w:w="{pad_v}" w:type="dxa"/>'
        f'<w:bottom w:w="{pad_v}" w:type="dxa"/>'
        f'<w:left w:w="{pad_h}" w:type="dxa"/>'
        f'<w:right w:w="{pad_h}" w:type="dxa"/>'
        '</w:tcMar>'
    ))
    return cell


def callout(d, text, italic=False, color=BLACK):
    cell = _borderless_cell(d.add_table(rows=1, cols=1), LIGHT_GREY_HEX)
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.italic = italic
    r.font.name, r.font.size, r.font.color.rgb = "Aptos", Pt(10.5), color
    _set_rfonts(r._r.get_or_add_rPr(), "Aptos")
    return p


def case_study_band(d):
    cell = _borderless_cell(d.add_table(rows=1, cols=1), NAVY_HEX, pad_v=200, pad_h=240)
    wp = cell.paragraphs[0]
    wp.add_run().add_picture(str(BRAND / "wordmark-white.png"), width=Mm(30))
    tp = cell.add_paragraph()
    tp.style = d.styles["Title"]
    r = tp.add_run("{{ c.title }}")
    r.bold = True
    r.font.name, r.font.size, r.font.color.rgb = "Aptos", Pt(20), WHITE
    _set_rfonts(r._r.get_or_add_rPr(), "Aptos")
    yellow_rule(tp)
    return tp


def note_p(d, text, italic=False):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.font.name, r.font.size, r.font.color.rgb = "Aptos", Pt(9.5), GREY
    _set_rfonts(r._r.get_or_add_rPr(), "Aptos")
    return p


THEME_XML = f"""<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="Synechron">
  <a:themeElements>
    <a:clrScheme name="Synechron">
      <a:dk1><a:srgbClr val="{BLACK_HEX}"/></a:dk1>
      <a:lt1><a:srgbClr val="{WHITE_HEX}"/></a:lt1>
      <a:dk2><a:srgbClr val="{NAVY_HEX}"/></a:dk2>
      <a:lt2><a:srgbClr val="{LIGHT_GREY_HEX}"/></a:lt2>
      <a:accent1><a:srgbClr val="{NAVY_HEX}"/></a:accent1>
      <a:accent2><a:srgbClr val="{TEAL_HEX}"/></a:accent2>
      <a:accent3><a:srgbClr val="{YELLOW_HEX}"/></a:accent3>
      <a:accent4><a:srgbClr val="{GREY_HEX}"/></a:accent4>
      <a:accent5><a:srgbClr val="{NAVY_HEX}"/></a:accent5>
      <a:accent6><a:srgbClr val="{TEAL_HEX}"/></a:accent6>
      <a:hlink><a:srgbClr val="{TEAL_HEX}"/></a:hlink>
      <a:folHlink><a:srgbClr val="{GREY_HEX}"/></a:folHlink>
    </a:clrScheme>
    <a:fontScheme name="Synechron">
      <a:majorFont><a:latin typeface="Aptos"/><a:ea typeface=""/><a:cs typeface=""/></a:majorFont>
      <a:minorFont><a:latin typeface="Aptos"/><a:ea typeface=""/><a:cs typeface=""/></a:minorFont>
    </a:fontScheme>
    <a:fmtScheme name="Synechron">
      <a:fillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:fillStyleLst>
      <a:lnStyleLst><a:ln w="9525"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln><a:ln w="9525"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln><a:ln w="9525"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln></a:lnStyleLst>
      <a:effectStyleLst><a:effectStyle><a:effectLst/></a:effectStyle><a:effectStyle><a:effectLst/></a:effectStyle><a:effectStyle><a:effectLst/></a:effectStyle></a:effectStyleLst>
      <a:bgFillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:bgFillStyleLst>
    </a:fmtScheme>
  </a:themeElements>
</a:theme>""".encode()

FONT_TABLE_XML = f"""<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<w:fonts xmlns:w="{W}">
  <w:font w:name="Aptos"><w:charset w:val="00"/><w:family w:val="swiss"/><w:pitch w:val="variable"/></w:font>
  <w:font w:name="Aptos Mono"><w:charset w:val="00"/><w:family w:val="modern"/><w:pitch w:val="fixed"/></w:font>
</w:fonts>""".encode()


def save_doc(d, path):
    buf = BytesIO()
    d.save(buf)
    with zipfile.ZipFile(BytesIO(buf.getvalue()), "r") as zin:
        styles_bytes = zin.read("word/styles.xml")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "word/stylesWithEffects.xml":
                    data = styles_bytes
                elif item.filename == "word/theme/theme1.xml":
                    data = THEME_XML
                elif item.filename == "word/fontTable.xml":
                    data = FONT_TABLE_XML
                zout.writestr(item, data)


def industry_section(d, page_break=True):
    d.add_paragraph("{%p if industry %}")
    d.add_paragraph("{%p if industry.show_head %}")
    d.add_heading("{{ industry.heading }}", 1).paragraph_format.page_break_before = page_break
    callout(d, "{{ industry.disclaimer }}", italic=True, color=GREY)
    d.add_paragraph("{%p for g in industry.groups %}")
    d.add_heading("{{ g.title }}", 2)
    d.add_paragraph().add_run("{{ g.lead }}").italic = True
    d.add_paragraph("{%p for q in g.quotes %}")
    d.add_paragraph("{%p if q.mono %}")
    d.add_paragraph("{{r q.rich }}", style="List Bullet")  # a technical note: Aptos Mono, line by line
    d.add_paragraph("{%p else %}")
    d.add_paragraph("\u201c{{ q.quote }}\u201d [{{ q.n }}]", style="List Bullet")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p if industry.sources %}")
    d.add_heading("Sources", 2)
    d.add_paragraph("{%p for s in industry.sources %}")
    d.add_paragraph("[{{ s.n }}] {{ s.title }}, {{ s.publisher }}, {{ s.url }}, retrieved {{ s.date }}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{%p if industry.note %}")
    note_p(d, "{{ industry.note }}")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{%p for b in industry.briefs %}")  # library briefs (#235): each with its own sources
    d.add_heading("{{ b.title }}", 1).paragraph_format.page_break_before = page_break
    callout(d, "{{ b.disclaimer }}", italic=True, color=GREY)
    d.add_paragraph("{%p for t in b.takeaways %}")
    d.add_paragraph("{{ t }}", style="List Bullet")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p for t in b.themes %}")
    d.add_heading("{{ t.heading }}", 2)
    d.add_paragraph("{{ t.text }}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p if b.sources %}")
    d.add_heading("Sources", 2)
    d.add_paragraph("{%p for s in b.sources %}")
    d.add_paragraph("[{{ s.n }}] {{ s.title }}, {{ s.publisher }}, {{ s.url }}, retrieved {{ s.date }}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{%p if b.note %}")
    note_p(d, "{{ b.note }}")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p endif %}")


if sys.argv[1] == "--case-study":
    d = style_doc(Document())
    d.sections[0].different_first_page_header_footer = True  # page 1: the blue band carries the wordmark, not the header
    footer(d.sections[0].first_page_footer.paragraphs[0])
    d.add_paragraph("{%p for c in cases %}")
    d.add_paragraph("{%p if c.logo %}")
    d.add_paragraph("{{ c.logo }}").alignment = WD_ALIGN_PARAGRAPH.RIGHT  # cover, top right
    d.add_paragraph("{%p endif %}")
    case_study_band(d)
    d.add_paragraph().add_run("{{ c.client }}").italic = True
    d.add_paragraph("{%p for x in c.details %}")
    d.add_paragraph("{{ x.label }}: {{ x.value }}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p if c.note %}")
    note_p(d, "{{ c.note }}")
    d.add_paragraph("{%p endif %}")
    d.add_heading("At a glance", 1)
    callout(d, "{{ c.summary }}")
    d.add_paragraph("{%p for l in c.lists %}")
    d.add_paragraph("{%p if l.heading == 'Outcomes' %}")
    d.add_heading("{{ l.heading }}", 2)
    d.add_paragraph("{%p for i in l.bullets %}")
    d.add_paragraph("{{ i }}", style="List Bullet")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p for b in c.blocks %}")  # Lessons learned is one of the blocks
    d.add_heading("{{ b.heading }}", 1)
    d.add_paragraph("{%p for p in b.paragraphs %}")
    d.add_paragraph("{{ p }}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p if c.arch_image %}")
    d.add_heading("Architecture", 1)
    d.add_paragraph("{{ c.arch_image }}")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{%p if c.diagrams %}")
    d.add_heading("Diagrams", 1)
    d.add_paragraph("{%p for d in c.diagrams %}")
    d.add_paragraph("{{ d.image }}")
    d.add_paragraph("{{ d.caption }}", style="Caption")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{%p for l in c.lists %}")
    d.add_paragraph("{%p if l.heading != 'Outcomes' %}")
    d.add_heading("{{ l.heading }}", 2)
    d.add_paragraph("{%p for i in l.bullets %}")
    d.add_paragraph("{{ i }}", style="List Bullet")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{%p endfor %}")
    note_p(d, "{{ c.footer }}")
    d.add_paragraph("{%p endfor %}")
    industry_section(d, page_break=False)
    save_doc(d, sys.argv[2])
    sys.exit()

d = style_doc(Document())

d.add_paragraph("{%p for c in cases %}")
d.add_paragraph("{%p if c.group %}")  # the first capability document opens "Our capability" (#213); the break sits on the first element
d.add_heading("{{ c.group }}", 1).paragraph_format.page_break_before = True
d.add_paragraph("{%p if c.logo %}")
d.add_paragraph("{{ c.logo }}").alignment = WD_ALIGN_PARAGRAPH.RIGHT
d.add_paragraph("{%p endif %}")
yellow_rule(d.add_heading("{{ c.title }}", 1))
d.add_paragraph("{%p elif c.logo %}")  # its own paragraph above the title, so the heading and Word's TOC stay as they were (#267)
logo = d.add_paragraph("{{ c.logo }}")
logo.alignment, logo.paragraph_format.page_break_before = WD_ALIGN_PARAGRAPH.RIGHT, True  # one case per page: the break moves up with it
yellow_rule(d.add_heading("{{ c.title }}", 1))
d.add_paragraph("{%p else %}")
h_title = yellow_rule(d.add_heading("{{ c.title }}", 1))
h_title.paragraph_format.page_break_before = True  # one case per page
d.add_paragraph("{%p endif %}")
d.add_paragraph().add_run("{{ c.client }}").italic = True
d.add_paragraph("{%p if c.note %}")
note_p(d, "{{ c.note }}")
d.add_paragraph("{%p endif %}")
callout(d, "{{ c.summary }}")
d.add_paragraph("{%p for x in c.details %}")
d.add_paragraph("{{ x.label }}: {{ x.value }}")
d.add_paragraph("{%p endfor %}")
d.add_paragraph("{%p for b in c.blocks %}")
d.add_heading("{{ b.heading }}", 2)
d.add_paragraph("{%p for p in b.paragraphs %}")
d.add_paragraph("{{ p }}")
d.add_paragraph("{%p endfor %}")
d.add_paragraph("{%p endfor %}")
d.add_paragraph("{%p for l in c.lists %}")
d.add_heading("{{ l.heading }}", 2)
d.add_paragraph("{%p for i in l.bullets %}")
d.add_paragraph("{{ i }}", style="List Bullet")
d.add_paragraph("{%p endfor %}")
d.add_paragraph("{%p endfor %}")
d.add_paragraph("{%p if c.arch_image %}")
d.add_heading("Architecture", 2)
d.add_paragraph("{{ c.arch_image }}")
d.add_paragraph("{%p endif %}")
d.add_paragraph("{%p for d in c.diagrams %}")
d.add_paragraph("{{ d.image }}")
d.add_paragraph("{{ d.caption }}", style="Caption")
d.add_paragraph("{%p endfor %}")
d.add_paragraph("{%p endfor %}")
industry_section(d)
save_doc(d, sys.argv[1])

if len(sys.argv) > 2:
    b = style_doc(Document())
    b.add_heading("{{ brief.title }}", 1)
    callout(b, "{{ brief.disclaimer }}", italic=True, color=GREY)
    b.add_paragraph("{%p if brief.takeaways %}")
    b.add_heading("Key takeaways", 2)
    b.add_paragraph("{%p for t in brief.takeaways %}")
    b.add_paragraph("{{ t }}", style="List Bullet")
    b.add_paragraph("{%p endfor %}")
    b.add_paragraph("{%p endif %}")
    b.add_paragraph("{%p for t in brief.themes %}")
    b.add_heading("{{ t.heading }}", 2)
    b.add_paragraph("{{ t.text }}")
    b.add_paragraph("{%p endfor %}")
    b.add_paragraph("{%p if brief.diagram %}")
    b.add_heading("Diagram", 2)
    b.add_paragraph("{{ brief.diagram }}")
    note_p(b, "{{ brief.credit }}", italic=True)
    b.add_paragraph("{%p endif %}")
    b.add_paragraph("{%p if brief.sources %}")
    b.add_heading("Sources", 2)
    b.add_paragraph("{%p for s in brief.sources %}")
    b.add_paragraph("[{{ s.n }}] {{ s.title }}, {{ s.publisher }}, {{ s.url }}, retrieved {{ s.date }}")
    b.add_paragraph("{%p endfor %}")
    b.add_paragraph("{%p if brief.images %}")
    b.add_heading("Images on the cited pages", 2)
    note_p(b, "Links only; each image is the copyright of its owner.")
    b.add_paragraph("{%p for i in brief.images %}")
    b.add_paragraph("{{ i.title }}, {{ i.publisher }}, {{ i.url }}")
    b.add_paragraph("{%p endfor %}")
    b.add_paragraph("{%p endif %}")
    b.add_heading("Quotes", 2)
    b.add_paragraph("{%p for q in brief.appendix %}")
    b.add_paragraph("{%p if q.mono %}")
    b.add_paragraph("{{r q.rich }}")  # a technical note: Aptos Mono
    b.add_paragraph("{%p else %}")
    b.add_paragraph("\u201c{{ q.quote }}\u201d [{{ q.n }}]")
    b.add_paragraph("{%p endif %}")
    b.add_paragraph("{%p endfor %}")
    b.add_paragraph("{%p endif %}")
    b.add_paragraph("{%p if brief.note %}")
    note_p(b, "{{ brief.note }}")
    b.add_paragraph("{%p endif %}")
    save_doc(b, sys.argv[2])

