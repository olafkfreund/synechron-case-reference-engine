"""Build brand/master.pptx from Synechron's official Light template (June 2026) (#282).

Usage: make_brand_master.py <Synechron_Light_Presentation_Template_June_2026.potx> <out.pptx>

The renderer fills layouts by name (app/render.py): "Reference case", "Industry context" and "Diagram"
with their exact placeholder names, plus "Cover" and "Closing". This keeps the template's master, theme
(Aptos), wordmark and artwork, drops its 24 sample slides and every layout we do not use, and:
- Cover and Closing are the template's yellow-device layouts on Synechron blue with the white wordmark:
  brand v2 cover option 2, mandatory for the closing slide (the Light template puts them on grey);
- the three content layouts are rebuilt on Light layouts whose only artwork is the wordmark.
Unreferenced media (the sample slides' photos and charts) are not saved: python-pptx writes only parts
reachable from the presentation.
"""
import io
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.util import Inches

NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
DEVICE = Path(__file__).resolve().parent.parent / "brand" / "cover-device.png"
BLUE, BLUE2, YELLOW, WHITE, BLACK, GREY = "002535", "0A5A78", "FBF321", "FFFFFF", "000000", "808285"

# name, x, y, w, h (inches on the 13.33 x 7.5 slide), size (pt), bold, colour, bulleted. Titles set their colour:
# the template's master title style is near white (its layouts override it).
# Nothing reaches into the top-right 50 x 25 mm: a permitted client logo goes there (render.to_pptx).
CASE = [("Title", .4, .4, 10.3, .7, 28, True, BLUE, False),
        ("Client", .4, 1.1, 10.3, .35, 14, False, BLUE2, False),
        ("Summary", .4, 1.55, 12.5, .9, 13, False, BLACK, False),
        ("Challenge", .4, 3.05, 6.0, 1.55, 11, False, BLACK, False),
        ("Solution", 6.85, 3.05, 6.0, 1.55, 11, False, BLACK, False),
        ("Outcomes", .4, 5.15, 6.0, 1.7, 11, False, BLACK, True),
        ("Technology", 6.85, 5.15, 6.0, 1.7, 11, False, BLACK, False)]
CASE_HEADINGS = [("Challenge", .4, 2.65), ("Solution", 6.85, 2.65), ("Outcomes", .4, 4.75), ("Technology", 6.85, 4.75)]
INDUSTRY = [("Title", .4, .4, 10.3, .7, 28, True, BLUE, False),
            ("Summary", .4, 1.15, 12.5, .6, 12, False, GREY, False),
            ("Statements", .4, 1.9, 12.5, 4.9, 11, False, BLACK, True)]
DIAGRAM = [("Title", .4, .4, 10.3, .7, 28, True, BLUE, False),
           ("Picture", .4, 1.25, 12.5, 5.0, 11, False, BLACK, False),
           ("Caption", .4, 6.35, 12.5, .45, 12, False, GREY, False)]
COVER_SUBTITLE = ("Subtitle", .41, 4.85, 5.0, .5, 16, False, WHITE, False)  # under the two-line title
CLOSING_KEEP = ("Graphic 8", "TextBox 2", "Picture 1", "© Copyright")  # wordmark, "Thank you", device, legal line


def e(v):
    return int(Inches(v))


def rpr(size, bold, colour):
    fill = f'<a:solidFill><a:srgbClr val="{colour}"/></a:solidFill>' if colour else ""
    sz = f'sz="{size * 100}" ' if size else ""
    return f'{sz}b="{int(bold)}">{fill}'


def shape(id_, name, x, y, w, h, nv, lvl1, body, autofit=True):
    fit = "<a:normAutofit/>" if autofit else ""  # shrink text that would overflow the box
    return parse_xml(
        f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{id_}" name="{name}"/><p:cNvSpPr/><p:nvPr>{nv}</p:nvPr></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{e(x)}" y="{e(y)}"/><a:ext cx="{e(w)}" cy="{e(h)}"/></a:xfrm></p:spPr>'
        f'<p:txBody><a:bodyPr wrap="square" lIns="0" rIns="0">{fit}</a:bodyPr><a:lstStyle>{lvl1}</a:lstStyle>{body}</p:txBody></p:sp>')


def keyline(id_, x, y):
    """The short yellow rule under a section heading: yellow as a controlled accent (brand v2)."""
    return parse_xml(
        f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{id_}" name="Keyline {id_}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{e(x)}" y="{e(y)}"/><a:ext cx="{e(.6)}" cy="{e(.04)}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="{YELLOW}"/></a:solidFill>'
        f'<a:ln><a:noFill/></a:ln></p:spPr></p:sp>')


def placeholder(i, spec):
    pname, x, y, w, h, size, bold, colour, bullets = spec
    ph = '<p:ph type="title"/>' if pname == "Title" else f'<p:ph type="body" sz="quarter" idx="{20 + i}"/>'
    marks = "" if bullets else "<a:buNone/>"
    indent = 'marL="228600" indent="-228600"' if bullets else 'marL="0" indent="0"'
    lvl1 = f'<a:lvl1pPr {indent} algn="l">{marks}<a:defRPr {rpr(size, bold, colour)}</a:defRPr></a:lvl1pPr>'
    return shape(100 + i, pname, x, y, w, h, ph, lvl1, f'<a:p><a:r><a:rPr lang="en-US"/><a:t>{pname}</a:t></a:r></a:p>')


def rebuild(layout, name, specs, headings=()):
    """Our placeholders only (unique idx), the layout's artwork (wordmark) kept."""
    layout.name = name
    tree = layout.shapes._spTree
    for ph in list(layout.placeholders):
        tree.remove(ph._element)
    for i, spec in enumerate(specs):
        tree.append(placeholder(i, spec))
    for i, (text, x, y) in enumerate(headings):
        tree.append(shape(150 + i, f"{text} heading", x, y, 6.0, .3, "", "",
                          f'<a:p><a:r><a:rPr lang="en-US" sz="1200" b="1"><a:solidFill><a:srgbClr val="{BLUE}"/></a:solidFill>'
                          f'</a:rPr><a:t>{text}</a:t></a:r></a:p>', autofit=False))
        tree.append(keyline(170 + i, x, y + .32))


def on_blue(layout, white_svg, device_png):
    """Cover option 2: solid Synechron blue, white wordmark, white text, the yellow device.

    The template draws the device as a cropped EMF, which the container's LibreOffice does not render:
    PDFs came out without it. brand/cover-device.png is the official Dark template's own blue cover,
    right half, rendered at 300 dpi (4000 x 2250 slide) by a full LibreOffice: the same artwork, already
    cropped to the picture box, so its srcRect goes."""
    bg = layout._element.cSld.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
    bg.getparent().replace(bg, parse_xml(
        f'<p:bg {NS}><p:bgPr><a:solidFill><a:srgbClr val="{BLUE}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>'))
    for blip in layout._element.iter("{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip"):
        if layout.part.rels[blip.get(EMBED)].target_part.partname.endswith("image3.svg"):  # the blue wordmark
            blip.set(EMBED, layout.part.relate_to(white_svg, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"))
    _, rid = layout.part.get_or_add_image_part(device_png)
    for blip in layout._element.iter(A + "blip"):
        if blip.get(EMBED) and layout.part.rels[blip.get(EMBED)].target_part.partname.endswith(".emf"):  # an SVG blip has none
            blip.set(EMBED, rid)
            src = blip.getparent().find(A + "srcRect")
            if src is not None:
                src.getparent().remove(src)
    for tag in ("rPr", "endParaRPr", "defRPr"):
        for r in layout._element.iter(A + tag):
            for old in r.findall(A + "solidFill") + r.findall(A + "schemeClr"):
                r.remove(old)
            ln = r.find(A + "ln")
            r.insert(0 if ln is None else list(r).index(ln) + 1, parse_xml(f'<a:solidFill {NS}><a:srgbClr val="{WHITE}"/></a:solidFill>'))


def open_template(path):
    """python-pptx refuses the .potx content type: repack it as a presentation, in memory."""
    src, buf = zipfile.ZipFile(path), io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
            out.writestr(item.filename, data)
    return Presentation(buf)


def main(src, dst):
    prs = open_template(src)
    for sld in list(prs.slides._sldIdLst):  # the template's sample slides must never ship
        prs.part.drop_rel(sld.rId)
        prs.slides._sldIdLst.remove(sld)
    by_name = {l.name: l for l in prs.slide_layouts}
    white_svg = next(p for p in prs.part.package.iter_parts() if str(p.partname).endswith("/image8.svg"))
    assert b'fill="#FFFFFF"' in white_svg.blob, "image8.svg is no longer the white wordmark"

    cover = by_name["Cover Slide with Logo"]
    cover.name = "Cover"
    title = next(ph for ph in cover.placeholders if ph.placeholder_format.type is not None)
    title.name = "Title"
    on_blue(cover, white_svg, str(DEVICE))
    cover.shapes._spTree.append(placeholder(10, COVER_SUBTITLE))

    closing = by_name["Thank You "]
    closing.name = "Closing"
    for sh in list(closing.shapes):  # the contact block and CTA pills: the engine has no contact to give
        if not sh.name.startswith(CLOSING_KEEP):
            sh._element.getparent().remove(sh._element)
    on_blue(closing, white_svg, str(DEVICE))

    rebuild(by_name["One column"], "Reference case", CASE, CASE_HEADINGS)
    rebuild(by_name["Table"], "Industry context", INDUSTRY)
    rebuild(by_name["Double Chart"], "Diagram", DIAGRAM)

    keep = {"Cover", "Closing", "Reference case", "Industry context", "Diagram"}
    for layout in list(prs.slide_layouts):
        if layout.name not in keep:
            prs.slide_layouts.remove(layout)
    p = prs.core_properties
    p.title, p.author, p.last_modified_by = "Reference cases", "Reference Engine", "Reference Engine"
    p.comments = p.subject = p.keywords = p.category = ""
    prs.save(dst)


if __name__ == "__main__":
    main(*sys.argv[1:3])
