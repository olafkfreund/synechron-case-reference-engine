"""Write the placeholder PowerPoint master. Marketing's real brand/master.pptx replaces it and needs:
a layout named "Reference case" whose placeholders are named exactly Title, Client, Summary,
Challenge, Solution, Outcomes, Technology (unique idx values; other placeholders are removed from
generated slides), and no slides of its own. Usage: make_reference_master.py <out.pptx>"""
import sys

from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.util import Inches

NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
GREY, RED = "595959", "C00000"
# name, x, y, w, h (inches on the 10 x 7.5 slide), font size (pt), bold, italic, colour, bulleted
PLACEHOLDERS = [("Title", .5, .3, 9, .8, 24, True, False, None, False),
                ("Client", .5, 1.1, 9, .35, 14, False, True, GREY, False),
                ("Summary", .5, 1.5, 9, 1.0, 14, False, False, None, False),
                ("Challenge", .5, 2.9, 4.4, 1.6, 11, False, False, None, False),
                ("Solution", 5.1, 2.9, 4.4, 1.6, 11, False, False, None, False),
                ("Outcomes", .5, 4.9, 4.4, 1.9, 11, False, False, None, True),
                ("Technology", 5.1, 4.9, 4.4, 1.9, 11, False, False, None, False)]
HEADINGS = [("Challenge", .5, 2.6), ("Solution", 5.1, 2.6), ("Outcomes", .5, 4.6), ("Technology", 5.1, 4.6)]


def e(v):
    return int(Inches(v))


def rpr(size, bold=False, italic=False, colour=None):
    fill = f'<a:solidFill><a:srgbClr val="{colour}"/></a:solidFill>' if colour else ""
    return f'sz="{size * 100}" b="{int(bold)}" i="{int(italic)}">{fill}'


def shape(id_, name, x, y, w, h, nv, lvl1, body):
    # normAutofit: shrink text that would overflow the box
    return parse_xml(
        f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{id_}" name="{name}"/><p:cNvSpPr/><p:nvPr>{nv}</p:nvPr></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{e(x)}" y="{e(y)}"/><a:ext cx="{e(w)}" cy="{e(h)}"/></a:xfrm></p:spPr>'
        f'<p:txBody><a:bodyPr wrap="square"><a:normAutofit/></a:bodyPr><a:lstStyle>{lvl1}</a:lstStyle>{body}</p:txBody></p:sp>')


prs = Presentation()
layout = prs.slide_layouts[5]
layout.name = "Reference case"
tree = layout.shapes._spTree
for ph in list(layout.placeholders):  # drop title/date/footer/number: only ours remain, idx unique
    tree.remove(ph._element)
for i, (name, x, y, w, h, size, bold, italic, colour, bullets) in enumerate(PLACEHOLDERS):
    ph = '<p:ph type="title"/>' if name == "Title" else f'<p:ph type="body" sz="quarter" idx="{20 + i}"/>'
    marks = "" if bullets else '<a:buNone/>'
    indent = 'marL="228600" indent="-228600"' if bullets else 'marL="0" indent="0"'
    lvl1 = f'<a:lvl1pPr {indent} algn="l">{marks}<a:defRPr {rpr(size, bold, italic, colour)}</a:defRPr></a:lvl1pPr>'
    tree.append(shape(100 + i, name, x, y, w, h, ph, lvl1, f'<a:p><a:r><a:rPr lang="en-US"/><a:t>{name}</a:t></a:r></a:p>'))
for i, (text, x, y) in enumerate(HEADINGS):
    tree.append(shape(150 + i, f"{text} heading", x, y, 4.4, .3, "", "",
                      f'<a:p><a:r><a:rPr lang="en-US" {rpr(12, True, colour=GREY)}</a:rPr><a:t>{text}</a:t></a:r></a:p>'))
tree.append(shape(200, "Draft notice", .5, 6.95, 9, .35, "", "",
                  f'<a:p><a:r><a:rPr lang="en-US" {rpr(12, True, colour=RED)}</a:rPr>'
                  f'<a:t>DRAFT TEMPLATE – replace with the marketing master</a:t></a:r></a:p>'))
p = prs.core_properties
p.title, p.author, p.last_modified_by = "Reference cases", "Reference Engine", "Reference Engine"
p.comments = p.subject = p.keywords = p.category = ""
prs.save(sys.argv[1])
