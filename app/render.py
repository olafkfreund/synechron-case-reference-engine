import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage, RichText
from fastapi import APIRouter, Depends, Form, HTTPException, Request
from fastapi.responses import Response
from jinja2 import Environment
from markupsafe import Markup
from pydantic import ValidationError
from pptx import Presentation

from app import anonymise, db, office
from app.archdiagram import layout, to_mermaid, to_pptx_shapes, to_png, to_svg
from app.main import User, require, wants_html
from app.review import PNG_HEADERS, VISIBLE, page
from app.writeup import PRICE_FLAG
from app.schema import ArchCallout, ArchComponent, ArchFlow, Architecture, ReferenceCase, fence, plain, vague_metric

router = APIRouter()
TEMPLATE = Path(__file__).resolve().parent.parent / "brand" / "reference.docx"
TEMPLATE_CASE_STUDY = TEMPLATE.with_name("case_study.docx")
MASTER = Path(__file__).resolve().parent.parent / "brand" / "master.pptx"
LAYOUT = "Reference case"
PLACEHOLDERS = ("Title", "Client", "Summary", "Challenge", "Solution", "Outcomes", "Technology")
WITHHELD = "output withheld: a protected client name is present"
CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # invalid in XML; also used to split names
SLIDE_TEXT = 300  # challenge/solution excerpt on a slide; the Word version has the full text
SLIDE_OUTCOMES = 5
ENGAGEMENT_NOTE = "Engagement reference: contracted scope; no outcomes are claimed."


class Withheld(Exception):
    pass


COMPANY = "Synechron"
INDUSTRY_HEADING = "Industry context (public sources)"
SOURCES_HEADING = "Best practice and public sources"
APPROACH_HEADING = "Approach and industry practice"  # a case study: after our delivery sections, under the disclaimer
INDUSTRY_DISCLAIMER = f"These are public statements by third parties, not {COMPANY} delivery evidence."
# the type filters claims; it never reaches an output (#23)
CLAIM_TYPES = frozenset({"out_of_the_box", "configuration", "industry_practice", "vendor_claim", "technical_note"})
NOTE = "technical_note"  # code or steps copied verbatim: shown monospace, never passed through plain()
SLIDE_ITEMS = 5


def _host(url) -> str:
    try:
        return urlparse(str(url)).hostname or ""
    except ValueError:  # e.g. "http://[x"
        return ""


def industry_section(claims: list[dict], clients, heading=INDUSTRY_HEADING, pages=()) -> dict | None:
    """Verbatim quotes grouped by source page, cited [n], with a numbered source list. apply() then blocked()
    PER CLAIM: one that still names a protected client is dropped (and counted), not the whole document."""
    page_titles = {p.get("url"): p.get("title") for p in pages}
    groups, order, omitted = {}, [], 0  # url -> group; order = urls by first appearance
    for c in claims:
        if c.get("type") not in CLAIM_TYPES:
            continue
        # only quote and title are anonymised: rewriting a publisher or URL would forge the citation
        note = c.get("type") == NOTE
        if note and not str(c.get("quote") or "").strip():  # nothing to show
            continue
        pl = (lambda t: t) if note else plain  # a note's code stays exact
        title = pl(str(c.get("title") or page_titles.get(c.get("url")) or "")) or _host(c.get("url", "")) or c.get("publisher", "")
        safe = {k: CTRL.sub("", str(v)) for k, v in
                {"quote": pl(str(c.get("quote", ""))), "title": title, "publisher": c.get("publisher", ""), "url": c.get("url", ""),
                 "retrieved_at": c.get("retrieved_at", "")}.items()}
        safe["quote"], safe["title"] = anonymise.apply(safe["quote"], clients), anonymise.apply(safe["title"], clients)
        safe["date"] = safe.pop("retrieved_at")[:10]
        if anonymise.blocked("\n".join(safe.values()), clients):  # incl. a client's own site as the source
            omitted += 1
            continue
        if safe["url"] not in groups:
            order.append(safe["url"])
            groups[safe["url"]] = {"title": safe["title"], "lead": f"{safe['publisher']} states:", "quotes": [],
                                   "source": {"n": len(order), "title": safe["title"], "publisher": safe["publisher"],
                                              "url": safe["url"], "date": safe["date"]}}
        g = groups[safe["url"]]
        g["quotes"].append({"quote": safe["quote"], "n": g["source"]["n"], "mono": note})
    if not groups and not omitted:
        return None
    gl = [groups[u] for u in order]
    sources = [g["source"] for g in gl]
    statements = [{"text": t, "mono": m} for g in gl for t, m in [(g["title"], False), *(
        (f"{q['quote']} [{q['n']}]", True) if q["mono"] else (f"\u201c{q['quote']}\u201d [{q['n']}]", False) for q in g["quotes"])]]
    if sources:
        statements.append({"text": "Sources: " + "; ".join(f"[{x['n']}] {x['title']}, {x['publisher']}, {x['url']}, retrieved {x['date']}"
                                                          for x in sources)})
    note = (f"{omitted} public statement{'s' if omitted != 1 else ''} omitted: "
            f"{'they' if omitted != 1 else 'it'} named a protected client") if omitted else ""
    return {"heading": heading, "disclaimer": INDUSTRY_DISCLAIMER,
            "groups": [{k: g[k] for k in ("title", "lead", "quotes")} for g in gl], "sources": sources,
            "statements": statements, "note": note}


def brief_section(body: dict, quotes: list[dict], clients, images=()) -> dict:
    """A brief for output: apply() on every string, then blocked() PER SENTENCE and per quote. What still names a
    protected client is dropped and counted, as in industry_section. Numbers [n] are the quote numbers the brief cites."""
    from app import brief  # brief imports render's neighbours; import here to keep render importable on its own
    clean = lambda t: anonymise.apply(CTRL.sub("", str(t or "")), clients)  # noqa: E731
    omitted, dead, sources, appendix = 0, set(), [], []
    for n, q in enumerate(quotes[:brief.MAX_QUOTES], 1):
        note = q.get("type") == NOTE
        item = {"n": n, "quote": clean(q.get("quote")) if note else clean(plain(str(q.get("quote") or ""))),
                "title": clean(plain(str(q.get("title") or "")) or _host(q.get("url")) or q.get("publisher")),
                "publisher": CTRL.sub("", str(q.get("publisher") or "")), "url": CTRL.sub("", str(q.get("url") or "")),
                "date": CTRL.sub("", str(q.get("retrieved_at") or ""))[:10]}
        if anonymise.blocked("\n".join(str(v) for v in item.values()), clients):
            dead.add(n)
            omitted += 1
            continue
        sources.append({k: item[k] for k in ("n", "title", "publisher", "url", "date")})
        appendix.append({"n": n, "quote": item["quote"], "mono": note})

    def prose(text):  # the sentences that survive, joined again
        nonlocal omitted
        kept = []
        for sent in brief.sentences(str(text or "")):
            c = clean(sent)
            if anonymise.blocked(c, clients) or dead & set(brief.cites(c)):
                omitted += 1
            else:
                kept.append(c)
        return " ".join(kept)

    links = []  # links only, never pixels; one naming a protected client is left out
    for i in images:
        link = {"title": clean(plain(str(i.get("title") or "")) or i.get("publisher")), "publisher": CTRL.sub("", str(i.get("publisher") or "")),
                "url": CTRL.sub("", str(i.get("url") or ""))}
        if not anonymise.blocked("\n".join(link.values()), clients):
            links.append(link)
    title = clean(body.get("title"))
    if not title or anonymise.blocked(title, clients):
        title = "Best practice brief"
    themes = []
    for t in body.get("themes", []):
        heading = clean(t.get("heading"))
        if anonymise.blocked(heading, clients):
            omitted += 1
        elif text := prose(t.get("text")):
            themes.append({"heading": heading, "text": text})
    out = {"title": title, "disclaimer": INDUSTRY_DISCLAIMER, "takeaways": [x for x in map(prose, body.get("takeaways", [])) if x],
           "themes": themes, "sources": sources, "images": links, "appendix": appendix,
           "note": (f"{omitted} statement{'s' if omitted != 1 else ''} omitted: "
                    f"{'they' if omitted != 1 else 'it'} named a protected client") if omitted else ""}
    if anonymise.blocked("\n".join(_strings(out)), clients):
        raise Withheld(WITHHELD)
    return out


BRIEF_TEMPLATE = TEMPLATE.with_name("brief.docx")


def to_brief_docx(sec: dict) -> bytes:
    tpl = DocxTemplate(BRIEF_TEMPLATE)
    sec = {**sec, "appendix": [{**q, "rich": _mono(q)} if q.get("mono") else q for q in sec["appendix"]]}
    tpl.render({"brief": sec}, autoescape=True)
    props = tpl.docx.core_properties  # never ship the template's own metadata
    props.title, props.author, props.last_modified_by = "Best practice brief", "Reference Engine", "Reference Engine"
    props.comments = props.subject = props.keywords = props.category = ""
    props.created = props.modified = datetime.now(timezone.utc).replace(tzinfo=None)
    buf = BytesIO()
    tpl.save(buf)
    return buf.getvalue()


SLIDE_QUOTE = 140


def slide_statements(ind: dict) -> list[str]:
    """Every slide line starts with its source and keeps the verbatim quote next to it, so a reader can
    check it. The omitted-statements note (if any) comes last, as in
    Word and Markdown (#123)."""
    def line(g):  # a model claim if the page has one; a note shows its first line only
        q = next((q for q in g["quotes"] if not q["mono"]), g["quotes"][0])
        if q["mono"]:
            first = next(x.strip() for x in q["quote"].splitlines() if x.strip())  # industry_section drops blank notes
            return f"{g['lead']} {excerpt(first, SLIDE_QUOTE)} (see source [{q['n']}])"
        return f"{g['lead']} \u201c{excerpt(q['quote'], SLIDE_QUOTE)}\u201d [{q['n']}]"
    lines = [line(g) for g in ind["groups"]]
    # a library brief is one line (its title and takeaways, no source numbers); briefs come first so they are not cut by the quotes
    from app.brief import CITE  # a brief's [n] points at its own sources, which the slide does not list
    nocite = lambda t: re.sub(r"\s+([.,;:])", r"\1", " ".join(CITE.sub("", t).split()))  # noqa: E731
    briefs = [excerpt(nocite(f"{b['title']}: {' '.join(b['takeaways'])}").strip(": "), SLIDE_QUOTE) for b in ind.get("briefs", [])]
    return (briefs + lines)[:SLIDE_ITEMS] + ([ind["note"]] if ind["note"] else [])


def section(case: ReferenceCase, client: str) -> dict:
    """Only the approved record: no quotes, organisations, needs_attention or AI-tailored text."""
    v = lambda s: None if s.unsourced else s.value  # noqa: E731
    period = " – ".join(p for p in (case.period.start, case.period.end) if p) if not case.period.unsourced else ""
    details = [("Industry", v(case.industry)), ("Region", v(case.region)),
               ("Engagement type", v(case.engagement_type)),
               ("Duration", f"{v(case.duration_months)} months" if v(case.duration_months) else None),
               ("Team size", v(case.team_size)), ("Period", period)]
    def prose(f, head):  # approved prose replaces the one-line value only when it is unflagged or the reviewer's own
        p = case.prose.get(f)
        if p and p.text.strip() and (not p.flags or p.by_reviewer):
            return p.text, [x.strip() for x in p.text.split("\n\n") if x.strip()]
        return head, [head]  # the headline value goes out unchanged: one paragraph, no split

    blocks = [("Challenge", prose("challenge", v(case.challenge))), ("Solution", prose("solution", v(case.solution))),
              ("Implementation approach", prose("implementation", v(case.implementation))),
              ("Configuration and integration", prose("configuration", v(case.configuration))),
              ("Lessons learned", (None, []) if case.basis == "engagement" else prose("lessons_learned", v(case.lessons_learned)))]
    lists = [("Capabilities", [v(c) for c in case.capabilities]),
             ("Technology", [v(t) for t in case.tech_stack]),
             ("Technical references", [v(t) for t in case.tech_references]),
             ("Outcomes", [(f"{o.metric}: {o.value}" if not vague_metric(o) else o.value) for o in case.outcomes if not o.unsourced])]
    engagement = case.basis == "engagement"
    if engagement:  # contracted scope claims no results, whatever the record holds
        lists = [x for x in lists if x[0] != "Outcomes"]
    return dict(
        title=v(case.title) or "", client=client, summary=case.summary,
        basis=case.basis, note=ENGAGEMENT_NOTE if engagement else "",
        details=[dict(label=a, value=str(b)) for a, b in details if b not in (None, "")],
        blocks=[dict(heading=a, text=t, paragraphs=ps) for a, (t, ps) in blocks if t],
        lists=[dict(heading=a, bullets=[str(i) for i in b if i]) for a, b in lists if any(b)],
        diagrams=[dict(id=d.id, caption=f"{d.caption} {d.credit}" if d.credit else d.caption, width=d.width, height=d.height) for d in case.diagrams if d.status == "ready"],
        architecture=architecture_dict(case))


def architecture_dict(case: ReferenceCase) -> dict | None:
    """The architecture as plain text and keys: only when it has components and no unresolved flag (#246).
    It goes through protect() with the rest, and every output draws from that protected dict, never from the case."""
    a = case.architecture
    if not a or not a.components or any(PRICE_FLAG in f for f in a.flags) or (a.flags and not a.by_reviewer):
        return None  # a reviewer's confirmation lets a flag pass, never a price
    shown = [i for i, o in enumerate(case.outcomes) if not o.unsourced]  # the Outcomes list as section() shows it
    # a contract claims no results (as in section()); a callout is numbered by its outcome's place in the shown list
    callouts = [] if case.basis == "engagement" else [{"text": c.text, "outcome": shown.index(c.outcome)} for c in a.callouts if c.outcome in shown]
    return {"title": a.title, "components": [c.model_dump(exclude={"cite"}) for c in a.components],
            "flows": [f.model_dump(exclude={"cite"}) for f in a.flows], "callouts": callouts, "footer": ""}


def architecture_of(sec: dict):
    """(Architecture, Scene) from a protected section; None when it has none, or protect() rewrote a key (a client called "core")."""
    a = sec.get("architecture")
    if not a:
        return None
    try:  # the structure is checked with placeholder text (lengths are input limits: an expanded client label may run past them)
        skel = Architecture.model_validate({
            "title": "t", "components": [{**c, "name": "n", "role": ""} for c in a["components"]],
            "flows": [{**f, "label": ""} for f in a["flows"]], "callouts": [{"text": "t", "outcome": c["outcome"]} for c in a["callouts"]]})
        if (len(skel.components), len(skel.flows), len(skel.callouts)) != (len(a["components"]), len(a["flows"]), len(a["callouts"])):
            return None  # an id was rewritten, or a cap cut something: not the architecture that was approved
    except (ValidationError, KeyError, TypeError):
        return None
    arch = Architecture.model_construct(
        title=a["title"], by_reviewer=False, flags=[],
        components=[ArchComponent.model_construct(id=s.id, name=c["name"], role=c["role"], zone=s.zone, kind=s.kind, cite="")
                    for s, c in zip(skel.components, a["components"])],
        flows=[ArchFlow.model_construct(src=s.src, dst=s.dst, label=f["label"], mode=s.mode, cite="") for s, f in zip(skel.flows, a["flows"])],
        callouts=[ArchCallout.model_construct(text=c["text"], outcome=s.outcome) for s, c in zip(skel.callouts, a["callouts"])])
    return arch, layout(arch, a.get("footer", ""))


ARCH_KEYS = frozenset({"id", "zone", "kind", "mode", "src", "dst"})  # the architecture's structure: never text, never rewritten


def _walk(o, f, keep=frozenset()):
    if isinstance(o, str):
        return f(o)
    if isinstance(o, dict):
        return {k: x if k in keep else _walk(x, f, ARCH_KEYS if k == "architecture" else keep) for k, x in o.items()}
    return [_walk(x, f, keep) for x in o] if isinstance(o, list) else o


def _strings(o, keep=frozenset()):
    if isinstance(o, str):
        yield o
    elif isinstance(o, dict):
        for k, x in o.items():
            if k not in keep:
                yield from _strings(x, ARCH_KEYS if k == "architecture" else keep)
    elif isinstance(o, list):
        for x in o:
            yield from _strings(x, keep)


def protect(sections: list[dict], clients) -> list[dict]:
    """apply() on every string, then fail closed: blocked() must find nothing in the whole output."""
    # strip control characters FIRST: stripped after the check, "Zo\x01rp" would turn back into "Zorp"
    out = _walk(sections, lambda s: anonymise.apply(CTRL.sub("", s), clients))
    if anonymise.blocked("\n".join(_strings(out)), clients):
        raise Withheld(WITHHELD)  # never name the client
    return out


def _mono(q: dict) -> RichText:
    rt = RichText()
    rt.add(q["quote"], font="Consolas")
    rt.add(f" [{q['n']}]")
    return rt


def to_docx(sections: list[dict], industry: dict | None = None, images: dict[int, bytes] = {}, template=None) -> bytes:
    tpl = DocxTemplate(template or TEMPLATE)
    # images are added after protect(): an InlineImage is not a string; one missing from `images` is skipped
    sections = [{**s, "diagrams": [{**d, "image": InlineImage(tpl, BytesIO(images[d["id"]]), width=Mm(min(160, d["width"] * 25.4 / 96)))}
                                   for d in s["diagrams"] if d["id"] in images],
                 "arch_image": InlineImage(tpl, BytesIO(to_png(a[1])), width=Mm(160)) if (a := architecture_of(s)) else None} for s in sections]
    if industry:  # a note is Consolas, line by line, in a run of its own (the template checks q.mono)
        industry = {**industry, "groups": [{**g, "quotes": [{**q, "rich": _mono(q)} if q.get("mono") else q for q in g["quotes"]]}
                                           for g in industry["groups"]]}
    tpl.render({"cases": sections, "industry": industry}, autoescape=True)  # autoescape: & and < must not break the XML
    props = tpl.docx.core_properties  # never ship the template's own metadata
    props.title, props.author, props.last_modified_by = "Reference cases", "Reference Engine", "Reference Engine"
    props.comments = props.subject = props.keywords = props.category = ""
    props.created = props.modified = datetime.now(timezone.utc).replace(tzinfo=None)  # not the template's dates
    buf = BytesIO()
    tpl.save(buf)
    return buf.getvalue()


def to_case_study_docx(sections: list[dict], industry: dict | None = None, images: dict[int, bytes] = {}) -> bytes:
    return to_docx(sections, industry, images, TEMPLATE_CASE_STUDY)


def excerpt(s: str, limit: int = SLIDE_TEXT) -> str:
    return s if len(s) <= limit else s[:limit].rsplit(" ", 1)[0].rstrip(",;:.") + " …"


INDUSTRY_LAYOUT = "Industry context"
INDUSTRY_PLACEHOLDERS = ("Title", "Summary", "Statements")


DIAGRAM_LAYOUT = "Diagram"
DIAGRAM_PLACEHOLDERS = ("Title", "Picture", "Caption")


def _layout(prs, name: str, required) -> tuple | None:
    """(layout, idx -> placeholder name) after validation; None when the master has no such layout.
    Slide placeholders get generic names when cloned, so the layout's names are the source of truth:
    its placeholder idx is only the join key, never a position we rely on."""
    layout = next((l for l in prs.slide_layouts if l.name == name), None)
    if layout is None:
        return None
    names = {ph.placeholder_format.idx: ph.name for ph in layout.placeholders}
    if len(names) != len(list(layout.placeholders)):
        raise ValueError(f"layout {name!r} has duplicate placeholder idx values")
    if missing := [n for n in required if n not in names.values()]:
        raise ValueError(f"layout {name!r} is missing placeholders: {', '.join(missing)}")
    return layout, names


def _fill(prs, layout, names, content: dict, required) -> None:
    slide = prs.slides.add_slide(layout)
    filled = set()
    for ph in list(slide.placeholders):
        name = names.get(ph.placeholder_format.idx)
        if name not in content:  # logo/subtitle/footer etc.: removed, so no "Click to add text"
            ph._element.getparent().remove(ph._element)
            continue
        first, *rest = content[name] or [""]  # no content: empty text
        ph.text_frame.text = first
        for item in rest:
            ph.text_frame.add_paragraph().text = item
        filled.add(name)
    if filled != set(required):
        raise ValueError(f"layout {layout.name!r} did not yield placeholders: {', '.join(sorted(set(required) - filled))}")


def _diagram_slide(prs, layout, names, title: str, d: dict, png: bytes) -> None:
    _fill(prs, layout, names, {"Title": [title], "Picture": [""], "Caption": [d["caption"]]}, DIAGRAM_PLACEHOLDERS)
    slide = prs.slides[-1]
    idx = next(i for i, n in names.items() if n == "Picture")
    box = next(ph for ph in slide.placeholders if ph.placeholder_format.idx == idx)
    left, top, w, h = box.left, box.top, box.width, box.height
    box._element.getparent().remove(box._element)
    # add_picture at the box, not insert_picture: a picture placeholder crops the image to fill it
    scale = min(w / d["width"], h / d["height"])
    pw, ph_ = int(d["width"] * scale), int(d["height"] * scale)
    slide.shapes.add_picture(BytesIO(png), left + (w - pw) // 2, top + (h - ph_) // 2, width=pw, height=ph_)


def _architecture_slide(prs, diagram_layout, title: str, scene) -> None:
    """One slide of native shapes. The Diagram layout's picture box gives the area; without that layout, a blank slide with a title."""
    from pptx.util import Emu, Inches, Pt
    heading = f"{title}: solution architecture"
    if diagram_layout:
        layout_, names = diagram_layout
        _fill(prs, layout_, names, {"Title": [heading], "Picture": [""], "Caption": [""]}, DIAGRAM_PLACEHOLDERS)
        slide = prs.slides[-1]
        idx = next(i for i, n in names.items() if n == "Picture")
        box = next(ph for ph in slide.placeholders if ph.placeholder_format.idx == idx)
        left, top, w, h = box.left, box.top, box.width, box.height
        box._element.getparent().remove(box._element)
    else:
        slide = prs.slides.add_slide(min(prs.slide_layouts, key=lambda l: len(l.placeholders)))
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        tb = slide.shapes.add_textbox(Inches(.5), Inches(.3), prs.slide_width - Inches(1), Inches(.7))
        tb.text_frame.text = heading
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(26)
        left, top, w, h = Inches(.5), Inches(1.2), prs.slide_width - Inches(1), prs.slide_height - Inches(1.7)
    to_pptx_shapes(slide, scene, left, top, w, h)


def to_pptx(sections: list[dict], industry: dict | None = None, images: dict[int, bytes] = {}) -> bytes:
    prs = Presentation(MASTER)
    if len(prs.slides):  # a master's sample slides (and their text) must never ship
        raise ValueError(f"{MASTER.name} must contain no slides, only layouts")
    case_layout = _layout(prs, LAYOUT, PLACEHOLDERS)
    if case_layout is None:
        raise ValueError(f"{MASTER.name} has no slide layout named {LAYOUT!r}")
    diagram_layout = None
    if any(d["id"] in images for c in sections for d in c["diagrams"]):  # a master without it fails only now
        diagram_layout = _layout(prs, DIAGRAM_LAYOUT, DIAGRAM_PLACEHOLDERS)
        if diagram_layout is None:
            raise ValueError(f"{MASTER.name} has no slide layout named {DIAGRAM_LAYOUT!r}")
    for c in sections:
        block = lambda h: next((b["text"] for b in c["blocks"] if b["heading"] == h), "")  # noqa: E731
        bullets = lambda h: next((l["bullets"] for l in c["lists"] if l["heading"] == h), [])  # noqa: E731
        _fill(prs, *case_layout, {
            "Title": [c["title"]], "Client": [c["client"]], "Summary": [c["summary"]],
            "Challenge": [excerpt(block("Challenge"))], "Solution": [excerpt(block("Solution"))],
            "Outcomes": [c["note"]] if c["note"] else bullets("Outcomes")[:SLIDE_OUTCOMES], "Technology": [", ".join(bullets("Technology"))]},
            PLACEHOLDERS)
        if a := architecture_of(c):
            _architecture_slide(prs, _layout(prs, DIAGRAM_LAYOUT, DIAGRAM_PLACEHOLDERS), c["title"], a[1])
        for d in c["diagrams"]:
            if d["id"] in images:
                _diagram_slide(prs, *diagram_layout, c["title"], d, images[d["id"]])
    if industry:
        if own := _layout(prs, INDUSTRY_LAYOUT, INDUSTRY_PLACEHOLDERS):
            _fill(prs, *own, {"Title": [industry["heading"]], "Summary": [industry["disclaimer"]],
                              "Statements": slide_statements(industry)}, INDUSTRY_PLACEHOLDERS)
        else:  # fallback: the case layout, its case-only boxes left empty
            refs = [f"[{x['n']}] {x['publisher']}" for x in industry["sources"]]
            _fill(prs, *case_layout, {
                "Title": [industry["heading"]], "Client": [""], "Summary": [industry["disclaimer"]],
                "Challenge": [""], "Solution": [""], "Outcomes": slide_statements(industry),
                "Technology": ["Sources: " + ", ".join(refs)] if refs else [""]}, PLACEHOLDERS)
    props = prs.core_properties
    props.title, props.author, props.last_modified_by = "Reference cases", "Reference Engine", "Reference Engine"
    props.comments = props.subject = props.keywords = props.category = ""
    props.created = props.modified = datetime.now(timezone.utc).replace(tzinfo=None)  # not the template's dates
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


PDF_TIMEOUT = 60  # seconds
PDF_WAIT = 10


class PdfError(RuntimeError):
    """Conversion failed. Messages never carry case content."""


def to_pdf(data: bytes, suffix: str) -> bytes:
    """Convert docx/pptx bytes with headless LibreOffice (see office.convert)."""
    name = f"in{suffix}"
    try:
        out = office.convert({name: data}, "pdf", PDF_TIMEOUT, PDF_WAIT)
    except RuntimeError as e:
        raise PdfError({"busy": "PDF service busy; try again shortly",
                        "timeout": "PDF conversion timed out"}.get(str(e), f"PDF conversion failed ({e})")) from None
    if name not in out:
        raise PdfError("PDF conversion failed (exit 0)")
    return out[name]


def md(s: str) -> str:
    """Markdown escaping: newlines become spaces; backslash, backtick, * _ [ ] < > # | ~ & are
    escaped; a leading list or numbered marker is escaped, so a value placed at the start of a
    paragraph or bullet cannot become a heading, list, table or strike-through."""
    s = re.sub(r"([\\`*_\[\]<>#|~&])", r"\\\1", " ".join(str(s).split()))  # &: no entity decodes to a name (#126)
    return re.sub(r"^(\d+)([.)])|^([-+=])", lambda m: f"{m[1]}\\{m[2]}" if m[1] else f"\\{m[3]}", s)


def link(url: str) -> str:
    """An autolink only for a plain http(s) URL: a '>' or space would end it and let the rest become Markdown."""
    return f"<{url}>" if re.fullmatch(r"https?://[^\s<>]+", url) else md(url)


_env = Environment(autoescape=False, trim_blocks=True, lstrip_blocks=True)  # Markdown, not HTML
_env.filters["md"] = md
_env.filters["link"] = link
_env.filters["fence"] = fence
MD = _env.from_string("""\
{% for c in cases %}
# {{ c.title | md }}

*{{ c.client | md }}*

{% if c.note %}
{{ c.note | md }}

{% endif %}
{{ c.summary | md }}

{% for x in c.details %}
- **{{ x.label }}:** {{ x.value | md }}
{% endfor %}

{% for b in c.blocks %}
## {{ b.heading }}

{% for p in b.paragraphs %}{{ p | md }}

{% endfor %}{% endfor %}
{% for l in c.lists %}
## {{ l.heading }}

{% for i in l.bullets %}
- {{ i | md }}
{% endfor %}

{% endfor %}
{% if c.mermaid %}
## Architecture

```mermaid
{{ c.mermaid }}```

{% endif %}
{% if c.diagrams %}
## Diagrams

{% for d in c.diagrams %}
- {{ d.caption | md }} (in the Word, PowerPoint and PDF versions)
{% endfor %}

{% endif %}
{% if not loop.last %}
---

{% endif %}
{% endfor %}
{% if industry %}
{% if cases %}
---

{% endif %}
{% if industry.show_head %}
# {{ industry.heading | md }}

*{{ industry.disclaimer | md }}*

{% for g in industry.groups %}
## {{ g.title | md }}

*{{ g.lead | md }}*

{% for q in g.quotes %}
{% if q.mono %}
{{ q.quote | fence }}

Source [{{ q.n }}]

{% else %}
> \u201c{{ q.quote | md }}\u201d [{{ q.n }}]

{% endif %}
{% endfor %}
{% endfor %}
{% if industry.sources %}
## Sources

{% for s in industry.sources %}
{{ s.n }}. {{ s.title | md }}, {{ s.publisher | md }}, {{ s.url | link }}, retrieved {{ s.date | md }}
{% endfor %}
{% endif %}
{% endif %}
{% if industry.note %}

{{ industry.note | md }}
{% endif %}
{% for b in industry.briefs %}

# {{ b.title | md }}

*{{ b.disclaimer | md }}*

{% if b.takeaways %}
## Key takeaways

{% for t in b.takeaways %}
- {{ t | md }}
{% endfor %}

{% endif %}
{% for t in b.themes %}
## {{ t.heading | md }}

{{ t.text | md }}

{% endfor %}
{% if b.sources %}
## Sources

{% for s in b.sources %}
{{ s.n }}. {{ s.title | md }}, {{ s.publisher | md }}, {{ s.url | link }}, retrieved {{ s.date | md }}
{% endfor %}
{% endif %}
{% if b.note %}

{{ b.note | md }}
{% endif %}
{% endfor %}
{% endif %}
""")

BRIEF_MD = _env.from_string("""\
# {{ b.title | md }}

*{{ b.disclaimer | md }}*

{% if b.takeaways %}
## Key takeaways

{% for t in b.takeaways %}
- {{ t | md }}
{% endfor %}

{% endif %}
{% for t in b.themes %}
## {{ t.heading | md }}

{{ t.text | md }}

{% endfor %}
{% if b.sources %}
## Sources

{% for s in b.sources %}
{{ s.n }}. {{ s.title | md }}, {{ s.publisher | md }}, {{ s.url | link }}, retrieved {{ s.date | md }}
{% endfor %}
{% if b.images %}

## Images on the cited pages

Links only; each image is the copyright of its owner.

{% for i in b.images %}
- {{ i.title | md }}, {{ i.publisher | md }}, {{ i.url | link }}
{% endfor %}
{% endif %}

## Quotes

{% for q in b.appendix %}
{% if q.mono %}
{{ q.quote | fence }}

Source [{{ q.n }}]

{% else %}
> \u201c{{ q.quote | md }}\u201d [{{ q.n }}]

{% endif %}
{% endfor %}
{% endif %}
{% if b.note %}
{{ b.note | md }}
{% endif %}
""")


def to_markdown(sections: list[dict], industry: dict | None = None) -> str:
    return MD.render(cases=[{**c, "mermaid": to_mermaid(a[0]) if (a := architecture_of(c)) else ""} for c in sections], industry=industry)


OFFICE = "application/vnd.openxmlformats-officedocument."
OUTPUTS = {"docx": (to_docx, OFFICE + "wordprocessingml.document"),
           "pptx": (to_pptx, OFFICE + "presentationml.presentation"),
           "md": (lambda s, i, images={}: to_markdown(s, i).encode(), "text/markdown; charset=utf-8"),
           # one format value per source keeps the form a single field
           "pdf_docx": (lambda s, i, images={}: to_pdf(to_docx(s, i, images), ".docx"), "application/pdf"),
           "pdf_pptx": (lambda s, i, images={}: to_pdf(to_pptx(s, i, images), ".pptx"), "application/pdf"),
           "case_study_docx": (to_case_study_docx, OFFICE + "wordprocessingml.document"),
           "case_study_pdf": (lambda s, i, images={}: to_pdf(to_case_study_docx(s, i, images), ".docx"), "application/pdf")}


LIBRARY_MAX = 3


def friendly(request, user, status, message, back):
    """A browser gets a portal page with a way back; an API client keeps the JSON error."""
    if wants_html(request):
        resp = page(request, "message.html", user, message=message, back=back)
        resp.status_code = status
        return resp
    raise HTTPException(status, message)


def build(conn, case_ids, research_id, industry_context_ack, user, fmt=None, library_ids=()):
    """The content of an output, shared by the download and the preview: (ids, sections, industry, anonymised, claims, library_used).
    Library items (#235) join the industry section (quotes) or render as briefs in industry["briefs"]; library_used lists those shown."""
    ids = list(dict.fromkeys(case_ids))
    if len(ids) > 3:
        raise HTTPException(400, "choose up to 3 cases")
    lib = list(dict.fromkeys(library_ids))
    if len(lib) > LIBRARY_MAX:
        raise HTTPException(400, f"choose up to {LIBRARY_MAX} library items")
    if not ids:  # industry context alone is never presented as our delivery without an explicit acknowledgement
        if research_id is None and not lib:
            raise HTTPException(400, "choose 1 to 3 cases, or a research result for an industry-context-only output")
        if not industry_context_ack:
            raise HTTPException(400, "an output with no case needs the acknowledgement that it contains public "
                                     "third-party statements only, not delivery evidence")
    if fmt is not None and fmt not in OUTPUTS:  # after the input checks, as on main
        raise HTTPException(400, f"format must be one of {', '.join(OUTPUTS)}")
    study = bool(fmt) and fmt.startswith("case_study")
    if study and len(ids) != 1:
        raise HTTPException(400, "a case study is one case")
    claims, pages = [], []
    if research_id is not None:  # only your own, finished research
        row = conn.execute("select results from research where id=%s and created_by=%s and status='done'",
                           (research_id, user.sub)).fetchone()
        if not row:
            raise HTTPException(404, "no such research")
        claims, pages = row[0].get("claims", []), row[0].get("pages", [])
    items = {r[0]: r for r in conn.execute("select id, kind, body from library_items where id = any(%s)", (lib,)).fetchall()} if lib else {}
    if len(items) != len(lib):
        raise HTTPException(404, "no such library item")  # every user may read every item, so an unknown id is the only miss
    # same restrictions as search: approved, in date, and the user can open the source document
    rows = conn.execute(
        "select c.id, c.data, cl.name, cl.anonymised_label, cl.referenceable, cl.id is not null, c.public_sources, c.member_count, c.approved_at "
        "from cases c left join clients cl on cl.id = c.client_id "
        f"where c.id = any(%s) and c.status = 'approved' and c.review_due > now() and {VISIBLE}",
        (ids, list(user.groups))).fetchall()
    by_id = {r[0]: r for r in rows}
    if len(by_id) != len(ids):
        raise HTTPException(404, "no such case")
    clients = anonymise.load_clients(conn)
    sections, anonymised, attached = [], False, []
    for i in ids:
        _, data, name, label, referenceable, linked, public_sources, members, approved = by_id[i]
        attached += public_sources
        shown = anonymise.shown(name, label, referenceable, linked)
        anonymised |= not (linked and referenceable) or anonymise.apply(shown, clients) != shown
        sec = section(ReferenceCase.model_validate(data), shown)
        if study:  # a role and a date, never a person; goes through protect() with the rest
            n = members or 1
            sec["footer"] = f"Based on {n} source document{'s' if n != 1 else ''} reviewed{f' on {approved:%Y-%m-%d}' if approved else ''}"
        if sec["architecture"]:
            sec["architecture"]["footer"] = f"Source: {sec['title']}" + (f" · reviewed {approved:%Y-%m-%d}" if approved else "")
            try:  # a blocked architecture is left out and counted on the preview; it never withholds the whole output
                protect([{"architecture": sec["architecture"]}], clients)
            except Withheld:
                sec["architecture"], sec["arch_withheld"] = None, True
        sections.append(sec)
    try:
        sections = protect(sections, clients)
    except Withheld as e:
        raise HTTPException(409, str(e)) from None
    from app import research  # here: research imports render's neighbours
    def key(c):  # one normal form for every claim, so the same quote from two places is one
        q = str(c.get("quote") or "")
        return c.get("url"), anonymise.apply(q if c.get("type") == "technical_note" else plain(q), clients)

    lib_claims, briefs, used, withheld, dropped = [], [], [], 0, 0
    for lid in lib:  # in the order chosen; text is re-checked now, whatever was true when it was saved
        _, kind, body = items[lid]
        ok = isinstance(body, dict) and isinstance(body.get("quotes"), list) and (kind == "quotes" or isinstance(body.get("brief"), dict))
        if kind not in ("quotes", "brief") or not ok:
            raise HTTPException(400, f"library item {lid} is damaged and cannot be used")
        if kind == "quotes":
            kept, n = research.visible_claims(body["quotes"], clients)
            dropped += n
            if (industry_section(kept, clients) or {}).get("statements"):  # used only if a quote survives to the page
                lib_claims += kept
                used.append(lid)
        else:
            try:
                briefs.append(brief_section(body["brief"], body["quotes"], clients))
                used.append(lid)
            except Withheld:
                withheld += 1
    seen, merged = set(), []
    for c in attached + claims + lib_claims:  # attached first: they win the slide's 5 places
        if (k := key(c)) not in seen:
            seen.add(k); merged.append(c)
    # the new heading only if an attached claim survives apply()/blocked() at render
    cited = bool(attached) and bool((industry_section(attached, clients, pages=pages) or {}).get("statements"))
    mine = {key(c) for c in attached}
    only_attached = not any(key(c) not in mine for c in claims + lib_claims)  # no ad-hoc research or library text merged in
    heading = (APPROACH_HEADING if study and only_attached else SOURCES_HEADING) if cited else INDUSTRY_HEADING
    industry = industry_section(merged, clients, heading=heading, pages=pages) \
        if merged else None
    if not ids and not (industry and industry["statements"]) and not briefs:
        raise HTTPException(400, "the research has no usable public statements")
    if briefs or withheld or dropped:
        industry = industry or {"heading": heading, "disclaimer": INDUSTRY_DISCLAIMER, "groups": [], "sources": [], "statements": [], "note": ""}
        for count, what in ((withheld, "library brief"), (dropped, "library quote")):
            if count:
                n = f"{count} {what}{'s' if count != 1 else ''} withheld: {'they' if count != 1 else 'it'} named a protected client"
                industry["note"] = f"{industry['note']}; {n}" if industry["note"] else n
    if industry:  # the heading and disclaimer come with quotes, or when there is no brief to stand alone
        industry["briefs"] = briefs
        industry["show_head"] = bool(industry["groups"]) or not (briefs or withheld or dropped)
    return ids, sections, industry, anonymised, claims, used


@router.post("/preview")
def preview(request: Request, case_ids: list[int] = Form([]), research_id: int | None = Form(None),
            industry_context_ack: bool = Form(False), library_ids: list[int] = Form([]), user: User = Depends(require("user"))):
    with db.connect() as conn:  # same content and errors as /generate; nothing is logged
        try:
            ids, sections, industry, _, _, _ = build(conn, case_ids, research_id, industry_context_ack, user, library_ids=library_ids)
        except HTTPException as e:
            if e.status_code not in (400, 404):
                raise
            return friendly(request, user, e.status_code, e.detail, "/search")
        footer, retitled = "", False  # what a case-study download adds; the same code path as /generate, so the text cannot drift
        if len(ids) == 1:
            _, study, study_industry, _, _, _ = build(conn, case_ids, research_id, industry_context_ack, user, "case_study_docx", library_ids)
            footer, retitled = study[0]["footer"], bool(study_industry) and study_industry["heading"] != (industry or {}).get("heading")
    for c in sections:  # a structure protect() could not keep intact is left out, and the page says so
        if c.get("architecture") and not architecture_of(c):
            c["architecture"], c["arch_withheld"] = None, True
    svgs = [Markup(to_svg(a[1], a[0].title)) if (a := architecture_of(c)) else "" for c in sections]  # Markup only after to_svg escaped every string
    return page(request, "output_preview.html", user, cases=sections, svgs=svgs, industry=industry, ids=ids,
                research_id=research_id, library_ids=list(dict.fromkeys(library_ids)), ack=industry_context_ack, footer=footer, retitled=APPROACH_HEADING if retitled else "")


def produce(conn, user: User, fmt: str, case_ids: list[int], research_id: int | None, ack: bool,
            library_ids: list[int]) -> tuple[bytes, str, str]:
    """Build, render and audit one output; raises HTTPException (400/404 from build, 503 from the PDF converter)."""
    ids, sections, industry, anonymised, claims, library_used = build(conn, case_ids, research_id, ack, user, fmt, library_ids)
    make, mime = OUTPUTS[fmt]
    pairs = [(ids[n], d["id"]) for n, sec in enumerate(sections) for d in sec["diagrams"]]  # (case, diagram)
    images = {i: bytes(png) for i, png in conn.execute(
        "select id, png from diagrams where (case_id, id) in (select * from unnest(%s::bigint[], %s::bigint[])) "
        "and png is not null", ([c for c, _ in pairs], [d for _, d in pairs])).fetchall()} if pairs else {}
    try:
        body = make(sections, industry, images=images)  # before the audit row: a broken template must not log a generation
    except PdfError as e:
        raise HTTPException(503, str(e)) from None
    conn.execute("insert into generations(user_id, format, case_ids, anonymised, industry_context_ack, research_id, library_ids) "
                 "values (%s,%s,%s,%s,%s,%s,%s)", (user.sub, fmt, ids, anonymised, ack and not ids,
                                                   research_id if (industry and claims) else None, library_used or None))
    name = f"{'case-study' if fmt.startswith('case_study') else 'reference-cases'}.{'pdf' if 'pdf' in fmt else fmt.rsplit('_', 1)[-1]}"
    return body, mime, name


@router.post("/generate")
def generate(request: Request, format: str = Form(), case_ids: list[int] = Form([]), research_id: int | None = Form(None),
             industry_context_ack: bool = Form(False), library_ids: list[int] = Form([]), user: User = Depends(require("user"))):
    with db.connect() as conn:
        try:
            body, mime, name = produce(conn, user, format, case_ids, research_id, industry_context_ack, library_ids)
        except HTTPException as e:
            if e.status_code not in (400, 404):
                raise
            return friendly(request, user, e.status_code, e.detail, "/search")
    return Response(body, media_type=mime, headers={"Content-Disposition": f'attachment; filename="{name}"', "X-Content-Type-Options": "nosniff"})


@router.get("/cases/{cid}/diagrams/{did}.png")
def diagram_png(cid: int, did: int, user: User = Depends(require("user"))):
    with db.connect() as conn:  # a ready diagram of a case /generate would also accept; raw bytes are never selected
        row = conn.execute(
            "select d.png from diagrams d join cases c on c.id = d.case_id where d.id = %s and c.id = %s "
            f"and c.status = 'approved' and c.review_due > now() and {VISIBLE} and d.png is not null "
            "and exists (select 1 from jsonb_array_elements(c.data->'diagrams') e "
            "where (e->>'id')::bigint = d.id and e->>'status' = 'ready')", (did, cid, list(user.groups))).fetchone()
    if not row:
        raise HTTPException(404, "no such diagram")
    return Response(bytes(row[0]), media_type="image/png", headers=PNG_HEADERS)


@router.get("/cases/{cid}/architecture.png")
def architecture_png(cid: int, user: User = Depends(require("user"))):
    """The architecture of one case at full size, through build(): approved, in date, visible, protected, never blocked."""
    with db.connect() as conn:
        try:
            _, sections, *_ = build(conn, [cid], None, False, user)
        except HTTPException as e:
            if e.status_code not in (400, 404, 409):
                raise
            sections = []
    a = architecture_of(sections[0]) if sections else None
    if not a:
        raise HTTPException(404, "no such diagram")  # the same answer for a hidden case, a missing one and a case with none
    return Response(to_png(a[1]), media_type="image/png", headers=PNG_HEADERS)  # no-store: a withdrawn case must not linger in a cache
