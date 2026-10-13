"""Layout and renderers for the solution architecture diagram (#246).

layout() turns an Architecture into a Scene: boxes, lines, icon shapes and text on a 1600x900 canvas, with
style keys instead of colours. Text is measured with the vendored IBM Plex Sans (brand/fonts), so the Scene is
the same on every machine, and text is wrapped, never cut: no "…" anywhere.
Flows are routed on an 8 px grid around the cards (shortest path, few bends, parallel flows on their own track).
"""
import heapq
import math
import re
from dataclasses import dataclass, field
from functools import lru_cache
from math import ceil
from pathlib import Path

from PIL import ImageFont

from app.schema import Architecture

W, H = 1600, 900
FONTS = Path(__file__).resolve().parent.parent / "brand" / "fonts"
WEIGHTS = {400: "Regular", 500: "Medium", 600: "SemiBold"}
# style key -> (size px, weight)
TEXT = {"title": (26, 600), "name": (15, 600), "role": (12, 400), "role-sm": (11, 400), "zone": (12, 600),
        "label": (11, 500), "callout": (13, 500), "legend": (11, 400), "footer": (11, 400), "badge": (12, 600),
        "chip": (13, 600), "note": (11, 400), "marker": (10, 600)}
LINE_H = {"title": 32, "name": 18, "role": 15, "role-sm": 14, "chip": 16, "note": 14, "callout": 17, "footer": 14, "label": 14}
LETTER = {"zone": 1.2}  # letter-spacing in px
ZONES = ["channels", "experience", "integration", "core", "data", "external"]  # left to right; platform is the bottom band

M = 24
SPAN = W - 2 * M
CARD_MIN_H, ROW_GAP, SUB_GAP, PAD, LABEL_H = 46, 22, 16, 14, 28
GMIN, GMAX, BAND_GAP, PILL_H, GRID = 80, 150, 30, 18, 8
BUDGET = 12000  # A* expansions per flow; 14 flows in a crowded zone must lay out in about a second


@dataclass
class Box:
    x: float
    y: float
    w: float
    h: float
    style: str
    ref: str = ""  # component id for cards, zone name for zones
    k: float = 1.0  # scale-to-fit factor for stroke widths and corner radii


@dataclass
class Line:
    points: list[tuple[float, float]]
    style: str
    arrow: bool = False
    refs: tuple[str, str] = ("", "")  # (src, dst) component ids
    k: float = 1.0


@dataclass
class Text:
    x: float
    y: float  # top of the line
    text: str
    style: str
    anchor: str = "start"  # start | middle
    k: float = 1.0  # font size factor


@dataclass
class Shape:
    """A piece of a kind icon, in canvas coordinates. fill/stroke: None, "w" (white) or "bg" (the icon square's colour)."""
    op: str  # e (ellipse x,y,w,h) | r (rect x,y,w,h) | p (polygon or polyline points)
    geom: list
    fill: str | None
    stroke: str | None
    kind: str
    closed: bool = True
    k: float = 1.0


@dataclass
class Scene:
    width: int = W
    height: int = H
    zones: list[Box] = field(default_factory=list)  # drawn first
    lines: list[Line] = field(default_factory=list)  # then the connectors, under the cards
    boxes: list[Box] = field(default_factory=list)
    shapes: list[Shape] = field(default_factory=list)
    texts: list[Text] = field(default_factory=list)
    area: tuple = (M, 0, W - M, H)  # where the diagram body was fitted
    scale: float = 1.0  # how far the body was scaled to fill it
    body: tuple = (0, 0, 0, 0)  # the fitted body's bounding box; header above it, footer below


# --- text ------------------------------------------------------------------------------------
@lru_cache(maxsize=None)
def _font_px(style: str, px: int):
    return ImageFont.truetype(str(FONTS / f"IBMPlexSans-{WEIGHTS[TEXT[style][1]]}.ttf"), px)


def _font(style: str):
    return _font_px(style, TEXT[style][0])


def measure(text: str, style: str) -> float:
    return _font(style).getlength(text) + LETTER.get(style, 0) * len(text)


def wrap(text: str, style: str, width: float) -> list[str]:
    """Greedy word wrap. A word wider than the line is broken. Never cuts or abbreviates."""
    lines, cur = [], ""
    for word in text.split():
        while measure(word, style) > width and len(word) > 1:
            n = len(word)
            while n > 1 and measure(word[:n], style) > width:
                n -= 1
            if cur:
                lines.append(cur)
                cur = ""
            lines.append(word[:n])
            word = word[n:]
        if not cur or measure(f"{cur} {word}", style) <= width:
            cur = f"{cur} {word}".strip()
        else:
            lines.append(cur)
            cur = word
    return [*lines, cur] if cur else lines


# --- icons (24x24, white on the kind colour) ----------------------------------------------------
def _arc(cx, cy, rx, ry, a0, a1, n=10):
    return [(cx + rx * math.cos(a0 + (a1 - a0) * i / n), cy + ry * math.sin(a0 + (a1 - a0) * i / n)) for i in range(n + 1)]


def _gear():
    pts = []
    for t in range(8):
        a = t * math.pi / 4
        for r, d in ((7, -.28), (9.6, -.17), (9.6, .17), (7, .28)):
            pts.append((12 + r * math.cos(a + d), 12 + r * math.sin(a + d)))
    return pts


ICONS = {  # kind -> [(op, geom, fill, stroke, closed)]
    "store": [("e", (6, 4, 12, 5), None, "w", True), ("p", [(6, 6.5), (6, 17)], None, "w", False), ("p", [(18, 6.5), (18, 17)], None, "w", False),
              ("p", _arc(12, 11, 6, 2.5, 0, math.pi), None, "w", False), ("p", _arc(12, 16.5, 6, 2.5, 0, math.pi), None, "w", False)],
    "queue": [("r", (4, 5, 16, 3.4), "w", None, True), ("r", (4, 10.3, 16, 3.4), "w", None, True), ("r", (4, 15.6, 16, 3.4), "w", None, True)],
    "user": [("e", (8.5, 3.5, 7, 7), "w", None, True), ("p", _arc(12, 20, 7.5, 8, math.pi, 2 * math.pi), "w", None, True)],
    "external": [("e", (3, 10, 9, 8), "w", None, True), ("e", (7, 5, 10, 10), "w", None, True), ("e", (12, 9, 9, 9), "w", None, True),
                 ("r", (7, 13, 10, 5), "w", None, True)],
    "service": [("p", _gear(), "w", None, True), ("e", (8.5, 8.5, 7, 7), "bg", None, True)],
    "app": [("r", (4, 5, 16, 14), None, "w", True), ("r", (4, 5, 16, 4.5), "w", None, True)],
    "platform": [("p", [(12, 3 + d), (21, 7 + d), (12, 11 + d), (3, 7 + d)], "bg", "w", True) for d in (8, 4, 0)],
}


def _icon(sc: Scene, x: float, y: float, kind: str):
    for op, g, fill, stroke, closed in ICONS[kind]:
        geom = [(x + a, y + b) for a, b in g] if op == "p" else [x + g[0], y + g[1], g[2], g[3]]
        sc.shapes.append(Shape(op, geom, fill, stroke, kind, closed))


# --- ordering ----------------------------------------------------------------------------------
def _order(arch: Architecture) -> dict[str, list[str]]:
    """Component ids per zone, ordered by barycentre of their neighbours' ranks (2 sweeps)."""
    order = {z: [c.id for c in arch.components if c.zone == z] for z in [*ZONES, "platform"]}
    nbrs: dict[str, set[str]] = {c.id: set() for c in arch.components}
    for f in arch.flows:
        nbrs[f.src].add(f.dst)
        nbrs[f.dst].add(f.src)
    zone_of = {c.id: c.zone for c in arch.components}
    for _ in range(2):
        for z, ids in order.items():
            rank = {i: (n + .5) / len(order[zone_of[i]]) for zone in order.values() for n, i in enumerate(zone)}
            bary = {i: sum(rank[j] for j in nbrs[i] if zone_of[j] != z) / k if (k := sum(zone_of[j] != z for j in nbrs[i])) else rank[i]
                    for i in ids}
            order[z] = sorted(ids, key=lambda i: (bary[i], ids.index(i)))
    return order


# --- router ------------------------------------------------------------------------------------
STEPS = ((1, 0), (0, 1), (-1, 0), (0, -1))  # right, down, left, up


class _Grid:
    """Orthogonal shortest paths on an 8 px grid around the obstacles. Used edges cost more, so parallel flows take their own track."""

    def __init__(self, obstacles, soft, bounds):
        self.bx0, self.by0 = math.ceil(bounds[0] / GRID), math.ceil(bounds[1] / GRID)
        self.bx1, self.by1 = int(bounds[2] // GRID), int(bounds[3] // GRID)
        self.blocked = set()
        for x0, y0, x1, y1 in obstacles:  # nodes strictly inside the rectangle grown by 3 px
            for ix in range(int((x0 - 3) // GRID) + 1, math.ceil((x1 + 3) / GRID)):
                for iy in range(int((y0 - 3) // GRID) + 1, math.ceil((y1 + 3) / GRID)):
                    self.blocked.add((ix, iy))
        self.soft = self._nodes(soft)
        self.used: set = set()
        self.cross: set = set()

    @staticmethod
    def _nodes(rects):
        return {(ix, iy) for x0, y0, x1, y1 in rects
                for ix in range(math.ceil(x0 / GRID), int(x1 // GRID) + 1) for iy in range(math.ceil(y0 / GRID), int(y1 // GRID) + 1)}

    def route(self, start, sdir, end, edir, box=None, avoid=(), budget=BUDGET):
        """start/end: grid nodes; sdir the outward direction at the start, edir the inward one at the end (0 right, 1 down, 2 left, 3 up).
        `box` limits the path to a rectangle in px, `avoid` lists rectangles in px it must stay out of. Gives up after `budget`
        expansions (a crowded diagram): None, and the caller draws a plain elbow. Returns the nodes or None."""
        x0, y0, x1, y1 = self.bx0, self.by0, self.bx1, self.by1
        if box:
            x0, y0, x1, y1 = max(x0, math.ceil(box[0] / GRID)), max(y0, math.ceil(box[1] / GRID)), min(x1, int(box[2] // GRID)), min(y1, int(box[3] // GRID))
        bad = self._nodes(avoid)
        back = (edir + 2) % 4
        heap = [(0, 0, 0, start[0], start[1], sdir)]
        best = {(start[0], start[1], sdir): 0}
        prev: dict = {}
        tick = 0
        while heap and budget:
            budget -= 1
            _, c, _, ix, iy, d = heapq.heappop(heap)
            if c > best.get((ix, iy, d), 1e18):
                continue
            if (ix, iy) == end and d != back:
                path, key = [(ix, iy)], (ix, iy, d)
                while key in prev:
                    key = prev[key]
                    path.append(key[:2])
                return path[::-1]
            for nd in (d, (d + 1) % 4, (d + 3) % 4):
                nx, ny = ix + STEPS[nd][0], iy + STEPS[nd][1]
                if not (x0 <= nx <= x1 and y0 <= ny <= y1) or (nx, ny) in self.blocked or (nx, ny) in bad:
                    continue
                edge = (min(ix, nx), min(iy, ny), nd % 2)
                nc = c + 1 + (6 if nd != d else 0) + (2 if (nx, ny) in self.soft else 0) + (20 if edge in self.used else 0) \
                    + (4 if (nx, ny, 1 - nd % 2) in self.cross else 0)
                if nc < best.get((nx, ny, nd), 1e18):
                    best[(nx, ny, nd)] = nc
                    prev[(nx, ny, nd)] = (ix, iy, d)
                    tick += 1
                    heapq.heappush(heap, (nc + 2 * (abs(end[0] - nx) + abs(end[1] - ny)), nc, tick, nx, ny, nd))  # weighted A*: a little off the optimum, much faster
        return None

    def mark(self, nodes):
        for (ax, ay), (bx, by) in zip(nodes, nodes[1:]):
            o = 0 if ay == by else 1
            self.used.add((min(ax, bx), min(ay, by), o))
            self.cross.update({(ax, ay, o), (bx, by, o)})


def _simplify(pts):
    out = [pts[0]]
    for p in pts[1:]:
        if p != out[-1]:
            out.append(p)
    i = 1
    while i < len(out) - 1:
        a, b, c = out[i - 1], out[i], out[i + 1]
        if a[0] == b[0] == c[0] or a[1] == b[1] == c[1]:
            del out[i]
        else:
            i += 1
    return out


def seg_hits_rect(a, b, r, margin=0.5) -> bool:
    """An axis-aligned segment passes through the inside of rect r=(x, y, w, h) shrunk by `margin`."""
    x0, y0, x1, y1 = r[0] + margin, r[1] + margin, r[0] + r[2] - margin, r[1] + r[3] - margin
    if a[0] == b[0]:
        return x0 < a[0] < x1 and min(a[1], b[1]) < y1 and max(a[1], b[1]) > y0
    return y0 < a[1] < y1 and min(a[0], b[0]) < x1 and max(a[0], b[0]) > x0


def _rects_hit(a, b, pad=0) -> bool:
    return not (a[0] + a[2] + pad <= b[0] or b[0] + b[2] + pad <= a[0] or a[1] + a[3] + pad <= b[1] or b[1] + b[3] + pad <= a[1])


# --- layout ------------------------------------------------------------------------------------
def _chips(arch, comp, plat):
    """Platform components as compact chips, packed in rows. Returns (placed, band height, self-flow lane height)."""
    plat_set = set(plat)
    notes: dict[str, list[str]] = {i: [] for i in plat}
    for f in arch.flows:
        for p, o in ((f.src, f.dst), (f.dst, f.src)):
            if p in plat_set and o not in plat_set:
                notes[p].append(f"Used by {comp[o].name}" + (f": {f.label}" if f.label else ""))
    n_self = sum(f.src in plat_set and f.dst in plat_set for f in arch.flows)
    items = []
    for i in plat:
        c = comp[i]
        name = wrap(c.name, "chip", 210)
        rest = [ln for t in ([c.role] if c.role else []) + notes[i] for ln in wrap(t, "note", 210)]
        tw = max([measure(x, "chip") for x in name] + [measure(x, "note") for x in rest])
        items.append((i, 52 + tw, 10 + 16 * len(name) + 14 * len(rest) + 10, name, rest))
    rows, cur, cw = [], [], 0.0
    for it in items:
        if cur and cw + it[1] + 14 > SPAN - 2 * PAD:
            rows.append(cur)
            cur, cw = [], 0.0
        cur.append(it)
        cw += it[1] + 14
    if cur:
        rows.append(cur)
    placed, y = [], LABEL_H
    for row in rows:
        x = M + (SPAN - (sum(it[1] for it in row) + 14 * (len(row) - 1))) / 2
        rh = max(44, max(it[2] for it in row))
        for i, w, h, name, rest in row:
            placed.append((i, x, y, w, rh, name, rest))
            x += w + 14
        y += rh + 36  # room for a label pill between the rows
    lane = 8 * min(n_self, 4) + 12 if n_self else 0
    return placed, y - 36 + lane + PAD, lane


def _config(comp, order, cols, rsize, w, subs, limits):
    rstyle = "role" if rsize == 12 else "role-sm"
    lines = {i: (wrap(comp[i].name, "name", w - 52), wrap(comp[i].role, rstyle, w - 52)) for z in cols for i in order[z]}
    if limits and any(len(n) > 2 or len(r) > 3 for n, r in lines.values()):
        return None
    rows = {z: ceil(len(order[z]) / subs[z]) for z in cols}
    # a row is as tall as its tallest card's text, and every card in that row (in any zone) has that height
    rowh = [max([CARD_MIN_H] + [10 + 18 * len(lines[i][0]) + LINE_H[rstyle] * len(lines[i][1]) + 10
                                for z in cols for n, i in enumerate(order[z]) if n % rows[z] == r]) for r in range(max(rows.values(), default=0))]
    zh = {z: LABEL_H + sum(rowh[:rows[z]]) + (rows[z] - 1) * ROW_GAP + PAD for z in cols}
    colw = {z: subs[z] * w + (subs[z] - 1) * SUB_GAP + 2 * PAD for z in cols}
    k = len(cols)
    return dict(w=w, rstyle=rstyle, rowh=rowh, subs=subs, rows=rows, zh=zh, colw=colw, lines=lines, band_h=max(zh.values(), default=0),
                G=min(GMAX, max(0, (SPAN - sum(colw.values())) / (k - 1))) if k > 1 else 0)


def _plan(comp, order, cols, avail, plat_h, res_top, gap):
    """Card width, role size and sub-columns: the widest cards first, the role font last, sub-columns only when a zone is too tall."""
    k = len(cols)
    cands = [(12, w, True) for w in range(300, 229, -10)] + [(11, w, True) for w in range(300, 229, -10)] + \
            [(11, w, False) for w in range(229, 129, -10)]
    for rsize, w, limits in cands:
        subs = {z: 1 for z in cols}
        while (cfg := _config(comp, order, cols, rsize, w, subs, limits)):
            if k > 1 and sum(cfg["colw"].values()) + (k - 1) * GMIN > SPAN or sum(cfg["colw"].values()) > SPAN:
                break
            if res_top + cfg["band_h"] + gap + plat_h <= avail:
                return cfg
            tall = [z for z in cols if cfg["rows"][z] > 1]
            if not tall:
                break
            subs[max(tall, key=lambda z: cfg["zh"][z])] += 1
    return None


_SCENES: dict = {}  # (architecture JSON, footer) -> Scene; layout is slow enough to want a cache, but the architecture may be unvalidated


def layout(arch: Architecture, footer: str = "") -> Scene:
    key = (arch.model_dump_json(), footer)
    if key not in _SCENES:
        if len(_SCENES) >= 64:
            _SCENES.clear()
        _SCENES[key] = _cached(arch, footer)
    return _SCENES[key]


def _key_rows(labels):
    out, kx, ky = [], M, 0
    for n, label, mode in labels:
        wd = 24 + measure(label, "label") + 24
        if kx + wd > W - M and kx > M:
            kx, ky = M, ky + 20
        out.append((n, label, mode, kx, ky))
        kx += wd
    return out, (ky + 26 if out else 0)


def _cached(arch: Architecture, footer: str) -> Scene:
    reserve: list = []
    for _ in range(3):
        sc, markers = _build(arch, footer, reserve)
        if _key_rows(markers)[1] <= _key_rows(reserve)[1]:
            return sc
        reserve = markers
    return sc


def _fit_body(sc: Scene, i0: tuple, area: tuple, cap: float = 1.6) -> None:
    """Scale the diagram body (zones, flows, cards, icons, their text) uniformly to fill `area`, centred; the header and footer keep their size."""
    z, ln, bx, sh, tx = i0
    zones, lines, boxes, shapes, texts = sc.zones[z:], sc.lines[ln:], sc.boxes[bx:], sc.shapes[sh:], sc.texts[tx:]
    xs = [b.x for b in zones + boxes] + [b.x + b.w for b in zones + boxes] + [p[0] for l in lines for p in l.points]
    ys = [b.y for b in zones + boxes] + [b.y + b.h for b in zones + boxes] + [p[1] for l in lines for p in l.points]
    if not xs:
        return
    bx0, by0, bx1, by1 = min(xs), min(ys), max(xs), max(ys)
    s = min(cap, (area[2] - area[0]) / (bx1 - bx0), (area[3] - area[1]) / (by1 - by0))
    ox = area[0] + ((area[2] - area[0]) - (bx1 - bx0) * s) / 2 - bx0 * s
    oy = area[1] + ((area[3] - area[1]) - (by1 - by0) * s) / 2 - by0 * s
    tr = lambda x, y: (x * s + ox, y * s + oy)
    for b in zones + boxes:
        b.x, b.y = tr(b.x, b.y)
        b.w, b.h, b.k = b.w * s, b.h * s, s
    for l in lines:
        l.points, l.k = [tr(*p) for p in l.points], s
    for sp in shapes:
        sp.k = s
        sp.geom = [tr(*p) for p in sp.geom] if sp.op == "p" else [*tr(sp.geom[0], sp.geom[1]), sp.geom[2] * s, sp.geom[3] * s]
    for t in texts:
        t.x, t.y = tr(t.x, t.y)
        t.k = s
    sc.area, sc.scale = area, s
    sc.body = (area[0] + ((area[2] - area[0]) - (bx1 - bx0) * s) / 2, oy + by0 * s, area[2] - ((area[2] - area[0]) - (bx1 - bx0) * s) / 2, oy + by1 * s)


def _build(arch: Architecture, footer: str, reserve: list) -> tuple[Scene, list]:
    sc = Scene()
    comp = {c.id: c for c in arch.components}
    order = _order(arch)
    cols = [z for z in ZONES if order[z]]
    plat = order["platform"]
    plat_set = set(plat)
    k = len(cols)
    cidx = {i: n for n, z in enumerate(cols) for i in order[z]}
    flows = list(arch.flows)
    drawn = [f for f in flows if not ((f.src in plat_set) ^ (f.dst in plat_set))]  # a flow to a platform chip is a note on the chip

    # header: wrapped title left, numbered outcome callouts top-right
    title = wrap(arch.title, "title", 1060 - M)
    head = 24 + 32 * len(title) + 6
    for n, ln in enumerate(title):
        sc.texts.append(Text(M, 20 + 32 * n, ln, "title"))
    y = 12
    for co in arch.callouts[:3]:
        lines = wrap(co.text, "callout", 410)
        h = 14 + 17 * len(lines)
        sc.boxes.append(Box(1116, y, 460, h, "callout"))
        sc.boxes.append(Box(1124, y + (h - 22) / 2, 22, 22, "badge"))
        sc.texts.append(Text(1135, y + (h - 22) / 2 + 4, str(co.outcome + 1), "badge", "middle"))
        for n, ln in enumerate(lines):
            sc.texts.append(Text(1156, y + 7 + 17 * n, ln, "callout"))
        y += h + 4
    head = max(head, y + 8)

    # footer and legend
    lx = W - M
    legend = []
    for m in reversed(sorted({f.mode for f in drawn})):
        lx -= measure(m, "legend")
        legend.append(Text(lx, 0, m, "legend"))
        lx -= 44
        legend.append(Line([(lx, 7), (lx + 36, 7)], f"flow-{m}", True))
        lx -= 18
    foot = wrap(footer, "footer", lx - M - 16) if footer else []
    foot_top = H - 12 - 14 * max(1, len(foot))
    for it in legend:
        if isinstance(it, Text):
            it.y = foot_top
        else:
            it.points = [(a, b + foot_top) for a, b in it.points]
        (sc.texts if isinstance(it, Text) else sc.lines).append(it)
    for n, ln in enumerate(foot):
        sc.texts.append(Text(M, foot_top + 14 * n, ln, "footer"))
    key_h = _key_rows(reserve)[1]
    bottom = foot_top - 10 - key_h

    body0 = (len(sc.zones), len(sc.lines), len(sc.boxes), len(sc.shapes), len(sc.texts))
    # platform chips first (their height is needed), then the zones
    chips, plat_h, _ = _chips(arch, comp, plat) if plat else ([], 0, 0)
    nlong = sum(1 for f in drawn if f.src in cidx and f.dst in cidx and abs(cidx[f.src] - cidx[f.dst]) >= 2)
    res_top = 8 * ceil(nlong / 2) + 12 if nlong else 0
    gap = max(BAND_GAP, res_top) if plat else res_top
    cfg = _plan(comp, order, cols, bottom - head, plat_h, res_top, gap)
    if cfg is None:  # nothing fits: the smallest cards, in as many sub-columns as the height needs
        subs = {z: max(1, ceil(len(order[z]) / 3)) for z in cols}
        cfg = _config(comp, order, cols, 11, 130, subs, False) or _config(comp, order, cols, 11, 130, {z: 1 for z in cols}, False)
    w, G, rstyle = cfg["w"], cfg["G"], cfg["rstyle"]
    y0 = head + max(0, (bottom - head - (res_top + cfg["band_h"] + gap + plat_h)) / 2)
    band_y = y0 + res_top
    plat_y = band_y + cfg["band_h"] + gap
    x = M + (SPAN - (sum(cfg["colw"].values()) + G * max(k - 1, 0))) / 2
    card: dict[str, Box] = {}
    sub_of: dict[str, tuple[int, int]] = {}
    for z in cols:
        zy = band_y  # every column zone has the tallest zone's height; cards sit at the top
        sc.zones.append(Box(x, zy, cfg["colw"][z], cfg["band_h"], "zone", z))
        sc.texts.append(Text(x + PAD, zy + 8, z.upper(), "zone"))
        for n, i in enumerate(order[z]):
            s, r = divmod(n, cfg["rows"][z])
            card[i] = Box(x + PAD + s * (w + SUB_GAP), zy + LABEL_H + sum(cfg["rowh"][:r]) + r * ROW_GAP, w, cfg["rowh"][r], f"card-{comp[i].kind}", i)
            sub_of[i] = (s, r)
        x += cfg["colw"][z] + G
    pbox = None
    if plat:
        pbox = Box(M, plat_y, SPAN, plat_h, "zone", "platform")
        sc.zones.append(pbox)
        sc.texts.append(Text(M + PAD, plat_y + 8, "PLATFORM", "zone"))
        for i, cx, cy, cw, ch, name, rest in chips:
            card[i] = Box(cx, plat_y + cy, cw, ch, "card-platform", i)
            sub_of[i] = (0, 0)
    for i, b in card.items():
        sc.boxes.append(b)
        sc.boxes.append(Box(b.x + 10, b.y + 10, 24, 24, f"glyph-{comp[i].kind}"))
        _icon(sc, b.x + 10, b.y + 10, comp[i].kind)
    for z in cols:
        for i in order[z]:
            b, (name, role) = card[i], cfg["lines"][i]
            sc.texts += [Text(b.x + 42, b.y + 10 + 18 * n, ln, "name") for n, ln in enumerate(name)]
            sc.texts += [Text(b.x + 42, b.y + 10 + 18 * len(name) + LINE_H[rstyle] * n, ln, rstyle) for n, ln in enumerate(role)]
    for i, cx, cy, cw, ch, name, rest in chips:
        b = card[i]
        sc.texts += [Text(b.x + 42, b.y + 10 + 16 * n, ln, "chip") for n, ln in enumerate(name)]
        sc.texts += [Text(b.x + 42, b.y + 10 + 16 * len(name) + 14 * n, ln, "note") for n, ln in enumerate(rest)]

    # --- flows: sides, ports, routes ---------------------------------------------------------------
    side: dict[tuple[int, str], str] = {}
    for n, f in enumerate(drawn):
        s, d = card[f.src], card[f.dst]
        if f.src in plat_set:  # between platform chips
            side[n, "s"], side[n, "d"] = ("b", "b") if abs(s.y - d.y) < 1 else (("b", "t") if d.y > s.y else ("t", "b"))
        elif comp[f.src].zone == comp[f.dst].zone:
            (ss, rs), (sd, rd) = sub_of[f.src], sub_of[f.dst]
            last = cfg["subs"][comp[f.src].zone] - 1
            if ss == sd and abs(rs - rd) == 1:
                side[n, "s"], side[n, "d"] = ("b", "t") if rd > rs else ("t", "b")
            elif ss == sd:  # around the column's outer side
                side[n, "s"] = side[n, "d"] = "l" if ss == 0 and last > 0 else "r"
            else:
                side[n, "s"], side[n, "d"] = ("r", "l") if sd > ss else ("l", "r")
        else:
            side[n, "s"], side[n, "d"] = ("r", "l") if cidx[f.dst] > cidx[f.src] else ("l", "r")
    users: dict[tuple[str, str], list[tuple[float, int, str]]] = {}
    for (n, e), sd in side.items():
        f = drawn[n]
        i, o = (f.src, card[f.dst]) if e == "s" else (f.dst, card[f.src])
        users.setdefault((i, sd), []).append(((o.y + o.h / 2) if sd in "rl" else (o.x + o.w / 2), n, e))
    ports: dict[tuple[int, str], tuple[float, float]] = {}
    for (i, sd), lst in users.items():
        b = card[i]
        lo, hi = (b.y, b.y + b.h) if sd in "rl" else (b.x, b.x + b.w)
        slots = list(range(math.ceil((lo + 10) / GRID) * GRID, int(hi - 10) + 1, GRID)) or [round((lo + hi) / 2 / GRID) * GRID]
        lst.sort()
        for p, (_, n, e) in enumerate(lst):
            q = min(len(slots) - 1, round((p + 1) * (len(slots) + 1) / (len(lst) + 1)) - 1) if len(lst) <= len(slots) else p % len(slots)
            v = slots[max(0, q)]
            ports[n, e] = {"r": (b.x + b.w, v), "l": (b.x, v), "t": (v, b.y), "b": (v, b.y + b.h)}[sd]

    grid = _Grid([(b.x, b.y, b.x + b.w, b.y + b.h) for b in card.values()], [(z.x, z.y, z.x + z.w, z.y + z.h) for z in sc.zones],
                 (M + 4, head, W - M - 4, bottom))
    band = (pbox.x, pbox.y, pbox.x + pbox.w, pbox.y + pbox.h) if pbox else None
    DIRS = {"r": 0, "b": 1, "l": 2, "t": 3}

    def stub(p, sd):
        x, y = p
        return {"r": (math.ceil((x + 4) / GRID), round(y / GRID)), "l": (math.floor((x - 4) / GRID), round(y / GRID)),
                "b": (round(x / GRID), math.ceil((y + 4) / GRID)), "t": (round(x / GRID), math.floor((y - 4) / GRID))}[sd]

    def dist(n):
        (sx, sy), (dx, dy) = ports[n, "s"], ports[n, "d"]
        return abs(sx - dx) + abs(sy - dy)

    routes: dict[int, list] = {}
    for n in sorted(range(len(drawn)), key=lambda n: (dist(n), n)):
        f, sd, dd = drawn[n], side[n, "s"], side[n, "d"]
        sp, dp = ports[n, "s"], ports[n, "d"]
        a, b = stub(sp, sd), stub(dp, dd)
        into = (DIRS[dd] + 2) % 4
        path = grid.route(a, DIRS[sd], b, into, box=band) if f.src in plat_set else grid.route(a, DIRS[sd], b, into, avoid=[band] if band else ())
        if path is None:  # no clear path: a plain elbow
            pts = [sp, (a[0] * GRID, a[1] * GRID), (a[0] * GRID, b[1] * GRID), (b[0] * GRID, b[1] * GRID), dp]
        else:
            grid.mark(path)
            pts = [sp, *[(ix * GRID, iy * GRID) for ix, iy in path], dp]
        routes[n] = _simplify(pts)
    for n, f in enumerate(drawn):
        sc.lines.append(Line(routes[n], f"flow-{f.mode}", True, (f.src, f.dst)))

    # --- labels: a pill on the best straight stretch of each route, else a numbered marker keyed along the bottom
    cardr = [(b.x, b.y, b.w, b.h) for b in card.values()]
    placed: list[tuple] = []
    markers: list[tuple[int, str, str]] = []
    allsegs = [(a, b, n) for n, pts in routes.items() for a, b in zip(pts, pts[1:])]

    def place(n, w, h, need_long, pad=2):
        segs = sorted(zip(routes[n], routes[n][1:]), key=lambda s: -(abs(s[0][0] - s[1][0]) + abs(s[0][1] - s[1][1])))
        ok = []
        for a, b in segs:
            ln = abs(a[0] - b[0]) + abs(a[1] - b[1])
            if need_long and ln < w + 8:
                continue
            for t in (.5, .4, .6, .3, .7, .2, .8, .1, .9):
                cx, cy = a[0] + (b[0] - a[0]) * t, a[1] + (b[1] - a[1]) * t
                r = (cx - w / 2, cy - h / 2, w, h)
                if not (M <= r[0] and r[0] + w <= W - M and head - 6 <= r[1] and r[1] + h <= foot_top - key_h - 4):
                    continue
                if any(_rects_hit(r, c, pad) for c in cardr) or any(_rects_hit(r, p, pad) for p in placed):
                    continue
                ok.append((sum(1 for p, q, m in allsegs if m != n and seg_hits_rect(p, q, r, 0)), -ln, r))
        return min(ok, key=lambda o: o[:2])[2] if ok else None

    for n, f in enumerate(drawn):
        if not f.label:
            continue
        one = measure(f.label, "label") + 14
        two = wrap(f.label, "label", max(60, one / 2 + 12))
        tries = [([f.label], one, PILL_H)]
        if len(two) == 2:
            tries.append((two, max(measure(x, "label") for x in two) + 14, 32))
        r = lines = None
        for need_long in (True, False):
            for lines, w, h in tries:
                r = place(n, w, h, need_long)
                if r:
                    break
            if r:
                break
        if r:
            placed.append(r)
            sc.boxes.append(Box(*r, f"pill-{f.mode}"))
            sc.texts += [Text(r[0] + r[2] / 2, r[1] + (r[3] - 14 * len(lines)) / 2 + 14 * j + 1, ln, "label", "middle") for j, ln in enumerate(lines)]
            continue
        r = place(n, 18, 18, False) or place(n, 14, 14, False, 1) or place(n, 12, 12, False, 0)
        if r is None:  # no clear spot: sit on the middle of the longest stretch
            a, b = max(zip(routes[n], routes[n][1:]), key=lambda s: abs(s[0][0] - s[1][0]) + abs(s[0][1] - s[1][1]))
            r = ((a[0] + b[0]) / 2 - 9, (a[1] + b[1]) / 2 - 9, 18, 18)
        placed.append(r)
        markers.append((len(markers) + 1, f.label, f.mode))
        sc.boxes.append(Box(*r, f"marker-{f.mode}"))
        sc.texts.append(Text(r[0] + r[2] / 2, r[1] + r[3] / 2 - 5, str(len(markers)), "marker", "middle"))
    _fit_body(sc, body0, (M, head, W - M, bottom))
    rows, kh = _key_rows(markers)
    for n, label, mode, kx, ky in rows:
        kyy = foot_top - 10 - kh + ky + 3
        sc.boxes.append(Box(kx, kyy, 18, 18, f"marker-{mode}"))
        sc.texts.append(Text(kx + 9, kyy + 4, str(n), "marker", "middle"))
        sc.texts.append(Text(kx + 24, kyy + 2, label, "label"))
    return sc, markers


# --- renderers -------------------------------------------------------------------------------
# brand palette (app/static/portal.css)
INK, GROUND, ACCENT, BORDER, MUTED, WHITE = "#1B2430", "#F6F5F1", "#0B6A6A", "#DAD7CF", "#55606E", "#FFFFFF"
KIND = {"user": "#3D4652", "app": "#0B6A6A", "service": "#1F5F8B", "queue": "#8A5A00", "store": "#5B4B8A",
        "external": "#7A6F5F", "platform": "#3D4652"}
MODE = {"api": "#0B6A6A", "event": "#8A5A00", "batch": "#5B4B8A", "file": "#7A6F5F", "ui": "#1F5F8B", "data": "#1E6B3A"}
DASH = {"event": (3, 4), "batch": (9, 5), "file": (9, 5)}
TEXT_COLOUR = {"role": MUTED, "role-sm": MUTED, "zone": MUTED, "legend": MUTED, "footer": MUTED, "note": MUTED, "badge": WHITE, "marker": WHITE}
LINE_W = 1.5


def box_paint(style: str) -> tuple[str, str, int]:
    """(fill, stroke, corner radius) for a box style key."""
    kind, _, sub = style.partition("-")
    return {"zone": (GROUND, BORDER, 12), "card": (WHITE, KIND.get(sub, INK), 8), "glyph": (KIND.get(sub, INK),) * 2 + (5,),
            "pill": (WHITE, MODE.get(sub, INK), 9), "callout": ("#E3F0F0", ACCENT, 8), "badge": (ACCENT, ACCENT, 11),
            "marker": (MODE.get(sub, INK), MODE.get(sub, INK), 9)}[kind]


def _arrow(pts, k=1.0):
    (x1, y1), (x2, y2) = pts[-2], pts[-1]
    d = math.dist(pts[-2], pts[-1]) or 1
    ux, uy = (x2 - x1) / d, (y2 - y1) / d
    bx, by = x2 - ux * 9 * k, y2 - uy * 9 * k
    return [(x2, y2), (bx - uy * 4.5 * k, by + ux * 4.5 * k), (bx + uy * 4.5 * k, by - ux * 4.5 * k)]


def _clean(pts):
    out = [pts[0]]
    for p in pts[1:]:
        if p != out[-1]:
            out.append(p)
    return out


def _round(pts, r=6, n=6):
    """The polyline with its corners rounded (radius r), as a denser polyline."""
    out = [pts[0]]
    for a, b, c in zip(pts, pts[1:], pts[2:]):
        la, lc = math.dist(a, b), math.dist(b, c)
        rr = min(r, la / 2, lc / 2)
        p1 = (b[0] + (a[0] - b[0]) / la * rr, b[1] + (a[1] - b[1]) / la * rr)
        p2 = (b[0] + (c[0] - b[0]) / lc * rr, b[1] + (c[1] - b[1]) / lc * rr)
        for i in range(n + 1):
            t = i / n
            out.append(tuple((1 - t) ** 2 * p1[j] + 2 * t * (1 - t) * b[j] + t * t * p2[j] for j in (0, 1)))
    out.append(pts[-1])
    return out


def _tok(tok, kind):
    return {"w": WHITE, "bg": KIND.get(kind, INK)}.get(tok)


def _text_box(t: Text):
    w = measure(t.text, t.style) * t.k
    x0 = t.x - w / 2 if t.anchor == "middle" else t.x
    return x0, t.y, x0 + w, t.y + TEXT[t.style][0] * 1.3 * t.k


def cropped(sc: Scene, margin: int = 32) -> Scene:
    """A copy for documents and the preview: the empty bands around the body are cut out (the footer moves up under it) and the
    canvas is trimmed to the content plus `margin`. The scene itself keeps the full 16:9 canvas, which the pptx slide fills."""
    import copy
    sc = copy.deepcopy(sc)
    top, bottom = sc.area[1], sc.area[3]
    if sc.body[3] > sc.body[1]:
        up = sc.body[1] - top  # the gap between the header and the body
        foot = up + (bottom - sc.body[3])  # and the one between the body and the footer
        def shift(y):  # a y below the body area belongs to the footer; one above the header stays
            return y - (foot if y >= bottom else up if y >= top else 0)
        for b in sc.zones + sc.boxes:
            b.y = shift(b.y)
        for t in sc.texts:
            t.y = shift(t.y)
        for ln in sc.lines:
            ln.points = [(x, shift(y)) for x, y in ln.points]
        for sp in sc.shapes:
            sp.geom = [(x, shift(y)) for x, y in sp.geom] if sp.op == "p" else [sp.geom[0], shift(sp.geom[1]), *sp.geom[2:]]
    xs, ys = [], []
    for b in sc.zones + sc.boxes:
        xs += [b.x, b.x + b.w]
        ys += [b.y, b.y + b.h]
    for ln in sc.lines:
        xs += [p[0] for p in ln.points]
        ys += [p[1] for p in ln.points]
    for t in sc.texts:
        x0, y0, x1, y1 = _text_box(t)
        xs += [x0, x1]
        ys += [y0, y1]
    x0, y0 = max(0, min(xs) - margin), max(0, min(ys) - margin)
    x1, y1 = min(sc.width, max(xs) + margin), min(sc.height, max(ys) + margin)
    for b in sc.zones + sc.boxes:
        b.x, b.y = b.x - x0, b.y - y0
    for t in sc.texts:
        t.x, t.y = t.x - x0, t.y - y0
    for ln in sc.lines:
        ln.points = [(x - x0, y - y0) for x, y in ln.points]
    for sp in sc.shapes:
        sp.geom = [(x - x0, y - y0) for x, y in sp.geom] if sp.op == "p" else [sp.geom[0] - x0, sp.geom[1] - y0, *sp.geom[2:]]
    sc.width, sc.height = math.ceil(x1 - x0), math.ceil(y1 - y0)
    return sc


def to_svg(sc: Scene, title: str = "Solution architecture", crop: bool = True) -> str:
    from xml.sax.saxutils import escape
    sc = cropped(sc) if crop else sc
    f = lambda v: f"{v:.1f}".rstrip("0").rstrip(".")
    pp = lambda pts: " ".join(f"{f(x)},{f(y)}" for x, y in pts)
    o = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {sc.width} {sc.height}" role="img" '
         f'font-family="IBM Plex Sans, system-ui, sans-serif"><title>{escape(title)}</title>'
         f'<rect width="{sc.width}" height="{sc.height}" fill="{WHITE}"/>']

    def box(b: Box):
        fill, stroke, r = box_paint(b.style)
        o.append(f'<rect x="{f(b.x)}" y="{f(b.y)}" width="{f(b.w)}" height="{f(b.h)}" rx="{f(r * b.k)}" fill="{fill}" stroke="{stroke}" stroke-width="{f(1.5 * b.k)}"/>')

    for b in sc.zones:
        box(b)
    for ln in sc.lines:
        pts = _clean(ln.points)
        if len(pts) < 2:
            continue
        mode = ln.style.split("-", 1)[1]
        dash = DASH.get(mode)
        o.append(f'<polyline points="{pp(_round(pts, 6 * ln.k))}" fill="none" stroke="{MODE[mode]}" stroke-width="{f(LINE_W * ln.k)}" stroke-linejoin="round"'
                 + (f' stroke-dasharray="{f(dash[0] * ln.k)} {f(dash[1] * ln.k)}"' if dash else "") + "/>")
        if ln.arrow:
            o.append(f'<polygon points="{pp(_arrow(pts, ln.k))}" fill="{MODE[mode]}"/>')
    for b in sc.boxes:
        box(b)
    for s in sc.shapes:
        fill, stroke = _tok(s.fill, s.kind), _tok(s.stroke, s.kind)
        paint = f'fill="{fill or "none"}"' + (f' stroke="{stroke}" stroke-width="{f(1.6 * s.k)}" stroke-linejoin="round"' if stroke else "")
        if s.op == "e":
            x, y, w, h = s.geom
            o.append(f'<ellipse cx="{f(x + w / 2)}" cy="{f(y + h / 2)}" rx="{f(w / 2)}" ry="{f(h / 2)}" {paint}/>')
        elif s.op == "r":
            x, y, w, h = s.geom
            o.append(f'<rect x="{f(x)}" y="{f(y)}" width="{f(w)}" height="{f(h)}" rx="1" {paint}/>')
        else:
            o.append(f'<{"polygon" if s.closed else "polyline"} points="{pp(s.geom)}" {paint}/>')
    for t in sc.texts:
        size, weight = TEXT[t.style]
        o.append(f'<text x="{f(t.x)}" y="{f(t.y)}" font-size="{f(size * t.k)}" font-weight="{weight}" fill="{TEXT_COLOUR.get(t.style, INK)}" '
                 f'dominant-baseline="hanging"' + (f' letter-spacing="{f(LETTER[t.style] * t.k)}"' if t.style in LETTER else "")
                 + (' text-anchor="middle"' if t.anchor == "middle" else "") + f">{escape(t.text)}</text>")
    o.append("</svg>")
    return "".join(o)


def to_png(sc: Scene, scale: int = 2, crop: bool = True) -> bytes:
    import io
    sc = cropped(sc) if crop else sc
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (sc.width * scale, sc.height * scale), WHITE)
    d = ImageDraw.Draw(im)
    S = lambda pts: [(x * scale, y * scale) for x, y in pts]

    def box(b: Box):
        fill, stroke, r = box_paint(b.style)
        d.rounded_rectangle([b.x * scale, b.y * scale, (b.x + b.w) * scale, (b.y + b.h) * scale], r * b.k * scale, fill, stroke, max(1, int(1.5 * b.k * scale)))

    for b in sc.zones:
        box(b)
    for ln in sc.lines:
        pts = _clean(ln.points)
        if len(pts) < 2:
            continue
        mode = ln.style.split("-", 1)[1]
        dash = DASH.get(mode)
        dense = _round(pts, 6 * ln.k)
        if not dash:
            d.line(S(dense), MODE[mode], max(1, int(LINE_W * ln.k * scale)), joint="curve")
        else:  # dashes walked along the rounded polyline
            at, on = 0.0, True
            for (x1, y1), (x2, y2) in zip(dense, dense[1:]):
                n = math.dist((x1, y1), (x2, y2))
                pos = 0.0
                while pos < n:
                    seg = min(dash[0 if on else 1] * ln.k - at, n - pos)
                    if on:
                        d.line(S([(x1 + (x2 - x1) * pos / n, y1 + (y2 - y1) * pos / n), (x1 + (x2 - x1) * (pos + seg) / n, y1 + (y2 - y1) * (pos + seg) / n)]),
                               MODE[mode], max(1, int(LINE_W * ln.k * scale)))
                    pos += seg
                    at += seg
                    if at >= dash[0 if on else 1] * ln.k - 1e-9:
                        on, at = not on, 0.0
        if ln.arrow:
            d.polygon(S(_arrow(pts, ln.k)), MODE[mode])
    for b in sc.boxes:
        box(b)
    for s in sc.shapes:
        fill, stroke = _tok(s.fill, s.kind), _tok(s.stroke, s.kind)
        lw = max(1, int(1.6 * s.k * scale))
        if s.op in "er":
            x, y, ww, hh = s.geom
            (d.ellipse if s.op == "e" else d.rectangle)([x * scale, y * scale, (x + ww) * scale, (y + hh) * scale], fill, stroke, lw)
        elif s.closed:
            d.polygon(S(s.geom), fill, stroke)
            if stroke:
                d.line(S([*s.geom, s.geom[0]]), stroke, lw, joint="curve")
        else:
            d.line(S(s.geom), stroke, lw, joint="curve")
    for t in sc.texts:
        size, weight = TEXT[t.style]
        font = _font_px(t.style, max(1, round(size * t.k * scale)))
        col = TEXT_COLOUR.get(t.style, INK)
        if t.style in LETTER:
            x = t.x * scale
            for ch in t.text:
                d.text((x, t.y * scale), ch, col, font, anchor="la")
                x += font.getlength(ch) + LETTER[t.style] * t.k * scale
        else:
            d.text((t.x * scale, t.y * scale), t.text, col, font, anchor="ma" if t.anchor == "middle" else "la")
    out = io.BytesIO()
    im.save(out, "PNG")
    return out.getvalue()


def to_pptx_shapes(slide, sc: Scene, left: int, top: int, width: int, height: int) -> None:
    """Native shapes: rounded rectangles, one freeform polyline per flow (arrow at the end), icon shapes and text boxes."""
    from lxml import etree
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt

    s = min(width / sc.width, height / sc.height)
    X, Y = (lambda v: int(left + v * s)), (lambda v: int(top + v * s))
    rgb = lambda h: RGBColor.from_string(h.lstrip("#"))

    def style(sh, fill, stroke, lw, name):  # lw in canvas px, already scaled by the caller
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if stroke:
            sh.line.color.rgb = rgb(stroke)
            sh.line.width = Emu(int(lw * s))
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        sh.name = name

    def box(b: Box):
        fill, stroke, r = box_paint(b.style)
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(b.x), Y(b.y), int(b.w * s), int(b.h * s))
        sh.adjustments[0] = min(.5, r / min(b.w, b.h))
        style(sh, fill, stroke, 1.5 * b.k, f"arch:{b.style}:{b.ref}" if b.ref else f"arch:{b.style}")

    def poly(pts, closed):
        fb = slide.shapes.build_freeform(X(pts[0][0]), Y(pts[0][1]))
        fb.add_line_segments([(X(x), Y(y)) for x, y in pts[1:]], close=closed)
        return fb.convert_to_shape()

    for b in sc.zones:
        box(b)
    for ln in sc.lines:
        pts = _clean(ln.points)
        if len(pts) < 2:
            continue
        mode = ln.style.split("-", 1)[1]
        sh = poly(_round(pts, 6 * ln.k), False)
        style(sh, None, MODE[mode], LINE_W * ln.k, f"arch:{ln.style}")
        if mode in DASH:
            sh.line.dash_style = MSO_LINE.DASH if DASH[mode][0] > 3 else MSO_LINE.ROUND_DOT
        if ln.arrow:
            etree.SubElement(sh.line._get_or_add_ln(), "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", type="triangle")
    for b in sc.boxes:
        box(b)
    for sp in sc.shapes:
        fill, stroke = _tok(sp.fill, sp.kind), _tok(sp.stroke, sp.kind)
        if sp.op in "er":
            x, y, w, h = sp.geom
            sh = slide.shapes.add_shape(MSO_SHAPE.OVAL if sp.op == "e" else MSO_SHAPE.RECTANGLE, X(x), Y(y), int(w * s), int(h * s))
        else:
            sh = poly(sp.geom, sp.closed)
        style(sh, fill if (sp.closed or sp.op != "p") else None, stroke, 1.6 * sp.k, f"arch:icon-{sp.kind}")
    for t in sc.texts:
        size, weight = TEXT[t.style]
        w = int((max(measure(t.text, t.style) * t.k, 10) + 6) * s)
        tb = slide.shapes.add_textbox(X(t.x) - (w // 2 if t.anchor == "middle" else 0), Y(t.y), w, int(size * t.k * 1.4 * s))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if t.anchor == "middle" else PP_ALIGN.LEFT
        r = p.add_run()
        r.text = t.text
        r.font.name = "IBM Plex Sans"
        r.font.size = Pt(size * t.k * s / 12700)
        r.font.bold = weight >= 600
        r.font.color.rgb = rgb(TEXT_COLOUR.get(t.style, INK))
        if t.style in LETTER:
            r._r.get_or_add_rPr().set("spc", str(int(LETTER[t.style] * t.k * s / 12700 * 100)))
        tb.name = f"arch:text:{t.style}"


def to_mermaid(arch: Architecture) -> str:
    esc = lambda s: re.sub(r'["\[\]|<>`]', lambda m: f"#{ord(m.group())};", " ".join(s.split()))  # no fence breaker, no line break
    ids = {c.id: f"n{n}" for n, c in enumerate(arch.components)}  # by position: "a-b" and "a_b" stay two nodes
    nid = ids.__getitem__
    out = ["flowchart LR"]
    for z in [*ZONES, "platform"]:
        cs = [c for c in arch.components if c.zone == z]
        if cs:
            out.append(f'  subgraph {z}["{z.capitalize()}"]')
            out += [f'    {nid(c.id)}["{esc(c.name)}' + (f'<br/>{esc(c.role)}' if c.role else "") + '"]' for c in cs]
            out.append("  end")
    for f in arch.flows:
        arrow = "-.->" if f.mode in DASH else "-->"
        out.append(f'  {nid(f.src)} {arrow}' + (f'|"{esc(f.label)}"|' if f.label else "") + f" {nid(f.dst)}")
    return "\n".join(out) + "\n"
